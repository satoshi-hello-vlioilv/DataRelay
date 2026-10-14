"""M 言語（Power Query の式）を読んで評価する（2.4.0。lib/navi_m.py・navi_mlib.py・navi_mconv.py・navi_msource.py）。

期待値は Microsoft の Power Query M 関数リファレンスの例と、Excel が画面操作で書き出す式の形に合わせる。
本物の Excel では動かせないので、「貼り付けた式がそのまま動く」ことを、Excel が実際に書く形（日本語版の
ステップ名・補助クエリ・meta 付きのパラメーター）で確かめる。

実行: python -m unittest tests.test_mquery -v
"""
import math
import sys
import tempfile
import unittest
from datetime import date, datetime, time, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'lib'))
import navi_m as M  # noqa: E402
import navi_mconv as C  # noqa: E402


def ev(src, **kw):
    """式を1本評価して、表ならセルを値に（誤りのセルは 'ERR'）、ほかはそのまま返す。"""
    v = M.Document.from_text(src).evaluate(host_=M.Host(**kw))
    return plain(v)


def plain(v):
    v = M.force(v) if not isinstance(v, M.ErrorCell) else v
    if isinstance(v, M.ErrorCell):
        return 'ERR'
    if isinstance(v, M.MTable):
        return {'columns': v.columns, 'rows': [[plain(x) for x in r] for r in v.rows]}
    if isinstance(v, M.MRecord):
        return {n: plain(v.raw(n)) for n in v.names}
    if isinstance(v, list):
        return [plain(x) for x in v]
    return v


class LanguageTest(unittest.TestCase):
    def test_literals_and_operators(self):
        self.assertEqual(ev('1 + 2 * 3'), 7)
        self.assertEqual(ev('(1 + 2) * 3'), 9)
        self.assertEqual(ev('6 / 3'), 2, '整数で表せる商は整数')
        self.assertEqual(ev('7 / 2'), 3.5)
        self.assertEqual(ev('1 / 0'), math.inf, '0 で割ると無限大（M の決まり）')
        self.assertEqual(ev('"a" & "b"'), 'ab')
        self.assertEqual(ev('{1, 2} & {3}'), [1, 2, 3])
        self.assertEqual(ev('[a = 1] & [b = 2]'), {'a': 1, 'b': 2})
        self.assertEqual(ev('"say ""hi""#(cr,lf)"'), 'say "hi"\r\n', '"" と #(cr,lf) の書き方')
        self.assertEqual(ev('0x1F'), 31)
        self.assertEqual(ev('1.5e2'), 150)
        self.assertIsNone(ev('null + 1'), 'null との計算は null')
        self.assertIsNone(ev('null & "a"'))
        self.assertEqual(ev('{1..4}'), [1, 2, 3, 4])
        self.assertEqual(ev('{"a".."d"}'), ['a', 'b', 'c', 'd'], '文字の範囲（列の分割で使う）')

    def test_comparison_and_logic(self):
        self.assertTrue(ev('1 = 1.0'))
        self.assertFalse(ev('1 = "1"'), '種類が違えば等しくない（誤りにはしない）')
        self.assertTrue(ev('null = null'))
        self.assertIsNone(ev('null > 1'), '大小は null になる')
        self.assertTrue(ev('"a" < "b"'))
        self.assertTrue(ev('#date(2024,1,2) > #date(2024,1,1)'))
        self.assertIsNone(ev('null and true'))
        self.assertFalse(ev('null and false'))
        self.assertTrue(ev('null or true'))
        self.assertTrue(ev('not false'))
        self.assertEqual(ev('null ?? 5'), 5)
        self.assertTrue(ev('false or (1 / 0 = #infinity)'))
        self.assertFalse(ev('false and error "使われない"'), 'and は右を読まずに決まれば読まない')

    def test_let_is_lazy(self):
        self.assertEqual(ev('let a = error "使わない", b = 2 in b'), 2, '使わないステップの誤りでは止まらない')
        self.assertEqual(ev('let b = a + 1, a = 1 in b'), 2, '後ろのステップを参照できる')
        with self.assertRaises(M.MError) as cm:
            ev('let a = b, b = a in a')
        self.assertIn('循環', str(cm.exception))

    def test_functions_and_each(self):
        self.assertEqual(ev('let f = (x, optional y) => x + (y ?? 10) in {f(1), f(1, 2)}'), [11, 3])
        self.assertEqual(ev('let f = (x as number) as number => x * 2 in f(4)'), 8)
        self.assertEqual(ev('List.Transform({1, 2, 3}, each _ * 10)'), [10, 20, 30])
        self.assertEqual(ev('List.Select({[a=1],[a=5]}, each [a] > 2)'), [{'a': 5}], '[a] は _[a]')
        self.assertEqual(ev('let fact = (n) => if n <= 1 then 1 else n * @fact(n - 1) in fact(5)'), 120)
        with self.assertRaises(M.MError) as cm:
            ev('let f = (x) => x in f(1, 2)')
        self.assertIn('引数', str(cm.exception))

    def test_records_and_lists(self):
        self.assertEqual(ev('[a = 1, b = a + 1][b]'), 2, 'レコードのフィールドは互いを参照できる')
        self.assertEqual(ev('[#"A b" = 1, Date accessed = 2][Date accessed]'), 2, '空白入りの名前')
        self.assertIsNone(ev('[a = 1][b]?'))
        self.assertEqual(ev('[a = 1, b = 2, c = 3][[a], [c]]'), {'a': 1, 'c': 3})
        self.assertEqual(ev('{10, 20, 30}{1}'), 20)
        self.assertIsNone(ev('{10}{5}?'))
        with self.assertRaises(M.MError):
            ev('{10}{5}')
        self.assertEqual(ev('[Attributes = [Hidden = true]][Attributes]?[Hidden]?'), True)
        self.assertIsNone(ev('[x = 1][Attributes]?[Hidden]?'), 'null に ? を続けても null')

    def test_try_and_errors(self):
        self.assertEqual(ev('try 1 / "a" otherwise 0'), 0)
        r = ev('try error "だめ"')
        self.assertTrue(r['HasError'])
        self.assertEqual(r['Error']['Message'], 'だめ')
        self.assertEqual(ev('try 5'), {'HasError': False, 'Value': 5})
        self.assertEqual(ev('try error [Reason = "R", Message = "M"] catch (e) => e[Reason]'), 'R')
        self.assertEqual(ev('(try Number.From("x"))[Error][Reason]'), 'DataFormat.Error')

    def test_types(self):
        self.assertTrue(ev('1 is number'))
        self.assertTrue(ev('null is nullable text'))
        self.assertFalse(ev('"a" is number'))
        self.assertEqual(ev('Value.Is(Date.From("2024/1/2"), type date)'), True)
        with self.assertRaises(M.MError):
            ev('"a" as number')
        self.assertEqual(ev('Table.ColumnNames(#table(type table [A = text, #"B c" = Int64.Type], {{"x", 1}}))'),
                         ['A', 'B c'])

    def test_meta_is_ignored(self):
        self.assertEqual(ev('"C:\\data" meta [IsParameterQuery = true, Type = "Text"]'), 'C:\\data')

    def test_syntax_errors_point_to_line(self):
        with self.assertRaises(M.MSyntaxError) as cm:
            M.read_queries('let\n    a = 1,\n    b = (2\nin\n    b')
        self.assertEqual(cm.exception.line, 4, '閉じていない括弧の次の行')
        with self.assertRaises(M.MSyntaxError):
            M.read_queries('')

    def test_comments(self):
        self.assertEqual(ev('let\n  // 説明\n  a = 1 /* 中 */ + 2\nin\n  a'), 3)

    def test_unknown_function_message(self):
        with self.assertRaises(M.MError) as cm:
            ev('Table.NoSuchThing(1)')
        self.assertIn('まだ使えません', str(cm.exception))
        with self.assertRaises(M.MError) as cm:
            ev('let a = 1 in b')
        self.assertIn('名前「b」', str(cm.exception))

    def test_error_location(self):
        with self.assertRaises(M.MError) as cm:
            M.Document.from_text('let\n a = 1,\n 変換 = Text.Upper(a)\nin\n 変換', 'Q').evaluate()
        self.assertEqual(cm.exception.step, '変換')
        self.assertEqual(cm.exception.query, 'Q')
        self.assertIn('ステップ「変換」', str(cm.exception))


class DocumentTest(unittest.TestCase):
    def test_comment_separated_queries(self):
        doc = M.Document.from_text('// 基準\n10\n\n// 計算\nlet\n    a = 基準 * 2\nin\n    a\n')
        self.assertEqual(doc.names(), ['基準', '計算'])
        self.assertEqual(doc.main_name(), '計算', 'ほかから参照されていない最後のクエリが主')
        self.assertEqual(doc.evaluate(), 20)

    def test_main_skips_excel_helper_queries(self):
        """Excel のエディターと同じ5本（見本用の「サンプル ファイルの変換」も、どこからも参照されない）。出すのはメイン。"""
        text = ('// 日報\nlet a = ファイルの変換(#"サンプル ファイル") in a\n\n// ファイルの変換\nlet ソース = (p as number) => p + 1 in ソース\n\n'
                '// サンプル ファイル\n1\n\n// パラメーター1\n#"サンプル ファイル" meta [IsParameterQuery=true]\n\n'
                '// サンプル ファイルの変換\nlet x = パラメーター1 in x\n')
        doc = M.Document.from_text(text)
        self.assertEqual(doc.main_name(), '日報')
        self.assertEqual(doc.helpers(), ['ファイルの変換', 'サンプル ファイル', 'パラメーター1', 'サンプル ファイルの変換'])
        for n in ('Transform Sample File', 'サンプル ファイルの変換 (2)', 'Parameter2', 'ファイルの変換 から 日報'):
            self.assertTrue(M.is_helper_name(n), n)
        self.assertFalse(M.is_helper_name('日報の集約'))

    def test_section_document(self):
        doc = M.Document.from_text('section Section1;\n\nshared A = 1;\nshared #"B 2" = A + 1;\n')
        self.assertEqual(doc.names(), ['A', 'B 2'])
        self.assertEqual(doc.evaluate('B 2'), 2)

    def test_single_query_keeps_header_comment_as_name(self):
        doc = M.Document.from_text('// 売上\nlet a = 1 in a')
        self.assertEqual(doc.names(), ['売上'])

    def test_step_spans_and_kinds(self):
        text = '// 主\nlet\n    ソース = 補助(1),\n    #"足す" = ソース + 1\nin\n    #"足す"\n\n// 補助\nlet ソース = (x as number) => x * 2 in ソース\n\n// 値\n"a" meta [IsParameterQuery=true]\n'
        doc = M.Document.from_text(text)
        spans = doc.step_spans('主')
        self.assertEqual([text[a:b] for _n, a, b in spans], ['補助(1)', 'ソース + 1'], '数式バーに出す式の場所（全体の中の位置）')
        self.assertEqual([doc.kind(n) for n in doc.names()], ['query', 'function', 'parameter'])
        sec = 'section S;\nshared A = 1;\nshared B = let x = A in x;'
        d2 = M.Document.from_text(sec)
        self.assertEqual([sec[a:b] for _n, a, b in d2.step_spans('B')], ['A'])

    def test_missing_names(self):
        """メインのクエリだけを貼ると、補助のクエリ（ファイルの変換 など）が無い。貼った時点で名前を挙げて知らせる。"""
        doc = M.Document.from_text('let\n ソース = Folder.Files("x"),\n 呼び = Table.AddColumn(ソース, "f", each ファイルの変換([Content])),\n'
                                   ' 名 = Table.ColumnNames(ファイルの変換(#"サンプル ファイル")),\n r = [a = 1, b = a],'
                                   ' f = (id) => id + 1\nin\n 名')
        self.assertEqual(doc.missing_names(), ['サンプル ファイル', 'ファイルの変換'], '関数・ステップ・引数・フィールド・each の _ は数えない')

    def test_steps(self):
        doc = M.Document.from_text('let\n ソース = {1,2,3},\n #"絞り込み" = List.Select(ソース, each _ > 1)\nin\n #"絞り込み"')
        self.assertEqual(doc.steps(doc.main_name()), ['ソース', '絞り込み'])
        self.assertEqual(plain(doc.evaluate(step='ソース')), [1, 2, 3])


T = '#table({"A","B"}, {{1,"x"},{2,"y"},{3,null},{2,"y"}})'


class TableTest(unittest.TestCase):
    def test_select_remove_rename_reorder(self):
        self.assertEqual(ev(f'Table.SelectRows({T}, each [A] > 1)')['rows'], [[2, 'y'], [3, None], [2, 'y']])
        self.assertEqual(ev(f'Table.SelectRows({T}, each [B] <> null)')['rows'], [[1, 'x'], [2, 'y'], [2, 'y']])
        self.assertEqual(ev(f'Table.RemoveColumns({T}, {{"A"}})')['columns'], ['B'])
        self.assertEqual(ev(f'Table.SelectColumns({T}, {{"B", "A"}})')['columns'], ['B', 'A'], '指定の順に並ぶ')
        self.assertEqual(ev(f'Table.SelectColumns({T}, {{"A", "Z"}}, MissingField.UseNull)')['columns'], ['A', 'Z'])
        self.assertEqual(ev(f'Table.RenameColumns({T}, {{{{"A", "番号"}}}})')['columns'], ['番号', 'B'])
        self.assertEqual(ev(f'Table.RenameColumns({T}, {{"A", "番号"}})')['columns'], ['番号', 'B'], '1組だけの書き方')
        self.assertEqual(ev(f'Table.ReorderColumns({T}, {{"B", "A"}})')['columns'], ['B', 'A'])
        with self.assertRaises(M.MError) as cm:
            ev(f'Table.RemoveColumns({T}, {{"Z"}})')
        self.assertIn('列「Z」', str(cm.exception))
        self.assertEqual(ev(f'Table.RemoveColumns({T}, {{"Z"}}, MissingField.Ignore)')['columns'], ['A', 'B'])

    def test_rows(self):
        self.assertEqual(ev(f'Table.FirstN({T}, 2)')['rows'], [[1, 'x'], [2, 'y']])
        self.assertEqual(ev(f'Table.Skip({T}, 3)')['rows'], [[2, 'y']])
        self.assertEqual(ev(f'Table.RemoveFirstN({T})')['rows'][0], [2, 'y'])
        self.assertEqual(ev(f'Table.RemoveLastN({T}, 2)')['rows'], [[1, 'x'], [2, 'y']])
        self.assertEqual(ev(f'Table.Distinct({T})')['rows'], [[1, 'x'], [2, 'y'], [3, None]])
        self.assertEqual(ev(f'Table.Distinct({T}, {{"B"}})')['rows'], [[1, 'x'], [2, 'y'], [3, None]])
        self.assertEqual(ev(f'Table.RowCount({T})'), 4)
        self.assertEqual(ev(f'Table.Range({T}, 1, 2)')['rows'], [[2, 'y'], [3, None]])
        self.assertEqual(ev(f'{T}{{1}}[B]'), 'y')
        self.assertEqual(ev(f'{T}{{[A=3]}}[B]'), None, 'レコードで1行を探す')
        self.assertEqual(ev(f'{T}[A]'), [1, 2, 3, 2], '列はリスト')
        with self.assertRaises(M.MError) as cm:
            ev(f'{T}{{[A=2]}}')
        self.assertIn('複数', str(cm.exception))

    def test_sort(self):
        t = '#table({"A","B"}, {{2,"b"},{1,"z"},{2,"a"},{null,"n"}})'
        self.assertEqual(ev(f'Table.Sort({t}, {{{{"A", Order.Descending}}, {{"B", Order.Ascending}}}})')['rows'],
                         [[2, 'a'], [2, 'b'], [1, 'z'], [None, 'n']])
        self.assertEqual(ev(f'Table.Sort({t}, "B")')['rows'][0], [2, 'a'])
        self.assertEqual(ev(f'Table.Sort({t}, {{"A", Order.Ascending}})')['rows'][0], [None, 'n'], 'null が先頭')
        self.assertEqual(ev(f'Table.Sort({t}, each [B])')['rows'][0], [2, 'a'])

    def test_add_and_transform_columns(self):
        r = ev(f'Table.AddColumn({T}, "C", each [A] * 10, type number)')
        self.assertEqual(r['columns'], ['A', 'B', 'C'])
        self.assertEqual([x[2] for x in r['rows']], [10, 20, 30, 20])
        r = ev(f'Table.AddColumn({T}, "C", each Text.Upper([B]) & "!")')
        self.assertEqual([x[2] for x in r['rows']], ['X!', 'Y!', None, 'Y!'])
        r = ev(f'Table.AddColumn({T}, "C", each 1 + [B])')
        self.assertEqual(r['rows'][0][2], 'ERR', '1行の失敗はセルの誤り（全体は止めない）')
        r = ev(f'Table.TransformColumns({T}, {{{{"B", Text.Upper, type text}}, {{"A", each _ + 1}}}})')
        self.assertEqual(r['rows'][0], [2, 'X'])
        r = ev(f'Table.AddIndexColumn({T}, "No", 1, 1, Int64.Type)')
        self.assertEqual([x[2] for x in r['rows']], [1, 2, 3, 4])
        self.assertEqual(ev(f'Table.DuplicateColumn({T}, "A", "A2")')['rows'][0], [1, 'x', 1])

    def test_transform_column_types(self):
        t = '#table({"N","D","T","I"}, {{"1,234.5","2024/01/05","12:30","2.5"},{"x","2024-02-29 08:00","",null}})'
        r = ev(f'Table.TransformColumnTypes({t}, {{{{"N", type number}}, {{"D", type date}}, {{"T", type time}}, '
               f'{{"I", Int64.Type}}}})')
        self.assertEqual(r['rows'][0], [1234.5, date(2024, 1, 5), time(12, 30), 2], '2.5 は偶数へ丸めて 2')
        self.assertEqual(r['rows'][1][0], 'ERR', '読めない値はセルの誤り')
        self.assertEqual(r['rows'][1][1], date(2024, 2, 29))
        self.assertEqual(r['rows'][1][3], None)
        r = ev('Table.TransformColumnTypes(#table({"N"}, {{"1.234,5"}}), {{"N", type number}}, "de-DE")')
        self.assertEqual(r['rows'][0][0], 1234.5, '文化圏で小数点が変わる')
        r = ev('Table.TransformColumnTypes(#table({"D"}, {{"01/02/2024"}}), {{"D", type date}}, "en-US")')
        self.assertEqual(r['rows'][0][0], date(2024, 1, 2))
        r = ev('Table.TransformColumnTypes(#table({"A"}, {{1}}), {{"Z", type text}}, [MissingField = MissingField.Ignore])')
        self.assertEqual(r['columns'], ['A'])

    def test_errors_rows(self):
        t = 'Table.TransformColumnTypes(#table({"A","B"}, {{"1","a"},{"x","b"}}), {{"A", Int64.Type}})'
        self.assertEqual(ev(f'Table.RemoveRowsWithErrors({t})')['rows'], [[1, 'a']])
        self.assertEqual(ev(f'Table.SelectRowsWithErrors({t}, {{"A"}})')['rows'][0][1], 'b')
        self.assertEqual(ev(f'Table.ReplaceErrorValues({t}, {{{{"A", 0}}}})')['rows'], [[1, 'a'], [0, 'b']])

    def test_promote_and_demote(self):
        t = '#table({"Column1","Column2","Column3"}, {{"日付", 2024, null},{"a", 1, 2}})'
        self.assertEqual(ev(f'Table.PromoteHeaders({t}, [PromoteAllScalars=true])')['columns'], ['日付', '2024', 'Column3'])
        self.assertEqual(ev(f'Table.PromoteHeaders({t})')['columns'], ['日付', 'Column2', 'Column3'],
                         'PromoteAllScalars が無ければ文字だけ見出しにする')
        self.assertEqual(ev('Table.PromoteHeaders(#table(3, {{"a","a","a"}}))')['columns'], ['a', 'a_1', 'a_2'])
        self.assertEqual(ev(f'Table.DemoteHeaders({T})')['rows'][0], ['A', 'B'])

    def test_replace_fill(self):
        t = '#table({"A","B"}, {{"a-b", 1},{null, 2},{"c", null}})'
        self.assertEqual(ev(f'Table.ReplaceValue({t}, "-", "_", Replacer.ReplaceText, {{"A"}})')['rows'][0][0], 'a_b')
        self.assertEqual(ev(f'Table.ReplaceValue({t}, null, 0, Replacer.ReplaceValue, {{"B"}})')['rows'][2][1], 0)
        self.assertEqual(ev(f'Table.ReplaceValue({t}, 1, 9, Replacer.ReplaceText, {{"B"}})')['rows'][0][1], 1,
                         'ReplaceText は文字にだけ効く')
        self.assertEqual([r[0] for r in ev(f'Table.FillDown({t}, {{"A"}})')['rows']], ['a-b', 'a-b', 'c'])
        self.assertEqual([r[1] for r in ev(f'Table.FillUp({t}, {{"B"}})')['rows']], [1, 2, None])
        r = ev(f'Table.ReplaceValue({t}, each [A], each [A] & "!", Replacer.ReplaceValue, {{"A"}})')
        self.assertEqual(r['rows'][0][0], 'a-b!', '置き換える値に each を使う形')

    def test_split_and_combine(self):
        t = '#table({"K"}, {{"A-1-x"},{"B-2"},{null}})'
        r = ev(f'Table.SplitColumn({t}, "K", Splitter.SplitTextByDelimiter("-", QuoteStyle.Csv), {{"K.1", "K.2"}})')
        self.assertEqual(r['columns'], ['K.1', 'K.2'])
        self.assertEqual(r['rows'], [['A', '1'], ['B', '2'], [None, None]])
        r = ev(f'Table.SplitColumn({t}, "K", Splitter.SplitTextByDelimiter("-"), 3)')
        self.assertEqual(r['rows'][1], ['B', '2', None])
        r = ev('Table.SplitColumn(#table({"K"},{{"AB123"}}), "K", Splitter.SplitTextByCharacterTransition('
               '(c) => not List.Contains({"0".."9"}, c), {"0".."9"}), {"K.1","K.2"})')
        self.assertEqual(r['rows'][0], ['AB', '123'])
        r = ev('Table.SplitColumn(#table({"K"},{{"abcdef"}}), "K", Splitter.SplitTextByPositions({0, 2, 4}), 3)')
        self.assertEqual(r['rows'][0], ['ab', 'cd', 'ef'])
        r = ev('Table.CombineColumns(#table({"A","B","C"},{{"x","y",1}}), {"A","B"}, '
               'Combiner.CombineTextByDelimiter("/", QuoteStyle.None), "AB")')
        self.assertEqual(r['columns'], ['AB', 'C'])
        self.assertEqual(r['rows'][0], ['x/y', 1])

    def test_group(self):
        t = '#table({"K","V"}, {{"a",1},{"b",2},{"a",3},{"A",4}})'
        r = ev(f'Table.Group({t}, {{"K"}}, {{{{"合計", each List.Sum([V]), type number}}, '
               f'{{"件数", each Table.RowCount(_), Int64.Type}}}})')
        self.assertEqual(r['rows'], [['a', 4, 2], ['b', 2, 1], ['A', 4, 1]])
        r = ev(f'Table.Group({t}, "K", {{"n", each Table.RowCount(_)}}, GroupKind.Global, Comparer.OrdinalIgnoreCase)')
        self.assertEqual(r['rows'], [['a', 3], ['b', 1]])
        r = ev(f'Table.Group({t}, "K", {{"n", each Table.RowCount(_)}}, GroupKind.Local)')
        self.assertEqual(len(r['rows']), 4, 'Local は連続する行だけをまとめる')

    def test_unpivot_pivot(self):
        t = '#table({"店","1月","2月"}, {{"A",10,null},{"B",30,40}})'
        r = ev(f'Table.UnpivotOtherColumns({t}, {{"店"}}, "属性", "値")')
        self.assertEqual(r['columns'], ['店', '属性', '値'])
        self.assertEqual(r['rows'], [['A', '1月', 10], ['B', '1月', 30], ['B', '2月', 40]], 'null は出さない')
        self.assertEqual(ev(f'Table.Unpivot({t}, {{"1月"}}, "属性", "値")')['rows'][0], ['A', None, '1月', 10])
        r = ev(f'let u = Table.UnpivotOtherColumns({t}, {{"店"}}, "属性", "値") in '
               f'Table.Pivot(u, List.Distinct(u[属性]), "属性", "値", List.Sum)')
        self.assertEqual(r['columns'], ['店', '1月', '2月'])
        self.assertEqual(r['rows'], [['A', 10, None], ['B', 30, 40]])
        r = ev('Table.Pivot(#table({"k","a","v"},{{1,"x",1},{1,"x",2}}), {"x"}, "a", "v")')
        self.assertEqual(r['rows'][0][1], 'ERR', '集計の関数なしで値が2つならセルの誤り')

    def test_join_and_expand(self):
        a = '#table({"id","name"}, {{1,"a"},{2,"b"},{3,"c"}})'
        b = '#table({"key","qty"}, {{1,10},{1,11},{3,30},{9,90}})'
        r = ev(f'Table.ExpandTableColumn(Table.NestedJoin({a}, {{"id"}}, {b}, {{"key"}}, "B", JoinKind.LeftOuter), '
               f'"B", {{"qty"}}, {{"B.qty"}})')
        self.assertEqual(r['columns'], ['id', 'name', 'B.qty'])
        self.assertEqual(r['rows'], [[1, 'a', 10], [1, 'a', 11], [2, 'b', None], [3, 'c', 30]])
        r = ev(f'Table.NestedJoin({a}, "id", {b}, "key", "B", JoinKind.Inner)')
        self.assertEqual([x[0] for x in r['rows']], [1, 3])
        r = ev(f'Table.NestedJoin({a}, "id", {b}, "key", "B", JoinKind.LeftAnti)')
        self.assertEqual([x[0] for x in r['rows']], [2])
        r = ev(f'Table.NestedJoin({a}, "id", {b}, "key", "B", JoinKind.FullOuter)')
        self.assertEqual(len(r['rows']), 4)
        r = ev(f'Table.Join({a}, "id", {b}, "key")')
        self.assertEqual(r['rows'], [[1, 'a', 1, 10], [1, 'a', 1, 11], [3, 'c', 3, 30]])
        r = ev(f'Table.Join({a}, "id", {b}, "key", JoinKind.RightOuter)')
        self.assertEqual(r['rows'][-1], [None, None, 9, 90])
        r = ev('Table.ExpandRecordColumn(#table({"r"}, {{[a=1,b=2]},{null}}), "r", {"a"})')
        self.assertEqual(r['rows'], [[1], [None]])
        r = ev('Table.ExpandListColumn(#table({"k","l"}, {{1,{1,2}},{2,{}}}), "l")')
        self.assertEqual(r['rows'], [[1, 1], [1, 2], [2, None]])

    def test_combine(self):
        r = ev('Table.Combine({#table({"A","B"},{{1,2}}), #table({"B","C"},{{3,4}})})')
        self.assertEqual(r['columns'], ['A', 'B', 'C'])
        self.assertEqual(r['rows'], [[1, 2, None], [None, 3, 4]])
        self.assertEqual(ev('#table({"A"},{{1}}) & #table({"A"},{{2}})')['rows'], [[1], [2]])

    def test_transpose_and_build(self):
        r = ev('Table.Transpose(#table({"A","B"},{{1,2},{3,4}}))')
        self.assertEqual(r['rows'], [[1, 3], [2, 4]])
        r = ev('Table.FromRecords({[a=1,b=2],[a=3,b=4]})')
        self.assertEqual(r, {'columns': ['a', 'b'], 'rows': [[1, 2], [3, 4]]})
        self.assertEqual(ev('Table.FromRows({{1,2}}, {"x","y"})')['columns'], ['x', 'y'])
        self.assertEqual(ev('Table.FromList({"a,b","c,d"}, Splitter.SplitTextByDelimiter(","))')['rows'],
                         [['a', 'b'], ['c', 'd']])
        self.assertEqual(ev('Table.ToRecords(#table({"a"},{{1}}))'), [{'a': 1}])
        self.assertEqual(ev('Record.ToTable([a=1])')['rows'], [['a', 1]])
        self.assertEqual(ev('Table.TransformColumnNames(#table({" a "},{{1}}), Text.Trim)')['columns'], ['a'])


class LibraryTest(unittest.TestCase):
    def test_text(self):
        self.assertEqual(ev('Text.Start("abcdef", 3)'), 'abc')
        self.assertEqual(ev('Text.End("abcdef", 2)'), 'ef')
        self.assertEqual(ev('Text.Middle("abcdef", 2, 2)'), 'cd')
        self.assertEqual(ev('Text.PadStart("7", 3, "0")'), '007')
        self.assertEqual(ev('Text.Contains("Hello", "ell")'), True)
        self.assertEqual(ev('Text.Contains("Hello", "ELL", Comparer.OrdinalIgnoreCase)'), True)
        self.assertEqual(ev('Text.BeforeDelimiter("a-b-c", "-")'), 'a')
        self.assertEqual(ev('Text.AfterDelimiter("a-b-c", "-", {0, RelativePosition.FromEnd})'), 'c')
        self.assertEqual(ev('Text.BetweenDelimiters("[x](y)", "(", ")")'), 'y')
        self.assertEqual(ev('Text.Split("a,b", ",")'), ['a', 'b'])
        self.assertEqual(ev('Text.Combine({"a", null, "b"}, "-")'), 'a-b', 'null は飛ばす')
        self.assertEqual(ev('Text.Remove("a1b2", {"0".."9"})'), 'ab')
        self.assertEqual(ev('Text.Select("a1b2", {"0".."9"})'), '12')
        self.assertEqual(ev('Text.Proper("hello WORLD")'), 'Hello World')
        self.assertEqual(ev('Text.Trim("  a  ")'), 'a')
        self.assertEqual(ev('Text.Clean("a#(lf)b")'), 'ab')
        self.assertEqual(ev('Text.Format("#{0}-#{1}", {"a", 1})'), 'a-1')
        self.assertEqual(ev('Text.Format("#[x]!", [x="y"])'), 'y!')
        self.assertEqual(ev('Text.PositionOf("abcabc", "c", Occurrence.Last)'), 5)
        self.assertEqual(ev('Text.Replace("aXbX", "X", "-")'), 'a-b-')
        self.assertIsNone(ev('Text.Upper(null)'))
        self.assertEqual(ev('Text.From(1234.5)'), '1234.5')
        self.assertEqual(ev('Text.From(#date(2024,1,5))'), '2024/01/05', 'ja-JP の既定の書き方')
        self.assertEqual(ev('Text.From(true)'), 'true')

    def test_number(self):
        self.assertEqual(ev('Number.Round(2.5)'), 2, '既定は偶数への丸め')
        self.assertEqual(ev('Number.Round(2.5, 0, RoundingMode.AwayFromZero)'), 3)
        self.assertEqual(ev('Number.Round(1.2345, 2)'), 1.23)
        self.assertEqual(ev('Number.RoundUp(1.21, 1)'), 1.3)
        self.assertEqual(ev('Number.RoundDown(-1.21)'), -2)
        self.assertEqual(ev('Number.Mod(-5, 3)'), -2)
        self.assertEqual(ev('Number.IntegerDivide(-7, 2)'), -3)
        self.assertEqual(ev('Number.FromText("1,000")'), 1000)
        self.assertEqual(ev('Number.From("12%")'), 0.12)
        self.assertEqual(ev('Number.ToText(1234.567, "N2")'), '1,234.57')
        self.assertEqual(ev('Number.ToText(0.5, "P0")'), '50%')
        self.assertEqual(ev('Number.ToText(5, "D3")'), '005')
        self.assertEqual(ev('Number.ToText(1234.5, "#,##0.00")'), '1,234.50')
        self.assertEqual(ev('Number.ToText(1234.5, "N1", "de-DE")'), '1.234,5')
        self.assertEqual(ev('Number.From(#date(2024,1,1))'), 45292, 'Excel と同じ日付の通し番号')
        self.assertEqual(ev('Int64.From("12.5")'), 12)
        self.assertEqual(ev('Number.Abs(-3)'), 3)
        self.assertEqual(ev('Number.Power(2, 10)'), 1024)

    def test_dates(self):
        self.assertEqual(ev('Date.AddMonths(#date(2024,1,31), 1)'), date(2024, 2, 29), '月末に寄せる')
        self.assertEqual(ev('Date.AddDays(#datetime(2024,1,1,9,0,0), 1)'), datetime(2024, 1, 2, 9, 0))
        self.assertEqual(ev('Date.EndOfMonth(#date(2024,2,10))'), date(2024, 2, 29))
        self.assertEqual(ev('Date.StartOfWeek(#date(2024,1,10))'), date(2024, 1, 7), '既定は日曜始まり')
        self.assertEqual(ev('Date.StartOfWeek(#date(2024,1,10), Day.Monday)'), date(2024, 1, 8))
        self.assertEqual(ev('Date.DayOfWeek(#date(2024,1,7))'), 0)
        self.assertEqual(ev('Date.DayOfWeekName(#date(2024,1,7))'), '日曜日')
        self.assertEqual(ev('Date.MonthName(#date(2024,3,1), "en-US")'), 'March')
        self.assertEqual(ev('Date.WeekOfYear(#date(2024,1,7))'), 2)
        self.assertEqual(ev('Date.QuarterOfYear(#date(2024,5,1))'), 2)
        self.assertEqual(ev('Date.ToText(#date(2024,1,5), "yyyy年M月d日（ddd）")'), '2024年1月5日（金）')
        self.assertEqual(ev('Date.ToText(#date(2024,1,5), [Format="MMM d, yyyy", Culture="en-US"])'), 'Jan 5, 2024')
        self.assertEqual(ev('DateTime.ToText(#datetime(2024,1,5,13,4,5), "yyyy/MM/dd HH:mm:ss")'), '2024/01/05 13:04:05')
        self.assertEqual(ev('Date.FromText("20240105", [Format="yyyyMMdd"])'), date(2024, 1, 5))
        self.assertEqual(ev('Date.From("2024年1月5日")'), date(2024, 1, 5))
        self.assertEqual(ev('DateTime.From(45292.5)'), datetime(2024, 1, 1, 12, 0))
        self.assertEqual(ev('#date(2024,1,3) - #date(2024,1,1)'), timedelta(days=2))
        self.assertEqual(ev('Duration.Days(#date(2024,1,3) - #date(2024,1,1))'), 2)
        self.assertEqual(ev('#datetime(2024,1,1,0,0,0) + #duration(1,2,0,0)'), datetime(2024, 1, 2, 2, 0))
        self.assertEqual(ev('DateTime.Date(#datetime(2024,1,1,5,0,0))'), date(2024, 1, 1))
        now = datetime(2026, 10, 14, 9, 30)
        self.assertEqual(ev('DateTime.LocalNow()', now=now), now, '今の時刻は外から渡せる（試験のため）')

    def test_lists(self):
        self.assertEqual(ev('List.Sum({1, null, 2.5})'), 3.5)
        self.assertIsNone(ev('List.Sum({})'))
        self.assertEqual(ev('List.Average({1, 2, 3, null})'), 2)
        self.assertEqual(ev('List.Max({3, null, 7})'), 7)
        self.assertEqual(ev('List.Min({"b", "a"})'), 'a')
        self.assertEqual(ev('List.Count({1, null})'), 2)
        self.assertEqual(ev('List.NonNullCount({1, null})'), 1)
        self.assertEqual(ev('List.Distinct({1, 1.0, "1", true})'), [1, '1', True], '1 と true は別')
        self.assertEqual(ev('List.Distinct({"a","A"}, Comparer.OrdinalIgnoreCase)'), ['a'])
        self.assertEqual(ev('List.Contains({"a","b"}, "b")'), True)
        self.assertEqual(ev('List.Sort({3,1,2}, Order.Descending)'), [3, 2, 1])
        self.assertEqual(ev('List.Accumulate({1,2,3}, 0, (s, x) => s + x)'), 6)
        self.assertEqual(ev('List.Generate(() => 1, each _ < 4, each _ + 1)'), [1, 2, 3])
        self.assertEqual(ev('List.Zip({{1,2},{"a"}})'), [[1, 'a'], [2, None]])
        self.assertEqual(ev('List.Numbers(1, 3, 2)'), [1, 3, 5])
        self.assertEqual(ev('List.Dates(#date(2024,1,30), 3, #duration(1,0,0,0))'),
                         [date(2024, 1, 30), date(2024, 1, 31), date(2024, 2, 1)])
        self.assertEqual(ev('List.Difference({1,2,2,3}, {2})'), [1, 2, 3])
        self.assertEqual(ev('List.Union({{1,2},{2,3}})'), [1, 2, 3])
        self.assertEqual(ev('List.Intersect({{1,2,3},{2,3,4}})'), [2, 3])
        self.assertEqual(ev('List.PositionOf({"a","b"}, "b")'), 1)
        self.assertEqual(ev('List.RemoveMatchingItems({1, "", null, 2}, {"", null})'), [1, 2])
        self.assertEqual(ev('List.Median({1, 3, 2, 4})'), 2.5)

    def test_records(self):
        self.assertEqual(ev('Record.Field([a=1], "a")'), 1)
        self.assertEqual(ev('Record.FieldOrDefault([a=1], "b", 0)'), 0)
        self.assertEqual(ev('Record.FieldNames([a=1, b=2])'), ['a', 'b'])
        self.assertEqual(ev('Record.FieldValues([a=1, b=2])'), [1, 2])
        self.assertEqual(ev('Record.AddField([a=1], "b", 2)'), {'a': 1, 'b': 2})
        self.assertEqual(ev('Record.RemoveFields([a=1, b=2], "a")'), {'b': 2})
        self.assertEqual(ev('Record.RenameFields([a=1], {"a", "x"})'), {'x': 1})
        self.assertEqual(ev('Record.SelectFields([a=1, b=2], {"b"})'), {'b': 2})
        self.assertEqual(ev('Record.TransformFields([a=1], {"a", each _ + 1})'), {'a': 2})
        self.assertEqual(ev('Record.Combine({[a=1], [b=2], [a=3]})'), {'a': 3, 'b': 2})

    def test_blank_row_removal_as_generated_by_excel(self):
        """「空白行の削除」で Excel が書く式。"""
        t = '#table({"A","B"}, {{"x", 1},{"", null},{null, null}})'
        r = ev(f'Table.SelectRows({t}, each not List.IsEmpty(List.RemoveMatchingItems(Record.FieldValues(_), {{"", null}})))')
        self.assertEqual(r['rows'], [['x', 1]])

    def test_value_and_expression(self):
        self.assertEqual(ev('Value.Type(1) = type number'), True)
        self.assertEqual(ev('Expression.Evaluate("1 + x", [x = 2])'), 3)
        self.assertEqual(ev('Value.ReplaceType(#table({"a"},{{1}}), type table [b = number])')['columns'], ['b'])
        self.assertEqual(ev('Function.Invoke((a, b) => a - b, {5, 2})'), 3)


class ConversionTest(unittest.TestCase):
    def test_parse_number(self):
        for text, want in (('1,234', 1234), ('¥1,234', 1234), ('(12.5)', -12.5), ('12-', -12), ('１２３', 123),
                           ('50%', 0.5), ('1e3', 1000), ('', None), ('abc', None)):
            self.assertEqual(C.parse_number(text), want, text)

    def test_parse_dates(self):
        for text, want in (('2024/1/5', date(2024, 1, 5)), ('2024-01-05', date(2024, 1, 5)),
                           ('2024年1月5日', date(2024, 1, 5)), ('20240105', date(2024, 1, 5)),
                           ('2024/1/5 13:04', date(2024, 1, 5))):
            self.assertEqual(C.parse_datetime_text(text, None, 'date'), want, text)
        self.assertEqual(C.parse_datetime_text('2024/1/5 午後1:04'), datetime(2024, 1, 5, 13, 4))
        self.assertEqual(C.parse_datetime_text('1/5/2024 1:04 PM', 'en-US'), datetime(2024, 1, 5, 13, 4))
        self.assertEqual(C.parse_datetime_text('5/1/2024', 'en-GB', 'date'), date(2024, 1, 5))
        self.assertIsNone(C.parse_datetime_text('2024/13/40', None, 'date'), '無い日付は読めない')

    def test_cell_text(self):
        self.assertEqual(C.cell_text(0.1 + 0.2), '0.3', 'ファイルへ書くときは15桁（Excel と同じ見え方）')
        self.assertEqual(C.cell_text(date(2024, 1, 5)), '2024-01-05')
        self.assertEqual(C.cell_text(datetime(2024, 1, 5, 9, 3)), '2024-01-05 09:03:00')
        self.assertEqual(C.cell_text(None), '')
        self.assertEqual(C.cell_text(True), 'TRUE')


class ExcelSourceTest(unittest.TestCase):
    """Excel の「フォルダーから結合」が書き出す式を、そのまま貼り付けて動かす。"""

    @classmethod
    def setUpClass(cls):
        import openpyxl
        cls.tmp = tempfile.TemporaryDirectory()
        root = Path(cls.tmp.name)
        cls.folder = root / '日報'
        (cls.folder / 'sub').mkdir(parents=True)
        books = {'日報_0101.xlsx': [(datetime(2024, 1, 1), 'A-01', 10, 120.5), (datetime(2024, 1, 1), 'B-02', None, 80)],
                 'sub/日報_0102.xlsx': [(datetime(2024, 1, 2), 'C-03', 'x', 99)]}
        for name, rows in books.items():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Sheet1'
            ws.append(('日付', '品番', '数量', '単価'))
            for r in rows:
                ws.append(r)
            memo = wb.create_sheet('メモ')
            memo['B3'] = '離れたセル'
            wb.save(cls.folder / name)
        # テーブル（ListObject）と名前付き範囲を持つブック
        from openpyxl.workbook.defined_name import DefinedName
        from openpyxl.worksheet.table import Table
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '集計'
        ws['A1'] = 'タイトル'
        for r in (('品番', '数'), ('A', 1), ('B', 2)):
            ws.append(r)
        ws.add_table(Table(displayName='売上表', ref='A2:B4'))
        try:
            wb.defined_names['範囲1'] = DefinedName('範囲1', attr_text="'集計'!$A$3:$B$4")
        except TypeError:
            wb.defined_names.append(DefinedName('範囲1', attr_text="'集計'!$A$3:$B$4"))
        cls.book = root / 'book.xlsx'
        wb.save(cls.book)
        (root / '壊れた.xlsx').write_bytes(b'not a zip')
        cls.broken = root / '壊れた.xlsx'
        cls.csv = root / 'a.csv'
        cls.csv.write_bytes('品番,数量\r\n"A,1",10\r\nB,20\r\n'.encode('cp932'))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def combine_text(self):
        p = str(self.folder).replace('"', '""')
        return f'''// サンプル ファイル
let
    ソース = Folder.Files("{p}"),
    ナビゲーション1 = ソース{{0}}[Content]
in
    ナビゲーション1

// パラメーター1
#"サンプル ファイル" meta [IsParameterQuery=true, BinaryIdentifier=#"サンプル ファイル", Type="Binary", IsParameterQueryRequired=true]

// サンプル ファイルの変換
let
    ソース = Excel.Workbook(パラメーター1, null, true),
    Sheet1_Sheet = ソース{{[Item="Sheet1",Kind="Sheet"]}}[Data],
    昇格されたヘッダー数 = Table.PromoteHeaders(Sheet1_Sheet, [PromoteAllScalars=true])
in
    昇格されたヘッダー数

// ファイルの変換
let
    ソース = (パラメーター1) => let
        ソース = Excel.Workbook(パラメーター1, null, true),
        Sheet1_Sheet = ソース{{[Item="Sheet1",Kind="Sheet"]}}[Data],
        昇格されたヘッダー数 = Table.PromoteHeaders(Sheet1_Sheet, [PromoteAllScalars=true])
    in
        昇格されたヘッダー数
in
    ソース

// 日報
let
    ソース = Folder.Files("{p}"),
    #"フィルター選択された非表示のファイル1" = Table.SelectRows(ソース, each [Attributes]?[Hidden]? <> true),
    #"カスタム関数の呼び出し1" = Table.AddColumn(#"フィルター選択された非表示のファイル1", "ファイルの変換", each #"ファイルの変換"([Content])),
    #"名前変更された列1" = Table.RenameColumns(#"カスタム関数の呼び出し1", {{"Name", "Source.Name"}}),
    #"削除された他の列1" = Table.SelectColumns(#"名前変更された列1", {{"Source.Name", "ファイルの変換"}}),
    #"展開されたテーブル列1" = Table.ExpandTableColumn(#"削除された他の列1", "ファイルの変換", Table.ColumnNames(#"サンプル ファイルの変換")),
    #"変更された型" = Table.TransformColumnTypes(#"展開されたテーブル列1",{{{{"Source.Name", type text}}, {{"日付", type date}}, {{"品番", type text}}, {{"数量", Int64.Type}}, {{"単価", type number}}}})
in
    #"変更された型"
'''

    def test_combine_files_as_written_by_excel(self):
        doc = M.Document.from_text(self.combine_text())
        self.assertEqual(doc.main_name(), '日報')
        self.assertEqual(doc.dependencies('日報'), ['サンプル ファイルの変換', 'ファイルの変換'])
        t = doc.evaluate()
        self.assertEqual(t.columns, ['Source.Name', '日付', '品番', '数量', '単価'])
        self.assertEqual([x.name() for x in t.types], ['text', 'date', 'text', 'Int64.Type', 'number'])
        rows = plain(t)['rows']
        self.assertEqual(rows[0], ['日報_0101.xlsx', date(2024, 1, 1), 'A-01', 10, 120.5])
        self.assertEqual(rows[1][3], None)
        self.assertEqual(rows[2][0], '日報_0102.xlsx', '下のフォルダーのファイルも読む')
        self.assertEqual(rows[2][3], 'ERR', '数でない値は、その1つのセルだけ誤りになる')

    def test_preview_row_cap(self):
        h = M.Host(row_cap=2)
        doc = M.Document.from_text(self.combine_text())
        t = doc.evaluate(host_=h)
        self.assertEqual(len(t.rows), 2, '見出し＋1行だけ読む（2ファイル×1行）')
        self.assertTrue(any('先頭 2 行' in n for n in h.notes), '省いたことを書き留める')

    def test_folder_files_shape(self):
        p = str(self.folder).replace('"', '""')
        t = M.Document.from_text(f'Folder.Files("{p}")').evaluate()
        self.assertEqual(t.columns, ['Content', 'Name', 'Extension', 'Date accessed', 'Date modified', 'Date created',
                                     'Attributes', 'Folder Path'])
        self.assertEqual([plain(r[1]) for r in t.rows], ['日報_0101.xlsx', '日報_0102.xlsx'])
        self.assertTrue(str(M.force(t.rows[1][7])).endswith(('sub\\', 'sub/')), 'Folder Path は区切りで終わる')
        t2 = M.Document.from_text(f'Folder.Contents("{p}")').evaluate()
        self.assertEqual([plain(r[1]) for r in t2.rows], ['sub', '日報_0101.xlsx'], '直下だけ（フォルダーも並ぶ）')

    def test_workbook_items(self):
        p = str(self.book).replace('"', '""')
        doc = M.Document.from_text(f'Excel.Workbook(File.Contents("{p}"), null, true)')
        t = doc.evaluate()
        kinds = [(plain(r[0]), plain(r[3])) for r in t.rows]
        self.assertIn(('集計', 'Sheet'), kinds)
        self.assertIn(('売上表', 'Table'), kinds)
        self.assertIn(('範囲1', 'DefinedName'), kinds)
        tbl = M.Document.from_text(f'Excel.Workbook(File.Contents("{p}")){{[Item="売上表",Kind="Table"]}}[Data]').evaluate()
        self.assertEqual(plain(tbl), {'columns': ['品番', '数'], 'rows': [['A', 1], ['B', 2]]})
        sheet = M.Document.from_text(f'Excel.Workbook(File.Contents("{p}"), true){{[Item="集計"]}}[Data]').evaluate()
        self.assertEqual(sheet.columns, ['タイトル', 'Column2'], '見出しを使う（useHeaders=true）')

    def test_used_range_starts_at_first_used_cell(self):
        p = str(self.folder / '日報_0101.xlsx').replace('"', '""')
        t = M.Document.from_text(f'Excel.Workbook(File.Contents("{p}")){{[Item="メモ",Kind="Sheet"]}}[Data]').evaluate()
        self.assertEqual(plain(t), {'columns': ['Column1'], 'rows': [['離れたセル']]})

    def test_missing_sheet_and_files(self):
        p = str(self.book).replace('"', '""')
        with self.assertRaises(M.MError) as cm:
            M.Document.from_text(f'Excel.Workbook(File.Contents("{p}")){{[Item="無い",Kind="Sheet"]}}[Data]').evaluate()
        self.assertIn('シート名', str(cm.exception))
        with self.assertRaises(M.MError) as cm:
            M.Document.from_text('Folder.Files("/無い/フォルダー")').evaluate()
        self.assertIn('フォルダーが見つかりません', str(cm.exception))
        b = str(self.broken).replace('"', '""')
        with self.assertRaises(M.MError) as cm:
            M.Document.from_text(f'Excel.Workbook(File.Contents("{b}"))').evaluate()
        self.assertIn('Excel のファイルとして読めません', str(cm.exception))
        with self.assertRaises(M.MError) as cm:
            M.Document.from_text('Excel.CurrentWorkbook(){[Name="t"]}[Content]').evaluate()
        self.assertIn('Excel.Workbook(File.Contents', str(cm.exception))
        with self.assertRaises(M.MError) as cm:
            M.Document.from_text('Web.Contents("http://example.com")').evaluate()
        self.assertIn('使えません', str(cm.exception))

    def test_csv_document(self):
        p = str(self.csv).replace('"', '""')
        t = M.Document.from_text(f'let s = Csv.Document(File.Contents("{p}"), [Delimiter=",", Columns=2, Encoding=932, '
                                 f'QuoteStyle=QuoteStyle.Csv]), h = Table.PromoteHeaders(s) in h').evaluate()
        self.assertEqual(plain(t), {'columns': ['品番', '数量'], 'rows': [['A,1', '10'], ['B', '20']]})

    def test_resolve_hook(self):
        """相対パスの読み替えは外（アプリ）から渡す。"""
        h = M.Host(resolve=lambda p: str(self.folder) if p == '.\\日報' else p)
        t = M.Document.from_text('Folder.Files(".\\日報")').evaluate(host_=h)
        self.assertEqual(len(t.rows), 2)


if __name__ == '__main__':
    unittest.main()


# ---- Excel が書き出したファイルからの取り込み（.odc・.xlsx）と、実物の .odc と同じ作りの式 ----------------
def raw_xlsx(path, sheets, strings=(), date1904=False, names=(), num_fmts=(), xfs=(0,)):
    """XML を直接組んだ .xlsx（Excel が書くが openpyxl では書けない形: ふりがな付きの共有文字・その場の文字・
    行や列の番号の省略・範囲の書き込みが無いシート・1904 年基準 など）。sheets は [(名前, sheetData の中身, dimension)]。"""
    import zipfile
    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rns = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    pr = 'http://schemas.openxmlformats.org/package/2006/relationships'
    od = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    ct = 'application/vnd.openxmlformats-officedocument.spreadsheetml'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml',
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   f'<Override PartName="/xl/workbook.xml" ContentType="{ct}.sheet.main+xml"/>'
                   f'<Override PartName="/xl/styles.xml" ContentType="{ct}.styles+xml"/>'
                   f'<Override PartName="/xl/sharedStrings.xml" ContentType="{ct}.sharedStrings+xml"/>'
                   + ''.join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="{ct}.worksheet+xml"/>'
                             for i in range(1, len(sheets) + 1)) + '</Types>')
        z.writestr('_rels/.rels', f'<Relationships xmlns="{pr}"><Relationship Id="rId1" Type="{od}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        pr1904 = ' date1904="1"' if date1904 else ''
        z.writestr('xl/workbook.xml', f'<workbook {ns} {rns}><workbookPr{pr1904}/><sheets>'
                   + ''.join(f'<sheet name="{n}" sheetId="{i}" r:id="rId{i}"/>' for i, (n, _, _) in enumerate(sheets, 1))
                   + '</sheets>' + (('<definedNames>' + ''.join(f'<definedName name="{n}">{v}</definedName>' for n, v in names)
                                     + '</definedNames>') if names else '') + '</workbook>')
        n = len(sheets)
        z.writestr('xl/_rels/workbook.xml.rels', f'<Relationships xmlns="{pr}">'
                   + ''.join(f'<Relationship Id="rId{i}" Type="{od}/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1, n + 1))
                   + f'<Relationship Id="rId{n + 1}" Type="{od}/styles" Target="styles.xml"/>'
                   + f'<Relationship Id="rId{n + 2}" Type="{od}/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
        z.writestr('xl/styles.xml', f'<styleSheet {ns}>'
                   + (f'<numFmts count="{len(num_fmts)}">' + ''.join(f'<numFmt numFmtId="{i}" formatCode="{c}"/>' for i, c in num_fmts)
                      + '</numFmts>' if num_fmts else '')
                   + '<fonts count="1"><font><sz val="11"/></font></fonts><fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
                   + '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
                   + '<cellStyleXfs count="1"><xf numFmtId="0"/></cellStyleXfs>'
                   + f'<cellXfs count="{len(xfs)}">' + ''.join(f'<xf numFmtId="{f}" xfId="0" applyNumberFormat="1"/>' for f in xfs)
                   + '</cellXfs></styleSheet>')
        z.writestr('xl/sharedStrings.xml', f'<sst {ns} count="{len(strings)}" uniqueCount="{len(strings)}">'
                   + ''.join(f'<si>{x}</si>' for x in strings) + '</sst>')
        for i, (_, data, dim) in enumerate(sheets, 1):
            z.writestr(f'xl/worksheets/sheet{i}.xml', f'<worksheet {ns}>' + (f'<dimension ref="{dim}"/>' if dim else '')
                       + f'<sheetData>{data}</sheetData></worksheet>')


class XlsxReaderTest(unittest.TestCase):
    """Excel を読む速い読み口（navi_mxlsx）が、今までの読み方（openpyxl）と同じ値を返すかの突き合わせ。"""

    @classmethod
    def setUpClass(cls):
        import openpyxl
        cls.tmp = tempfile.TemporaryDirectory()
        root = Path(cls.tmp.name)
        cls.books = []
        # 1) openpyxl で書いたブック: 日付・時刻・小数・整数・真偽・空の文字・数式（値なし）・抜けた行と列・名前付き範囲
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '実績'
        ws.append(('日付', '時刻', '数', '小数', '真偽', '文字'))
        ws.append((datetime(2024, 1, 5, 8, 30), time(9, 15), 10, 1.5, True, '検査 済'))
        ws.append((date(2024, 2, 29), None, -3, 1e-7, False, ''))
        ws['H6'] = '離れたセル'
        ws['B8'] = '=SUM(C2:C3)'
        ws['C2'].number_format = '#,##0'
        ws['D3'].number_format = 'yyyy"年"m"月"d"日"'
        s2 = wb.create_sheet('空')
        s2.sheet_state = 'hidden'
        s3 = wb.create_sheet('表 2')
        for r in range(1, 30):
            s3.append([f'行{r}', r, r * 0.25] if r % 7 else [])
        from openpyxl.workbook.defined_name import DefinedName
        dn = DefinedName('範囲A', attr_text="'表 2'!$A$2:$C$10")
        if hasattr(wb.defined_names, 'add'):
            wb.defined_names.add(dn)
        else:
            wb.defined_names['範囲A'] = dn
        p1 = root / 'openpyxl.xlsx'
        wb.save(p1)
        cls.books.append(p1)
        # 2) Excel の書き方: ふりがな付き・飾り付きの共有文字、その場の文字、誤りのセル、行・列の番号の省略、
        #    範囲の書き込みの無いシート、時間の長さ（[h]:mm）、時刻だけ、文字の日付（t="d"）
        strings = ['<t>品番</t>', '<r><t>東京</t></r><r><rPr/><t>工場</t></r><rPh sb="0" eb="2"><t>トウキョウ</t></rPh>',
                   '<t xml:space="preserve"> 前後に空白 </t>', '<t></t>']
        data = ('<row r="2"><c r="B2" t="s"><v>0</v></c><c r="C2" t="s"><v>1</v></c><c r="E2" t="inlineStr"><is><t>その場</t></is></c></row>'
                '<row r="3"><c r="B3" s="1"><v>45306</v></c><c r="C3" s="2"><v>1.75</v></c><c r="D3" s="3"><v>0.5</v></c>'
                '<c r="E3" t="e"><v>#N/A</v></c></row>'
                '<row><c t="s"><v>2</v></c><c><v>3.5E+2</v></c><c t="b"><v>1</v></c><c t="str"><v>式の文字</v></c><c t="s"><v>3</v></c></row>'
                '<row r="7"><c r="F7" t="d"><v>2024-03-01T00:00:00</v></c><c r="G7" s="4"><v>99999999</v></c></row>')
        p2 = root / 'excel.xlsx'
        raw_xlsx(p2, [('生', data, None), ('範囲あり', data, 'B2:G7'), ('1セル', data, 'A1')], strings=strings,
                 num_fmts=[(164, 'yyyy/mm/dd'), (165, '[h]:mm')], xfs=(0, 14, 165, 20, 164),
                 names=[('名前付き', "'範囲あり'!$B$2:$E$4")])
        cls.books.append(p2)
        p3 = root / 'mac1904.xlsx'
        raw_xlsx(p3, [('S', '<row r="1"><c r="A1" s="1"><v>0</v></c><c r="B1" s="1"><v>366</v></c><c r="C1"><v>12</v></c></row>', 'A1:C1')],
                 date1904=True, xfs=(0, 14))
        cls.books.append(p3)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def read_all(self, path, fast, cap=None):
        """ブックの全項目（シート・テーブル・名前付き範囲）を、見出しなし・ありの両方で読んだ値。"""
        import navi_msource as S
        import navi_mxlsx as X
        real = X.Book
        if not fast:
            def refuse(*a, **k):
                raise X.Unsupported('試験: openpyxl で読む')
            X.Book = refuse
        try:
            out = {}
            for head in ('false', 'true'):
                src = (f'let wb = Excel.Workbook(File.Contents("{str(path).replace(chr(34), chr(34) * 2)}"), {head}) in '
                       'Table.TransformColumns(wb, {"Data", each _})')
                v = M.Document.from_text(src).evaluate(host_=M.Host(row_cap=cap))
                for row in v.rows:
                    t = M.force(row[1])
                    out[(head, row[0], row[3], row[4])] = (t.columns, [[('ERR', c.error.message) if isinstance(c, M.ErrorCell) else c
                                                                        for c in r] for r in t.rows])
            S._cache().clear() if hasattr(S._cache(), 'clear') else None
            return out
        finally:
            X.Book = real

    def test_same_values_as_openpyxl(self):
        for p in self.books:
            for cap in (None, 3):
                with self.subTest(book=p.name, cap=cap):
                    import navi_mxlsx as X
                    self.assertTrue(X.Book(str(p)).sheets, '速い読み口で読めている（openpyxl に退いていない）')
                    slow, fast = self.read_all(p, False, cap), self.read_all(p, True, cap)
                    self.assertEqual(sorted(fast), sorted(slow), '同じ項目（シート・名前付き範囲・隠し）')
                    for k in slow:
                        self.assertEqual(fast[k], slow[k], k)

    def test_east_asian_builtin_date_formats(self):
        """日本語版の組み込みの書式（31: yyyy"年"m"月"d"日"、32: h"時"mm"分"）は日付・時刻（openpyxl は数のまま読む）。"""
        import navi_mxlsx as X
        p = Path(self.tmp.name) / 'ja.xlsx'
        raw_xlsx(p, [('S', '<row r="1"><c r="A1" s="1"><v>45383</v></c><c r="B1" s="2"><v>0.25</v></c><c r="C1" s="3"><v>45383</v></c></row>', 'A1:C1')],
                 xfs=(0, 31, 32, 58))
        rows = X.Book(str(p)).rows('S')
        self.assertEqual(rows, [(1, [(1, datetime(2024, 4, 1)), (2, time(6, 0)), (3, datetime(2024, 4, 1))])])

    def test_falls_back_for_unknown_shape(self):
        import navi_mxlsx as X
        with self.assertRaises(X.Unsupported):
            X.Book(io_bytes(b'not a zip'))
        self.assertEqual(ev('Table.ColumnNames(Excel.Workbook(File.Contents("' + str(self.books[0]).replace('"', '""') + '")))'),
                         ['Name', 'Data', 'Item', 'Kind', 'Hidden'])


def io_bytes(b):
    import io
    return io.BytesIO(b)


# 実物（利用者から受け取った「接続ファイルのエクスポート」の .odc）と同じ書き方: 引数に型の付いた関数クエリ（as binary）、
# 空白入りの列名 [#"   本"]、Table.ColumnNames(関数(#"サンプル ファイル"))、ファイル名を置き換えて分け「2024年10月」を日付にする、
# 同じ月は先のファイルだけ残す Table.Distinct、#N/A のセルの行を落とす RemoveRowsWithErrors。場所と名前は試験用に変えてある。
# 利用者のブック（.xlsx）にあった5本の組（出すクエリ・パラメーター1（meta 付きの引数）・サンプル ファイル・
# サンプル ファイルの変換・ファイルの変換（関数））と同じ作り。
PLAN_FORMULAS = [
    ('計画の集約', '''let
    ソース = Folder.Files("{folder}"),
    フィルターされた行 = Table.SelectRows(ソース, each Text.Contains([Name], "生産計画") and Text.Contains([Name], "●")),
    フィルターされた行1 = Table.SelectRows(フィルターされた行, each ([Extension] = ".xlsx")),
    置き換えられた値 = Table.ReplaceValue(フィルターされた行1,"_生産計画SP","",Replacer.ReplaceText,{{"Name"}}),
    置き換えられた値1 = Table.ReplaceValue(置き換えられた値,".xlsx","",Replacer.ReplaceText,{{"Name"}}),
    置き換えられた値2 = Table.ReplaceValue(置き換えられた値1,"_稼働日","",Replacer.ReplaceText,{{"Name"}}),
    置き換えられた値3 = Table.ReplaceValue(置き換えられた値2,"●","",Replacer.ReplaceText,{{"Name"}}),
    置き換えられた値4 = Table.ReplaceValue(置き換えられた値3,"【場内用】","",Replacer.ReplaceText,{{"Name"}}),
    置き換えられた値5 = Table.ReplaceValue(置き換えられた値4,"_R1","",Replacer.ReplaceText,{{"Name"}}),
    置き換えられた値6 = Table.ReplaceValue(置き換えられた値5," ","",Replacer.ReplaceText,{{"Name"}}),
    削除された他の列 = Table.SelectColumns(置き換えられた値6,{{"Name", "Content"}}),
    区切り記号による列の分割 = Table.SplitColumn(削除された他の列, "Name", Splitter.SplitTextByDelimiter("_", QuoteStyle.Csv), {{"Name.1", "Name.2"}}),
    削除された重複 = Table.Distinct(区切り記号による列の分割, {{"Name.2"}}),
    変更された型 = Table.TransformColumnTypes(削除された重複,{{{{"Name.1", Int64.Type}}, {{"Name.2", type date}}}}),
    #"フィルター選択された非表示の File1" = Table.SelectRows(変更された型, each [Attributes]?[Hidden]? <> true),
    カスタム関数の呼び出し1 = Table.AddColumn(#"フィルター選択された非表示の File1", "ファイルの変換", each ファイルの変換([Content])),
    #"名前が変更された列 " = Table.RenameColumns(カスタム関数の呼び出し1,{{{{"Name.1", "作成日"}}, {{"Name.2", "計画年月"}}}}),
    削除された他の列1 = Table.SelectColumns(#"名前が変更された列 ", {{"作成日","計画年月","ファイルの変換"}}),
    展開されたテーブル列1 = Table.ExpandTableColumn(削除された他の列1, "ファイルの変換", Table.ColumnNames(ファイルの変換(#"サンプル ファイル"))),
    削除された列 = Table.RemoveColumns(展開されたテーブル列1,{{"計画", "重量"}})
in
    削除された列'''),
    ('ファイルの変換', '''let
    ソース = (パラメーター1 as binary) => let
        ソース = Excel.Workbook(パラメーター1, null, true),
        日当量_Sheet = ソース{{[Item="日当量",Kind="Sheet"]}}[Data],
        削除された他の列 = Table.SelectColumns(日当量_Sheet,{{"Column1", "Column2", "Column6", "Column7"}}),
        削除された最初の行 = Table.Skip(削除された他の列,4),
        昇格されたヘッダー数 = Table.PromoteHeaders(削除された最初の行, [PromoteAllScalars=true]),
        変更された型 = Table.TransformColumnTypes(昇格されたヘッダー数,{{{{"設備名", type text}}, {{"計画", type number}}, {{"Column3", type any}}, {{"   本", type any}}}}),
        #"名前が変更された列 " = Table.RenameColumns(変更された型,{{{{"Column3", "重量"}}}}),
        フィルターされた行 = Table.SelectRows(#"名前が変更された列 ", each ([設備名] <> " CF計")),
        削除されたエラー = Table.RemoveRowsWithErrors(フィルターされた行, {{"重量"}}),
        フィルターされた行1 = Table.SelectRows(削除されたエラー, each ([設備名] <> null and [設備名] <> " AP2")),
        置き換えられた値 = Table.ReplaceValue(フィルターされた行1," ","",Replacer.ReplaceText,{{"設備名"}}),
        フィルターされた行2 = Table.SelectRows(置き換えられた値, each ([#"   本"] <> null))
    in
        フィルターされた行2
in
    ソース'''),
    ('パラメーター1', '#"サンプル ファイル" meta [IsParameterQuery=true, BinaryIdentifier=#"サンプル ファイル", Type="Binary", IsParameterQueryRequired=true]'),
    ('サンプル ファイルの変換', '''let
    ソース = Excel.Workbook(パラメーター1, null, true),
    日当量_Sheet = ソース{{[Item="日当量",Kind="Sheet"]}}[Data],
    削除された他の列 = Table.SelectColumns(日当量_Sheet,{{"Column1", "Column2", "Column6", "Column7"}}),
    削除された最初の行 = Table.Skip(削除された他の列,4),
    昇格されたヘッダー数 = Table.PromoteHeaders(削除された最初の行, [PromoteAllScalars=true]),
    変更された型 = Table.TransformColumnTypes(昇格されたヘッダー数,{{{{"設備名", type text}}, {{"計画", type number}}, {{"Column3", type any}}, {{"   本", type any}}}}),
    #"名前が変更された列 " = Table.RenameColumns(変更された型,{{{{"Column3", "重量"}}}}),
    削除されたエラー = Table.RemoveRowsWithErrors(#"名前が変更された列 ", {{"重量"}}),
    フィルターされた行 = Table.SelectRows(削除されたエラー, each ([設備名] <> null and [設備名] <> " CF計" and [設備名] <> " AP2")),
    削除されたエラー1 = Table.RemoveRowsWithErrors(フィルターされた行, {{"計画"}}),
    フィルターされた行2 = Table.SelectRows(削除されたエラー1, each ([#"   本"] <> null))
in
    フィルターされた行2'''),
    ('サンプル ファイル', '''let
    ソース = Folder.Files("{folder}"),
    フィルターされた行1 = Table.SelectRows(ソース, each ([Extension] = ".xlsx") and Text.Contains([Name], "●")),
    ナビゲーション1 = フィルターされた行1{{0}}[Content]
in
    ナビゲーション1'''),
]


def odc_text(pairs, main, old=False):
    """Excel の「接続ファイルのエクスポート」と同じ作りの .odc。old なら式を接続文字列に包む古い形。"""
    import base64
    import html as H
    import io
    import struct
    import zipfile
    if old:
        sec = 'section Section1;\r\n\r\n' + ''.join(f'shared #"{n}" = {f};\r\n\r\n' for n, f in pairs)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z:
            z.writestr('Formulas/Section1.m', sec)
        pkg = struct.pack('<II', 0, len(buf.getvalue())) + buf.getvalue() + struct.pack('<III', 0, 0, 0)
        conn = f'Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location={main};Extended Properties=&quot;{base64.b64encode(pkg).decode()}&quot;'
        mash = ''
    else:
        conn = f'Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location={main};Extended Properties=&quot;&quot;'
        items = ''.join(f'<Query Name="{H.escape(n)}"><Formula><![CDATA[{f}]]></Formula><IsParameterQuery xsi:nil="true" />'
                        f'<IsDirectQuery xsi:nil="true" /></Query>' for n, f in pairs)
        mashup = ('<Mashup xmlns:xsd="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
                  'xmlns="http://schemas.microsoft.com/DataMashup"><Client>EXCEL</Client><Version>2.157.151.0</Version>'
                  f'<Culture>ja-JP</Culture><SafeCombine>true</SafeCombine><Items>{items}</Items></Mashup>')
        mash = f'<odc:PowerQueryMashupData>{H.escape(mashup)}</odc:PowerQueryMashupData>'
    return ('<html xmlns:o="urn:schemas-microsoft-com:office:office"\r\nxmlns="http://www.w3.org/TR/REC-html40">\r\n<head>\r\n'
            '<meta http-equiv=Content-Type content="text/x-ms-odc; charset=utf-8">\r\n<meta name=ProgId content=ODC.Database>\r\n'
            f'<title>クエリ - {main}</title>\r\n<xml id=msodc><odc:OfficeDataConnection xmlns:odc="urn:schemas-microsoft-com:office:odc">'
            f'<odc:PowerQueryConnection odc:Type="OLEDB"><odc:ConnectionString>{conn}</odc:ConnectionString>'
            f'<odc:CommandType>SQL</odc:CommandType><odc:CommandText>SELECT * FROM [{main}]</odc:CommandText>'
            f'</odc:PowerQueryConnection>{mash}</odc:OfficeDataConnection></xml></head></html>').encode('utf-8')


class ExcelFileImportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import openpyxl
        cls.tmp = tempfile.TemporaryDirectory()
        cls.folder = Path(cls.tmp.name) / '計画'
        cls.folder.mkdir()

        def book(name, rows):
            wb = openpyxl.Workbook()
            wb.active.title = '表紙'
            ws = wb.create_sheet('日当量')
            for r in [('生産計画',), (None,), ('2024',), (None,)]:
                ws.append(r)
            ws.append(('設備名', '計画', None, None, None, None, '   本'))
            for r in rows:
                ws.append(r)
            wb.save(cls.folder / name)
        book('●20241001_生産計画SP_2024年10月_稼働日.xlsx',
             [(' LS3', 120.5, '', None, None, 3.2, 10), (' CF計', 300, None, None, None, 1, 1), (' AP2', 1, None, None, None, 1, 1),
              (' LS4', 80, None, None, None, 2.5, None), (' DL2', 50, None, None, None, '#N/A', 4)])
        book('●20241015_生産計画SP_2024年11月_稼働日_R1.xlsx', [(' LS3', 130, None, None, None, 3.0, 12)])
        book('●20241020_生産計画SP_2024年11月_稼働日_R2.xlsx', [(' LS3', 999, None, None, None, 3.0, 12)])
        book('【場内用】●20240901_生産計画SP_2024年 9月_稼働日.xlsx', [(' KEN', 10, None, None, None, 1.0, 2)])
        (cls.folder / 'メモ.txt').write_text('x')
        p = str(cls.folder).replace('"', '""')
        cls.pairs = [(n, f.format(folder=p)) for n, f in PLAN_FORMULAS]

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def check_plan(self, code):
        doc = M.Document.from_text(code)
        self.assertEqual(doc.main_name(), '計画の集約')
        self.assertEqual(doc.helpers(), ['ファイルの変換', 'パラメーター1', 'サンプル ファイルの変換', 'サンプル ファイル'])
        self.assertEqual((doc.kind('パラメーター1'), doc.kind('ファイルの変換')), ('parameter', 'function'))
        self.assertEqual(doc.missing_names(), [])
        s = doc.evaluate('サンプル ファイルの変換')
        self.assertEqual((s.columns, plain(s)['rows']), (['設備名', '計画', '重量', '   本'], [[' LS3', 120.5, 3.2, 10]]),
                         'サンプル ファイル（先頭のファイル）を パラメーター1 経由で読む')
        t = doc.evaluate('計画の集約')
        self.assertEqual(t.columns, ['作成日', '計画年月', '設備名', '   本'])
        self.assertEqual(plain(t)['rows'], [[20241001, date(2024, 10, 1), 'LS3', 10],
                                            [20241015, date(2024, 11, 1), 'LS3', 12],
                                            [20240901, date(2024, 9, 1), 'KEN', 2]],
                         '合計行・AP2・本が空の行・#N/A の行を落とし、同じ月は先のファイル（R1）だけを残す')

    def test_odc(self):
        import navi_mimport as I
        blob = odc_text(self.pairs, '計画の集約')
        self.assertEqual(I.kind_of(blob), 'odc')
        pairs, main = I.extract(blob)
        self.assertEqual([n for n, _ in pairs], [n for n, _ in PLAN_FORMULAS])
        self.assertEqual(main, '計画の集約', '出すクエリは接続文字列の Location')
        self.check_plan(I.blocks(pairs, main))

    def test_old_odc_with_package(self):
        import navi_mimport as I
        pairs, main = I.extract(odc_text(self.pairs, '計画の集約', old=True))
        self.assertEqual(sorted(n for n, _ in pairs), sorted(n for n, _ in self.pairs))
        self.check_plan(I.blocks(pairs, main))

    def test_workbook_with_datamashup(self):
        import base64
        import io
        import struct
        import zipfile
        import navi_mimport as I
        import openpyxl
        sec = 'section Section1;\r\n\r\n' + ''.join(f'shared #"{n}" = {f};\r\n\r\n' for n, f in self.pairs)
        inner = io.BytesIO()
        with zipfile.ZipFile(inner, 'w') as z:
            z.writestr('[Content_Types].xml', '<Types/>')
            z.writestr('Formulas/Section1.m', sec)
        pkg = struct.pack('<II', 0, len(inner.getvalue())) + inner.getvalue() + struct.pack('<III', 0, 0, 0)
        out = io.BytesIO()
        openpyxl.Workbook().save(out)
        book = io.BytesIO(out.getvalue())
        with zipfile.ZipFile(book, 'a') as z:
            xml = f'<?xml version="1.0" encoding="utf-16"?><DataMashup xmlns="http://schemas.microsoft.com/DataMashup">{base64.b64encode(pkg).decode()}</DataMashup>'
            z.writestr('customXml/item1.xml', xml.encode('utf-16'))
        blob = book.getvalue()
        self.assertEqual(I.kind_of(blob), 'workbook')
        pairs, main = I.extract(blob)
        self.assertEqual([n for n, _ in pairs], [n for n, _ in self.pairs])
        self.assertEqual(main, '', '読み込み先が無ければ、出すクエリは名前から決める（ヘルパーを除いた先頭）')
        self.check_plan(I.blocks(pairs))
        # Excel がシートへ読み込んだブック: 接続（Location は名前を "…" で囲むこともある）と queryTable がある
        ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        conns = ''.join(f'<connection id="{i}" name="クエリ - {n}" type="5"><dbPr connection="Provider=Microsoft.Mashup.OleDb.1;'
                        f'Data Source=$Workbook$;Location=&quot;{n}&quot;;Extended Properties=&quot;&quot;"/></connection>'
                        for i, n in enumerate(['サンプル ファイル', '計画の集約'], 1))
        loaded = io.BytesIO(blob)
        with zipfile.ZipFile(loaded, 'a') as z:
            z.writestr('xl/connections.xml', f'<connections {ns}>{conns}</connections>')
            z.writestr('xl/queryTables/queryTable1.xml', f'<queryTable {ns} name="ExternalData_1" connectionId="2"/>')
        pairs, main = I.extract(loaded.getvalue())
        self.assertEqual(main, '計画の集約', '出すクエリは、シートに読み込んでいるもの（接続のみのクエリではない）')
        self.check_plan(I.blocks(pairs, main))
        with self.assertRaises(ValueError):
            I.extract(out.getvalue())

    def test_year_month_dates(self):
        for text, want in (('2024年10月', date(2024, 10, 1)), ('2024/10', date(2024, 10, 1)), ('24/10/01', date(2024, 10, 1)),
                           ('2024年 9月', date(2024, 9, 1))):
            self.assertEqual(C.parse_datetime_text(text, None, 'date'), want, text)
