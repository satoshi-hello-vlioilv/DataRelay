"""クエリ（Power Query の M の式）の出し入れと、対象1件ぶんの中間CSVを作る係（2.4.0）。

言語そのもの（読む・評価する・関数）は navi_m・navi_mlib・navi_msource にある。こちらは、それを控えDB
（m_queries の表）へ出し入れし、下読み（画面の「適用したステップ」とプレビュー）と、実行のときの中間CSVを作る。
仕上げ（変換・公開・控え・同時出力）は、RNE・テキスト・結合と同じ finish_one_job をそのまま通す。

クエリは読取マスタ・結合マスタと同じく、名前を付けて別に持ち、対象から選ぶ。1つのクエリに複数のクエリ
（Excel の補助クエリ一式）を入れられ、どれを出すかは対象ごとに選べる（空ならクエリの既定）。
本体の名前を借りるので、navi_joinrun と同じく app.py の終わりで取り込む。
"""
import csv
import json
import re
import time
from datetime import datetime
from pathlib import Path

import app
import navi_m as M
import navi_mconv as C
import navi_mimport
import navi_paths
import navi_post
from app import log, settings_sync_lock


def _borrow(name):
    """本体の関数は、取り込んだ時点の実体ではなく、呼ばれた時点で引く（navi_split と同じ約束）。"""
    def call(*a, **k):
        return getattr(app, name)(*a, **k)
    call.__name__ = name
    call.__qualname__ = 'app.' + name
    return call


for _n in ('_mark_settings_dirty', 'flush_local_to_master_async', 'init_settings_db', 'load', 'normalize_job_source',
           'phase_log', 'read_extract', 'resolve_path', 'settings_connection'):
    globals()[_n] = _borrow(_n)
del _n

EXPORT_KIND = 'datarelay-m-queries'
EXPORT_VERSION = 1
MAX_CODE = 2_000_000          # 1つのクエリの式の上限（文字数）。貼り間違いで何十MBも入れないための歯止め
PREVIEW_ROWS = 200            # 下読みで画面へ返す行の上限
PREVIEW_ROW_CAP = 1000        # 下読みで1つのシート・CSV から読む行の上限（Power Query の下読みと同じ考え方）


# ---- 取り決め（1件）の形 ---------------------------------------------------------
def normalize_query(d):
    d = dict(d or {})
    code = str(d.get('code') or '').replace('\r\n', '\n').replace('\r', '\n')
    return {'id': str(d.get('id') or '').strip(), 'name': str(d.get('name') or '').strip(),
            'description': str(d.get('description') or '').strip(), 'code': code,
            'output': str(d.get('output') or '').strip(), 'updated_at': str(d.get('updated_at') or '')}


def validate_query(q):
    """保存してよいか。駄目な理由の一覧（空なら保存できる）。式が読めなければ何行目かを付けて断る。"""
    bad = []
    if not q['name']:
        bad.append('名前がありません')
    if not q['code'].strip():
        bad.append('式がありません（Power Query の詳細エディターの中身を貼り付けてください）')
        return bad
    if len(q['code']) > MAX_CODE:
        bad.append(f'式が長すぎます（{len(q["code"]):,}文字。上限 {MAX_CODE:,}）')
        return bad
    try:
        doc = M.Document.from_text(q['code'], q['name'] or 'クエリ1')
    except M.MError as e:
        bad.append(f'式が読めません: {e}')
        return bad
    if q['output'] and q['output'] not in doc.names():
        bad.append(f'出すクエリ「{q["output"]}」が式の中にありません（あるのは {"・".join(doc.names())}）')
    return bad


def document(q):
    return M.Document.from_text(q['code'], q.get('name') or 'クエリ1')


def output_name(doc, q, job=None):
    """どのクエリを出すか。対象で選んでいればそれ、無ければクエリの既定、それも無ければ式から決める。"""
    want = str((job or {}).get('query_out') or '').strip() or q.get('output') or ''
    if want and want in doc.names():
        return want
    if want:
        raise M.MError(f'出すクエリ「{want}」が、クエリ「{q.get("name")}」の式の中にありません'
                       f'（あるのは {"・".join(doc.names())}）。')
    return doc.main_name()


# ---- 控えDBとの出し入れ ----------------------------------------------------------
def _row(r):
    return normalize_query({'id': r['id'], 'name': r['name'], 'description': r['description'], 'code': r['code'],
                            'output': r['output_query'], 'updated_at': r['updated_at']})


def _load_m_queries(c):
    try:
        return [_row(r) for r in c.execute('SELECT * FROM m_queries ORDER BY display_order,name')]
    except Exception:
        log.exception('M_QUERY_LOAD_FAILED')
        return []


def load_m_queries():
    init_settings_db()
    with settings_connection() as c:
        return _load_m_queries(c)


def find_m_query(query_id):
    qid = str(query_id or '').strip()
    if not qid:
        return None
    return next((x for x in load_m_queries() if x['id'] == qid), None)


def save_m_query(d):
    """1件を登録・更新する。id が無ければ新規。保存できない理由があれば ValueError。"""
    q = normalize_query(d)
    bad = validate_query(q)
    if bad:
        raise ValueError('／'.join(bad))
    init_settings_db()
    qid = q['id'] or _uuid4()
    now = datetime.now().isoformat(timespec='seconds')
    with settings_sync_lock, settings_connection() as c:
        same = c.execute('SELECT id FROM m_queries WHERE name=? AND id<>?', (q['name'], qid)).fetchone()
        if same:
            raise ValueError(f'同じ名前のクエリ「{q["name"]}」がもうあります（名前で結び直すので、重ねられません）')
        order = c.execute('SELECT COALESCE(MAX(display_order),-1)+1 n FROM m_queries').fetchone()['n']
        cur = c.execute('SELECT display_order FROM m_queries WHERE id=?', (qid,)).fetchone()
        c.execute('INSERT OR REPLACE INTO m_queries (id,display_order,name,description,code,output_query,updated_at)'
                  ' VALUES(?,?,?,?,?,?,?)',
                  (qid, cur['display_order'] if cur else order, q['name'], q['description'], q['code'], q['output'], now))
        _mark_settings_dirty()
    flush_local_to_master_async('m-query-save')
    log.info('M_QUERY_SAVE id=%s name=%s chars=%s queries=%s output=%s', qid, q['name'], len(q['code']),
             len(document(q).names()), q['output'] or '(既定)')
    return dict(q, id=qid, updated_at=now)


def m_query_usage(query_id, cfg=None):
    """このクエリを使っている対象の名前。消してよいかの判断に要る。"""
    qid = str(query_id or '').strip()
    if not qid:
        return []
    jobs = (cfg or load()).get('jobs') or []
    # 入力に使う対象と、後処理（2.6.0）に使う対象の両方。後処理で使っている物を消すと、その対象は仕上げの前で止まる
    return [j.get('name', '') for j in jobs
            if (normalize_job_source(j.get('source')) == 'mquery' and str(j.get('query_id') or '') == qid)
            or (navi_post.normalize(j.get('post')).get('mode') == 'query' and navi_post.normalize(j.get('post')).get('query_id') == qid)]


def delete_m_query(query_id):
    """使っている対象があるうちは消さない。消すと、その対象は式を失って必ず失敗する。"""
    qid = str(query_id or '').strip()
    used = m_query_usage(qid)
    if used:
        raise ValueError(f'このクエリは{len(used)}件の対象が使っています（{"、".join(used[:3])}）。先に対象側を切り替えてください')
    init_settings_db()
    with settings_sync_lock, settings_connection() as c:
        n = c.execute('DELETE FROM m_queries WHERE id=?', (qid,)).rowcount
        _mark_settings_dirty()
    flush_local_to_master_async('m-query-delete')
    log.info('M_QUERY_DELETE id=%s deleted=%s', qid, n)
    return bool(n)


def _uuid4():
    import uuid
    return str(uuid.uuid4())


# ---- 持ち出しと取り込み ----------------------------------------------------------
def m_queries_export(items):
    return {'kind': EXPORT_KIND, 'version': EXPORT_VERSION, 'exported_at': datetime.now().isoformat(timespec='seconds'),
            'count': len(items or []),
            'queries': [{k: v for k, v in normalize_query(x).items() if k not in ('id', 'updated_at')}
                        for x in (items or [])]}


def m_queries_import(payload):
    """取り込む。読めたものと、読めなかった理由を返す。印（kind）を必ず確かめる。"""
    if isinstance(payload, (str, bytes)):
        try:
            payload = json.loads(payload if isinstance(payload, str) else payload.decode('utf-8-sig'))
        except Exception as e:
            return [], [f'JSONとして読めません: {e}']
    if isinstance(payload, list):
        payload = {'kind': EXPORT_KIND, 'queries': payload}
    if not isinstance(payload, dict) or str(payload.get('kind') or '') != EXPORT_KIND:
        return [], [f'クエリの持ち出しファイルではありません（kind={(payload or {}).get("kind") if isinstance(payload, dict) else "なし"}）']
    out, bad = [], []
    for i, x in enumerate(payload.get('queries') or [], 1):
        q = normalize_query(x)
        q['id'] = ''
        problems = validate_query(q)
        if problems:
            bad.append(f'{i}件目「{q["name"] or "名前なし"}」: ' + '／'.join(problems))
            continue
        out.append(q)
    if not out and not bad:
        bad.append('取り込めるクエリが1件もありませんでした')
    return out, bad


def import_file(blob, filename=''):
    """取り込むファイル → (クエリの取り決めの並び, 読めなかった理由)。

    .odc・.xlsx・.xlsm は Excel が書いたまま受ける（中の式を「// 名前」で区切った1つのクエリにまとめる。navi_mimport）。
    JSON は持ち出したもの、それ以外は式そのもののテキスト（ファイル名をクエリ名にする）。"""
    stem = Path(str(filename or 'クエリ')).stem or 'クエリ'
    kind = navi_mimport.kind_of(blob)
    if kind == 'json':
        return m_queries_import(blob.decode('utf-8-sig', errors='replace'))
    if kind == 'text':
        q = normalize_query({'name': stem, 'code': blob.decode('utf-8-sig', errors='replace')})
        bad = validate_query(q)
        return ([q], []) if not bad else ([], [f'「{stem}」: ' + '／'.join(bad)])
    try:
        pairs, main = navi_mimport.extract(blob)
    except Exception as e:
        return [], [f'「{filename or stem}」から式を取り出せませんでした: {e}']
    names = [n for n, _ in pairs]
    name = main if (kind == 'odc' and main) else stem
    q = normalize_query({'name': name, 'code': navi_mimport.blocks(pairs, main),
                         'output': main if main in names else '',
                         'description': f'{Path(str(filename)).name or stem} から取り込み（クエリ {len(pairs)}本）'})
    bad = validate_query(q)
    if bad:
        return [], [f'「{name}」: ' + '／'.join(bad)]
    log.info('M_QUERY_FROM_FILE kind=%s file=%s queries=%s output=%s', kind, filename, names, q['output'] or '(既定)')
    return [q], []


# ---- 評価 ------------------------------------------------------------------------
def resolve_m_path(raw):
    """式の中のファイル・フォルダーの場所。決まりはほかの入力と同じ（絶対・UNC はそのまま、相対は配布の置き場から）。"""
    return str(resolve_path(navi_paths.expand(raw)))


def host(preview=False, now=None):
    return M.Host(resolve=resolve_m_path, row_cap=PREVIEW_ROW_CAP if preview else None, now=now)


def as_table(v, name):
    """クエリの結果を表にする。リストは1列、レコードは Name・Value の表（Excel へ読み込むときと同じ）。"""
    v = M.force(v)
    if isinstance(v, M.MTable):
        return v
    if isinstance(v, list):
        return M.MTable(['Column1'], [[x] for x in v])
    if isinstance(v, M.MRecord):
        return M.MTable(['Name', 'Value'], [[n, v.raw(n)] for n in v.names])
    if isinstance(v, M.MFunction):
        raise M.MError(f'クエリ「{name}」の結果は関数です。ファイルへ出すのは表を返すクエリを選んでください。')
    return M.MTable(['Column1'], [[v]])


_TYPE_OF_KIND = {'date': 'date', 'datetime': 'datetime', 'datetimezone': 'datetime'}
_INT_FACETS = ('Int64', 'Int32', 'Int16', 'Int8', 'Byte')


def column_type(t, values):
    """M の列の型 → 出力の列の型（navi_output の語彙: text・integer・numeric・date・datetime）。

    型を決めていない列（any）は、中身から決める（Excel のセルの数・日付はそのまま数・日付として出す）。"""
    if t is not None and t.kind != 'any':
        if t.kind == 'number':
            return 'integer' if t.facet in _INT_FACETS else 'numeric'
        return _TYPE_OF_KIND.get(t.kind, 'text')
    seen = set()
    for v in values:
        if v is None or isinstance(v, M.ErrorCell):
            continue
        if isinstance(v, bool) or isinstance(v, str):
            return 'text'
        if isinstance(v, int):
            seen.add('integer')
        elif isinstance(v, float):
            seen.add('numeric')
        elif isinstance(v, datetime):
            seen.add('datetime')
        elif hasattr(v, 'year') and hasattr(v, 'day'):
            seen.add('date')
        else:
            return 'text'
    if not seen:
        return 'text'
    if seen <= {'integer'}:
        return 'integer'
    if seen <= {'integer', 'numeric'}:
        return 'numeric'
    if seen <= {'date'}:
        return 'date'
    if seen <= {'date', 'datetime'}:
        return 'datetime'
    return 'text'


def _forced_rows(t):
    """セルの遅延の値を求めておく（誤りはセルの誤りのまま）。"""
    rows = []
    for r in t.rows:
        out = []
        for v in r:
            if isinstance(v, M.Thunk):
                try:
                    v = v.get()
                except M.MError as e:
                    v = M.ErrorCell(e)
            out.append(v)
        rows.append(out)
    return rows


def _out_cell(v):
    if isinstance(v, M.ErrorCell):
        return ''
    if isinstance(v, (M.MTable, M.MRecord, list, M.MFunction)) or getattr(v, 'is_binary', False):
        return C.cell_text(None) if v is None else M.type_name(v).capitalize()
    return C.cell_text(v)


def write_table_csv(t, dst, encoding='utf-8-sig'):
    """表を、変換に渡す中間CSV（見出し1行）へ書く。→ (行, 列の型, 誤りのセル数, 列ごとの誤りの例, 入れ子の列)。

    誤りのセルは空にして数え（Excel へ読み込んだときと同じく、行は残す）、入れ子（表・レコード・リスト）は型の名前で書く。
    クエリの入力と後処理（2.6.0）で同じ書き方にする（片方だけ直る壊れ方をしないため）。"""
    rows = _forced_rows(t)
    types = {c: column_type(ty, [r[i] for r in rows]) for i, (c, ty) in enumerate(zip(t.columns, t.types))}
    errors, samples, nested = 0, {}, set()
    dst = Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    with dst.open('w', encoding=encoding, newline='') as f:
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        w.writerow(t.columns)
        for r in rows:
            for i, v in enumerate(r):
                if isinstance(v, M.ErrorCell):
                    errors += 1
                    samples.setdefault(t.columns[i], v.error.message)
                elif isinstance(v, (M.MTable, M.MRecord, list)):
                    nested.add(t.columns[i])
            w.writerow([_out_cell(v) for v in r])
    return rows, types, errors, samples, sorted(nested)


def write_m_intermediate(q, cfg, dst, job=None, reject_zero=False, encoding='utf-8-sig', now=None):
    """クエリを評価し、変換に渡す中間CSVへ書く。→ (行数, 列数, stat)。

    誤りのセルは空にして数え（Excel へ読み込んだときと同じく、行は残す）、どの列のどんな誤りかを stat に残す。"""
    started = time.perf_counter()
    doc = document(q)
    name = output_name(doc, q, job)
    h = host(False, now)
    t = as_table(doc.evaluate(name, host_=h), name)
    if reject_zero and not t.rows:
        raise ValueError(f'クエリ「{name}」の結果が0件のため出力を中止しました')
    rows, types, errors, samples, nested = write_table_csv(t, dst, encoding)
    stat = {'query': q['name'], 'output': name, 'rows': len(rows), 'columns': len(t.columns), 'error_cells': errors,
            'error_samples': samples, 'nested_columns': nested, 'types': types, 'folders': list(h.files_read),
            'elapsed': round(time.perf_counter() - started, 3)}
    log.info('M_QUERY_READ query=%s output=%s rows=%s columns=%s error_cells=%s nested=%s elapsed=%.2fs 中間CSV=%s',
             q['name'], name, len(rows), len(t.columns), errors, nested, stat['elapsed'], dst)
    if errors:
        log.warning('M_QUERY_ERROR_CELLS query=%s 箇所=%s 誤りのセルを空にしました 例=%s', q['name'], errors,
                    '、'.join(f'{k}: {v}' for k, v in list(samples.items())[:3]))
    typed = {k: v for k, v in types.items() if v != 'text'}
    if typed:
        log.info('M_QUERY_COLUMN_TYPES query=%s 型を決めた列=%s 内訳=%s', q['name'], len(typed), typed)
    return len(rows), len(t.columns), stat


# ---- 下読み（画面）--------------------------------------------------------------
def _show(v):
    """画面に出すセル。null は None、誤りは {'e': 説明}、入れ子は {'n': 'Table'} など（Power Query の見え方）。"""
    if isinstance(v, M.Thunk):
        try:
            v = v.get()
        except M.MError as e:
            return {'e': e.message}
    if isinstance(v, M.ErrorCell):
        return {'e': v.error.message}
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return C.cell_text(v) if (v != v or v in (float('inf'), float('-inf'))) else v
    if isinstance(v, (M.MTable, M.MRecord, list, M.MFunction, M.MType)) or getattr(v, 'is_binary', False):
        return {'n': M.type_name(v).capitalize()}
    return C.to_text(v)


def _error_view(e):
    out = {'message': e.message if isinstance(e, M.MError) else str(e), 'reason': getattr(e, 'reason', ''),
           'query': getattr(e, 'query', None) or '', 'step': getattr(e, 'step', None) or '',
           'where': e.where() if isinstance(e, M.MError) else ''}
    origin = getattr(e, 'origin', None)
    if origin and origin != (out['query'], out['step']):
        out['origin_query'], out['origin_step'] = origin
    if isinstance(e, M.MSyntaxError):
        out.update(line=e.line, col=e.col)
    return out


def preview(code, name='クエリ1', query=None, step=None, limit=PREVIEW_ROWS, full=False, now=None, inputs=None):
    """画面の下読み。式を読み、クエリの並び・各クエリのステップ・選んだステップの結果（先頭 limit 行）を返す。

    読めない・評価できないときも ok=False で、どこ（何行目・どのクエリのどのステップ）かを返す。"""
    started = time.perf_counter()
    out = {'ok': False, 'queries': [], 'query': '', 'step': '', 'columns': [], 'rows': [], 'row_count': 0,
           'error_cells': 0, 'notes': [], 'kind': ''}
    try:
        doc = M.Document.from_text(code, name or 'クエリ1', inputs=inputs)
    except M.MError as e:
        out['error'] = _error_view(e)
        return out
    out['uses_input'] = doc.uses_input()
    main = doc.main_name()
    helpers = set(doc.helpers())
    out['main'] = main
    out['queries'] = [{'name': n, 'steps': doc.steps(n), 'deps': doc.dependencies(n), 'main': n == main,
                       'kind': doc.kind(n), 'spans': [[a, b] for _s, a, b in doc.step_spans(n)],
                       'helper': n in helpers}
                      for n in doc.names()]
    out['missing'] = doc.missing_names()
    if out['missing']:
        out['notes'].append(missing_note(out['missing']))
    qn = query if query in doc.names() else main
    steps = doc.steps(qn)
    st = step if step in steps else None
    out['query'], out['step'] = qn, st or (steps[-1] if steps else '')
    h = host(not full, now)
    for n in out['notes']:
        h.note(n)
    try:
        v = doc.evaluate(qn, host_=h, step=st)
        v = M.force(v)
    except M.MError as e:
        out['error'] = _error_view(e)
        out['notes'] = list(h.notes)
        out['elapsed'] = round(time.perf_counter() - started, 3)
        return out
    out['kind'] = M.type_name(v)
    if isinstance(v, (M.MTable, list, M.MRecord)):
        t = as_table(v, qn)
        rows = _forced_rows(t)
        out['columns'] = [{'name': c, 'type': (ty.name() if ty is not None else 'any'),
                           'out': column_type(ty, [r[i] for r in rows])} for i, (c, ty) in enumerate(zip(t.columns, t.types))]
        out['row_count'] = len(rows)
        out['error_cells'] = sum(1 for r in rows for x in r if isinstance(x, M.ErrorCell))
        out['rows'] = [[_show(x) for x in r] for r in rows[:max(1, int(limit or PREVIEW_ROWS))]]
    else:
        out['value'] = _show(v)
    out['ok'] = True
    out['notes'] = list(h.notes)
    out['elapsed'] = round(time.perf_counter() - started, 3)
    return out


def missing_note(names):
    """足りない名前の知らせ。詳細エディターからはメインのクエリしか写せないので、取り込み方まで言う。"""
    shown = '」「'.join(names[:5]) + ('」ほか' if len(names) > 5 else '」')
    return (f'この式は「{shown}を使っていますが、ここに入っていません。詳細エディターからはメインのクエリしか写せないので、'
            'Excel の「接続ファイルのエクスポート」（.odc）か、ブック（.xlsx）を「取り込み」で読み込むと、補助のクエリも一緒に入ります'
            '（クエリの一覧で各クエリを選んで「// 名前」の行に続けて貼っても入ります）。')


# ---- 見本のファイル（「フォルダーから結合」のサンプル ファイル）--------------------------
_NAV = re.compile(r'^(?P<base>.+?)\{\s*(?P<i>\d+)\s*\}\s*\[\s*Content\s*\]\s*$', re.S)


def _unquote(name):
    name = name.strip()
    return name[2:-1].replace('""', '"') if name.startswith('#"') and name.endswith('"') else name


def sample_choices(code, name='クエリ1', query=None, limit=500):
    """見本のファイルを選ぶための候補。サンプル ファイルのクエリの最後のステップ（…{番号}[Content]）を探し、その元の表の
    ファイルを、元のファイル名で並べる（名前を置き換えたあとでも、どのファイルか分かるように）。

    → {'ok', 'query', 'step', 'index', 'span': [書き換える番号の始まり, 終わり], 'files': [{'i', 'name', 'folder'}]}。
    画面は span の数字を選んだ番号に書き換える（Power Query で見本のファイルを変えるのと同じ。式の中に残る）。"""
    out = {'ok': False, 'files': []}
    try:
        doc = M.Document.from_text(code, name or 'クエリ1')
    except M.MError as e:
        out['error'] = str(e)
        return out
    names = [query] if query in doc.names() else sorted(doc.names(), key=lambda n: 0 if re.match(r'^(サンプル ファイル|Sample File)', n) else 1)
    for qn in names:
        spans = doc.step_spans(qn)
        if not spans:
            continue
        step, a, b = spans[-1]
        text = code[a:b]
        m = _NAV.match(text.strip())
        if not m:
            continue
        base = _unquote(m.group('base'))
        if base not in doc.steps(qn) or base == step:
            continue
        lead = len(text) - len(text.lstrip())
        out.update(query=qn, step=step, base=base, index=int(m.group('i')),
                   span=[a + lead + m.start('i'), a + lead + m.end('i')])
        try:
            t = as_table(doc.evaluate(qn, host_=host(True), step=base), qn)
        except M.MError as e:
            out['error'] = str(e)
            return out
        ci = t.index().get('Content')
        ni = t.index().get('Name')
        fi = t.index().get('Folder Path')
        for i, r in enumerate(t.rows[:limit]):
            b_ = M.force(r[ci]) if ci is not None else None
            path = getattr(b_, 'path', None)
            label = Path(path).name if path else (str(M.force(r[ni])) if ni is not None else f'{i + 1}番目')
            folder = str(Path(path).parent) if path else (str(M.force(r[fi])) if fi is not None else '')
            out['files'].append({'i': i, 'name': label, 'folder': folder})
        out['total'] = len(t.rows)
        out['ok'] = True
        return out
    out['error'] = '見本のファイルを選ぶステップ（…{0}[Content]）がありません'
    return out


# ---- フォルダーとシートから式を作る（画面の「フォルダーから作る」）---------------------
def _m_text(s):
    return '"' + str(s).replace('"', '""') + '"'


def _m_name(s):
    s = str(s)
    import re
    if re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', s) and s not in M.KEYWORDS:
        return s
    return '#' + _m_text(s)


def build_folder_query(folder, sheet='', pattern='', extensions=('.xlsx', '.xlsm'), header=True, file_column=True,
                       subfolders=True, item_kind='Sheet'):
    """「フォルダーの Excel の、このシートをつなぐ」式を書く。Power Query の「フォルダーから結合」と同じ道筋。

    書いた式はそのまま画面に出して直せる（ここで作るのは式だけ。実行は貼り付けた式と同じ道を通る）。
    sheet が空なら各ブックの先頭のシート。pattern はファイル名に含む文字（空なら絞らない）。"""
    exts = '{' + ', '.join(_m_text(e.lower()) for e in extensions) + '}'
    lines = ['let',
             f'    ソース = Folder.Files({_m_text(folder)}),']
    prev = 'ソース'
    if not subfolders:
        lines.append(f'    #"直下のファイル" = Table.SelectRows({prev}, each [Folder Path] = {_m_text(_folder_path(folder))}),')
        prev = '#"直下のファイル"'
    cond = [f'List.Contains({exts}, Text.Lower([Extension]))', '[Attributes]?[Hidden]? <> true',
            'not Text.StartsWith([Name], "~$")']
    if pattern:
        cond.append(f'Text.Contains([Name], {_m_text(pattern)})')
    lines.append(f'    #"抽出したファイル" = Table.SelectRows({prev}, each {" and ".join(cond)}),')
    if sheet:
        pick = f'Excel.Workbook([Content], {"true" if header else "null"}, true){{[Item={_m_text(sheet)},Kind={_m_text(item_kind)}]}}[Data]'
    else:
        pick = f'Table.SelectRows(Excel.Workbook([Content], {"true" if header else "null"}, true), each [Kind] = "Sheet"){{0}}[Data]'
    lines.append(f'    #"シートのデータ" = Table.AddColumn(#"抽出したファイル", "データ", each {pick}),')
    if file_column:
        lines.append('    #"ファイル名の列" = Table.SelectColumns(Table.RenameColumns(#"シートのデータ", {"Name", "Source.Name"}), '
                     '{"Source.Name", "データ"}),')
        lines.append('    #"列の名前" = List.Distinct(List.Combine(List.Transform(#"ファイル名の列"[データ], Table.ColumnNames))),')
        lines.append('    #"展開したデータ" = Table.ExpandTableColumn(#"ファイル名の列", "データ", #"列の名前")')
        last = '#"展開したデータ"'
    else:
        lines.append('    #"結合したデータ" = Table.Combine(#"シートのデータ"[データ])')
        last = '#"結合したデータ"'
    lines += ['in', f'    {last}']
    return '\n'.join(lines) + '\n'


def _folder_path(folder):
    s = str(folder)
    return s if s.endswith(('\\', '/')) else s + '\\'


def inspect_folder(folder, extensions=('.xlsx', '.xlsm'), limit=200):
    """フォルダーの中の Excel と、先頭のブックのシート名（画面で選ぶための候補）。"""
    root = Path(resolve_m_path(folder))
    out = {'ok': False, 'folder': str(root), 'files': [], 'sheets': [], 'total': 0}
    if not root.is_dir():
        out['error'] = f'フォルダーが見つかりません: {root}'
        return out
    files = []
    for p in sorted(root.rglob('*'), key=lambda x: str(x).casefold()):
        if p.is_file() and p.suffix.lower() in extensions and not p.name.startswith('~$'):
            files.append(p)
    out['total'] = len(files)
    out['files'] = [str(p.relative_to(root)) for p in files[:limit]]
    sheets, seen = [], set()
    for p in files[:5]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            for ws in wb.worksheets:
                if ws.title not in seen:
                    seen.add(ws.title)
                    sheets.append(ws.title)
            wb.close()
        except Exception as e:
            out.setdefault('unreadable', []).append(f'{p.name}: {e}')
    out['sheets'] = sheets
    out['ok'] = True
    return out


# ---- 材料（ほかの対象の出力を読んでいるか）-----------------------------------------
_PATH_FUNCS = {'File.Contents': 'file', 'Folder.Files': 'folder', 'Folder.Contents': 'folder'}


def literal_paths(code):
    """式に書いてある場所（File.Contents("…")・Folder.Files("…") の文字）。→ (ファイル, フォルダー)。"""
    files, folders = [], []
    try:
        queries = M.read_queries(code)
    except M.MError:
        return files, folders

    def walk(node):
        if isinstance(node, tuple) and node:
            if node[0] == 'call' and isinstance(node[1], tuple) and node[1][0] == 'id' and node[1][1] in _PATH_FUNCS:
                args = node[2]
                if args and args[0][0] == 'const' and isinstance(args[0][1], str):
                    (files if _PATH_FUNCS[node[1][1]] == 'file' else folders).append(args[0][1])
            for x in node[1:]:
                walk(x)
        elif isinstance(node, list):
            for x in node:
                walk(x)
    for q in queries:
        walk(q.node)
    return files, folders


def job_inputs(job):
    """対象が読む材料の場所（結合の待ち合わせに使う）。→ (ファイル, フォルダー)。"""
    q = find_m_query(job.get('query_id'))
    if not q:
        return [], []
    files, folders = literal_paths(q['code'])
    return ([Path(resolve_m_path(x)) for x in files], [Path(resolve_m_path(x)) for x in folders])


# ---- 後処理（2.6.0）: 取り出した表に M の式を当ててから出す ------------------------------
POST_SAMPLE_ROWS = 1000       # 後処理の下読みのために控える、後処理の前の表の行数


def _typed(v, kind):
    """中間CSVの文字を、列の型（読取マスタ・クエリで決めた型）の値へ戻す。戻せなければ文字のまま（誤りにしない）。"""
    try:
        if kind == 'integer':
            return int(v.replace(',', ''))
        if kind == 'numeric':
            x = float(v.replace(',', ''))
            return int(x) if x.is_integer() and '.' not in v and 'e' not in v.lower() else x
        if kind == 'datetime':
            return datetime.fromisoformat(v.replace('/', '-'))
        if kind == 'date':
            return datetime.fromisoformat(v.replace('/', '-')).date()
    except ValueError:
        pass
    return v


_M_TYPE = {'integer': ('number', 'Int64'), 'numeric': ('number', None), 'date': ('date', None), 'datetime': ('datetime', None)}


def input_table(headers, body, types=None):
    """取り出した表（中間データの見出しと行。どれも文字）→ M の表。

    空のセルは null にする（Power Query で読み込んだ空のセルと同じ。ピボット解除で空のセルを出さない・合計で数えない）。
    列の型が分かっている列（読取マスタ・クエリの入力）は、その型の値に戻して型も付ける。RNE の列は文字のまま。"""
    types = types or {}
    kinds = [str(types.get(h) or 'text') for h in headers]
    rows = []
    for r in body:
        row = []
        for i, k in enumerate(kinds):
            v = r[i] if i < len(r) else ''
            v = '' if v is None else str(v)
            row.append(None if v == '' else (v if k == 'text' else _typed(v, k)))
        rows.append(row)
    mtypes = [M.MType(_M_TYPE[k][0], True, _M_TYPE[k][1]) if k in _M_TYPE else None for k in kinds]
    return M.MTable([str(h) for h in headers], rows, mtypes)


def input_cell(row, i):
    """見本の1セルを、型を見積もるための値に（数に見える文字は数に。input_table と違い、列の型を知らない見本の列に使う）。"""
    v = row[i] if i < len(row) else ''
    if v is None or v == '':
        return None
    try:
        return int(str(v).replace(',', ''))
    except ValueError:
        try:
            return float(str(v).replace(',', ''))
        except ValueError:
            return str(v)


def post_document(post, inputs=None):
    """後処理の取り決め → (Document, 出すクエリの名前, 見出しの言葉, 手順とステップ名の対応)。使えなければ ValueError。"""
    p = navi_post.normalize(post)
    if p.get('mode') == 'query':
        q = find_m_query(p.get('query_id'))
        if not q:
            raise ValueError('後処理のクエリが選ばれていません（または削除されています）。対象の「後処理」で選び直してください')
        bad = validate_query(q)
        if bad:
            raise ValueError(f'後処理のクエリ「{q["name"]}」が使えません: ' + '／'.join(bad))
        doc = M.Document.from_text(q['code'], q['name'], inputs=inputs)
        return doc, output_name(doc, q, {'query_out': p.get('query_out')}), f'クエリ「{q["name"]}」', []
    bad = navi_post.problems(p)
    if bad:
        raise ValueError('／'.join(bad))
    code, made = navi_post.compile_steps(p['steps'])
    return M.Document.from_text(code, '後処理', inputs=inputs), '後処理', '手順 ' + str(len(p['steps'])), made


def run_post(post, headers, body, types, dst, reject_zero=False, now=None):
    """取り出した表に後処理を当て、変換に渡す中間CSVへ書く。→ (行数, 列数, stat)。"""
    started = time.perf_counter()
    src = input_table(headers, body, types)
    doc, name, label, _made = post_document(post, {M.INPUT_NAME: src})
    t = as_table(doc.evaluate(name, host_=host(False, now)), name)
    if reject_zero and not t.rows:
        raise ValueError(f'後処理（{label}）の結果が0件のため出力を中止しました')
    rows, out_types, errors, samples, nested = write_table_csv(t, dst)
    stat = {'label': label, 'summary': navi_post.summary(post, label), 'rows_in': len(body), 'columns_in': len(headers),
            'rows': len(rows), 'columns': len(t.columns), 'error_cells': errors, 'error_samples': samples,
            'nested_columns': nested, 'types': out_types, 'elapsed': round(time.perf_counter() - started, 3)}
    return len(rows), len(t.columns), stat


def _safe_id(job_id):
    return re.sub(r'[^0-9A-Za-z_-]', '_', str(job_id or ''))[:80] or 'job'


def post_sample_path(job_id):
    return Path(app.LOCAL_RUNTIME) / 'post_sample' / f'{_safe_id(job_id)}.json'


def save_post_sample(job, headers, body, types):
    """後処理の前の表の先頭を、この PC に控える（後処理の下読みと列の候補に使う）。失敗しても実行は止めない。"""
    try:
        path = post_sample_path(job.get('id'))
        path.parent.mkdir(parents=True, exist_ok=True)
        data = {'headers': list(headers), 'rows': [list(r) for r in body[:POST_SAMPLE_ROWS]], 'total_rows': len(body),
                'types': dict(types or {}), 'value': str(job.get('_variant_value') or ''),
                'saved_at': datetime.now().isoformat(timespec='seconds')}
        tmp = path.with_suffix('.tmp')
        tmp.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        tmp.replace(path)
    except Exception as e:
        log.warning('POST_SAMPLE_SAVE_FAILED job=%s error=%s', job.get('name'), e)


def post_process(j, cfg, intermediate, expected_rows, expected_cols, line=''):
    """finish_one_job の入口で呼ぶ。取り出した表（中間データ）に後処理を当て、新しい中間データに差し替える。

    → (後処理を当てた対象の写し, 新しい中間CSV, 行数, 列数, stat)。対象の写しは、見出しが1行の中間CSV を読む形にそろえる
    （RNE の集計表の段は、ここで読んだときに畳んである。もう一度畳むと見出しを読み飛ばす）。"""
    zero = bool(cfg['settings'].get('reject_zero_rows'))
    t = phase_log('post_process', job=j.get('name'), line=line) if line else phase_log('post_process', job=j.get('name'))
    headers, body = read_extract(intermediate, j, zero, expected_rows, expected_cols)
    types = dict(j.get('_column_types') or {}) if isinstance(j.get('_column_types'), dict) else {}
    save_post_sample(j, headers, body, types)
    dst = Path(intermediate).with_name(Path(intermediate).stem + '_post.csv')
    rows, cols, stat = run_post(j.get('post'), headers, body, types, dst, reject_zero=zero)
    j2 = {**j, 'type': '詳細データ', '_rne_layout': None, '_intermediate_encoding': 'utf-8-sig', '_column_types': stat['types']}
    phase_log('post_process', t, job=j.get('name'), rows=rows, columns=cols)
    log.info('POST_PROCESS job=%s value=%s how=%s rows=%s→%s columns=%s→%s error_cells=%s elapsed=%.2fs 中間CSV=%s',
             j.get('name'), j.get('_variant_value') or '', stat['summary'], stat['rows_in'], rows, stat['columns_in'], cols,
             stat['error_cells'], stat['elapsed'], dst)
    if stat['error_cells']:
        log.warning('POST_PROCESS_ERROR_CELLS job=%s 箇所=%s 誤りのセルを空にしました 例=%s', j.get('name'), stat['error_cells'],
                    '、'.join(f'{k}: {v}' for k, v in list(stat['error_samples'].items())[:3]))
    return j2, dst, rows, cols, stat


def post_note(stat):
    """実行結果の1行に添える言葉（例: / 後処理 ピボット解除 1,200件→85件）。"""
    if not stat:
        return ''
    note = f' / {stat["summary"]} {stat["rows_in"]:,}件→{stat["rows"]:,}件'
    if stat.get('error_cells'):
        note += f'（誤りのセル {stat["error_cells"]}か所を空で出しました）'
    return note


def post_sample(job, cfg):
    """後処理の下読みに使う、後処理の前の表。→ {'headers','rows','types','total_rows','source','saved_at'} か None。

    この PC で後処理を当てた回の控え（いちばん確か）→ 後処理を使っていない対象なら、いまの出力ファイル（＝後処理の前の表）。"""
    path = post_sample_path(job.get('id'))
    if path.is_file():
        try:
            d = json.loads(path.read_text(encoding='utf-8'))
            d['source'] = 'sample'
            return d
        except Exception as e:
            log.warning('POST_SAMPLE_READ_FAILED job=%s error=%s', job.get('name'), e)
    if navi_post.active(job):
        return None
    try:
        files = app.viewer_files(job, cfg)
        one = next((x for x in files if Path(x['path']).is_file()), None)
        if not one:
            return None
        p = Path(one['path'])
        _fmt, headers, rows, total, _cols = app.read_preview_data(app.read_copy(p), job, limit=POST_SAMPLE_ROWS)
        return {'headers': headers, 'rows': rows, 'types': {}, 'total_rows': total, 'source': 'output',
                'saved_at': datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec='seconds'), 'value': one.get('value') or ''}
    except Exception as e:
        log.info('POST_SAMPLE_OUTPUT_UNREADABLE job=%s error=%s', job.get('name'), e)
        return None


def _step_info(code, made, count, inputs, stop=None):
    """手順ごとの、当てたあとの表の大きさと列（画面の流れの札と、次の手順の列の候補に使う）。止まった手順から先は None。"""
    info = [None] * count
    if inputs is None:
        return info
    doc = M.Document.from_text(code, '後処理', inputs=inputs)
    last = {}
    for i, n in made:
        last[i] = n
    for i in range(count):
        if stop is not None and i >= stop:
            break
        try:
            t = as_table(doc.evaluate('後処理', host_=host(True), step=last[i]), '後処理')
        except (M.MError, KeyError):
            break
        info[i] = {'rows': len(t.rows), 'columns': list(t.columns),
                   'types': [column_type(ty, [r[k] for r in _forced_rows(t)]) for k, ty in enumerate(t.types)]}
    return info


def preview_post(post, sample, step=None, limit=PREVIEW_ROWS):
    """後処理の下読み。見本（後処理の前の表）に当てて、step 番目の手順（0 から。None なら最後）までの結果を返す。

    手順で組むときは、手順の番号ごとに「どこまで当てたか」を返す（画面で手順を押すと、その手順のあとの表を見せる）。"""
    p = navi_post.normalize(post)
    src = input_table(sample['headers'], sample['rows'], sample.get('types')) if sample else None
    inputs = {M.INPUT_NAME: src} if src is not None else None
    if p.get('mode') == 'query':
        q = find_m_query(p.get('query_id'))
        if not q:
            return {'ok': False, 'error': {'message': 'クエリを選んでください'}, 'columns': [], 'rows': [], 'row_count': 0}
        out = preview(q['code'], q['name'], query=p.get('query_out') or q.get('output') or None, limit=limit, inputs=inputs)
        out['code'] = q['code']
    else:
        steps = p.get('steps') or []
        code, made = navi_post.compile_steps(steps)
        names = [n for _i, n in made]
        want = None
        if step is not None and steps:
            k = max(0, min(int(step), len(steps) - 1))
            hit = [n for i, n in made if i == k]
            want = hit[-1] if hit else None
        out = preview(code, '後処理', step=want, limit=limit, inputs=inputs)
        out['code'] = code
        out['step_names'] = [[i, n] for i, n in made]
        out['problems'] = navi_post.problems({**p, 'enabled': True}) if steps else []
        # 止まったステップを、画面の手順の番号へ戻す
        err = out.get('error') or {}
        if err.get('step') in names:
            out['error_step'] = next(i for i, n in made if n == err['step'])
        out['step_info'] = _step_info(code, made, len(steps), inputs, out.get('error_step'))
    out['input'] = {'columns': list(sample['headers']), 'types': [column_type(None, [input_cell(r, i) for r in sample['rows']]) if not (sample.get('types') or {}).get(h) else sample['types'][h] for i, h in enumerate(sample['headers'])], 'rows': len(sample['rows']), 'total_rows': sample.get('total_rows'),
                    'source': sample.get('source'), 'saved_at': sample.get('saved_at'), 'value': sample.get('value') or ''} if sample else None
    return out


def post_inputs_for(job_id='', query_id='', cfg=None):
    """クエリの編集で「取り出した表」を下読みするための見本。→ (inputs か None, 見本の説明か None)。

    対象を指定されればその対象の見本、無ければこのクエリを後処理に使っている対象のうち、見本があるもの。"""
    jobs = (cfg or load()).get('jobs') or []
    cands = [j for j in jobs if j.get('id') == job_id] if job_id else []
    if not cands and query_id:
        cands = [j for j in jobs if navi_post.normalize(j.get('post')).get('query_id') == query_id]
    for j in cands:
        sample = post_sample(j, cfg or load())
        if sample:
            info = {'job': j.get('name', ''), 'job_id': j.get('id'), 'rows': len(sample['rows']), 'total_rows': sample.get('total_rows'),
                    'source': sample.get('source'), 'saved_at': sample.get('saved_at'), 'value': sample.get('value') or ''}
            return {M.INPUT_NAME: input_table(sample['headers'], sample['rows'], sample.get('types'))}, info
    return None, None


_IN_TAIL = re.compile(r'^(\s*,?\s*)in\b\s*(.*?)\s*$', re.S)


def add_step(code, name='クエリ1', query=None, step=None):
    """式の中のクエリの最後に、手順（navi_post の手順1つ）を足す。→ {'ok', 'code', 'step'（足した最後のステップ名）} か {'ok': False, 'error'}。

    Power Query エディターで「変換」の札を押したときと同じく、いちばん最後のステップのあとに足し、in をそのステップにする。
    let で書かれていないクエリ（1つの式だけ）は、それを「ソース」にして let に包む。"""
    s = navi_post.normalize_step(step)
    if not s:
        return {'ok': False, 'error': '手順が分かりません'}
    bad = navi_post.step_problems(s, 1)
    if bad:
        return {'ok': False, 'error': bad[0].split(': ', 1)[-1]}
    try:
        doc = M.Document.from_text(code, name or 'クエリ1')
    except M.MError as e:
        return {'ok': False, 'error': f'式が読めません: {e.message}'}
    qn = query if query in doc.names() else doc.main_name()
    q = doc.by_name[qn]
    steps = doc.steps(qn)
    if q.node and q.node[0] == 'let' and q.node[2][0] == 'id' and q.node[2][1] == steps[-1]:
        last_end = doc.step_spans(qn)[-1][2]
        end = q.start + len(q.text)
        tail = _IN_TAIL.match(code[last_end:end])
        if not tail:
            return {'ok': False, 'error': 'クエリの終わり（in のあと）を読めません。詳細エディターで直してください'}
        parts = navi_post.step_parts(s, set(steps), navi_post.m_ref(steps[-1]))
        add = ''.join(f',\n    {navi_post.m_name(n)} = {x}' for n, x in parts)
        new = code[:last_end] + add + '\nin\n    ' + navi_post.m_name(parts[-1][0]) + code[end:]
    else:
        # let でない（または in が最後のステップでない）クエリは、全体を「ソース」にして包む
        body = code[q.start:q.start + len(q.text)]
        parts = navi_post.step_parts(s, {'ソース'}, 'ソース')
        lines = ['    ソース = ' + body.strip().replace('\n', '\n    ')] + [f'    {navi_post.m_name(n)} = {x}' for n, x in parts]
        new = code[:q.start] + 'let\n' + ',\n'.join(lines) + '\nin\n    ' + navi_post.m_name(parts[-1][0]) + code[q.start + len(q.text):]
    try:
        M.Document.from_text(new, name or 'クエリ1')
    except M.MError as e:
        return {'ok': False, 'error': f'足したあとの式が読めません: {e.message}'}
    return {'ok': True, 'code': new, 'query': qn, 'step': parts[-1][0], 'added': [n for n, _x in parts]}


POST_STARTER = 'let\n    ソース = ' + M.INPUT_NAME + '\nin\n    ソース'


def post_starter(steps=()):
    """後処理のクエリの書き始め（取り出した表を「ソース」に受け取る）。steps を渡すと、その手順を足した式。"""
    code = POST_STARTER
    for st in steps:
        r = add_step(code, '後処理', None, st)
        if r.get('ok'):
            code = r['code']
    return code
