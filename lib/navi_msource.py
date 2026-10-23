"""M のファイルを読む口（2.4.0）。Folder.Files・File.Contents・Excel.Workbook・Csv.Document など。

返す表の形（列名・並び・型）は Power Query と同じにする。貼り付けた式が [Extension]・[Attributes]?[Hidden]?・
{[Item="Sheet1",Kind="Sheet"]}[Data] のように列や行を名前で指すため、形が違うとそこで止まる。

Excel は openpyxl の読み取り専用の形で開く（大きなブックでも全部をメモリーへ広げない）。式の値（data_only）を読み、
#N/A などの誤りのセルは、Power Query と同じくセルの誤りにする。1回の評価の中では、同じファイルを2回開かない
（「フォルダーから結合」の式は、見本のファイルと本番の全ファイルで同じブックを開くため）。

ネットワークやデータベースへの口（Web.Contents・Sql.Database など）は持たない。使われていれば、名前を挙げて断る。
"""
import csv
import io
import json
import os
import re
import stat as stat_mod
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

import navi_mxlsx
from navi_m import ErrorCell, MError, MFunction, MRecord, MTable, MType, Thunk, force, host, type_name
from navi_mlib import _decode, _names, m


class MBinary:
    """バイナリの値。ファイルの場所だけを持ち、中身は読むときに初めて読む（フォルダーの全ファイルを先に読まない）。"""
    is_binary = True
    __slots__ = ('path', '_data')

    def __init__(self, path=None, data=None):
        self.path = str(path) if path is not None else None
        self._data = data

    def data(self):
        if self._data is None:
            try:
                self._data = Path(self.path).read_bytes()
            except OSError as e:
                raise MError(f'ファイルを読めませんでした: {self.path}（{e.strerror or e}）', 'DataSource.Error', self.path)
        return self._data

    def open(self):
        """openpyxl・csv へ渡す口。ファイルならその場所、そうでなければバイト列の流れ。"""
        if self._data is None and self.path:
            return self.path
        return io.BytesIO(self.data())

    def __repr__(self):
        return f'<binary {self.path or len(self._data or b"")}>'


def _resolve(path):
    p = force(path)
    if not isinstance(p, str) or not p.strip():
        raise MError('ファイル・フォルダーの場所（文字）を渡します。')
    h = host()
    got = h.resolve(p.strip()) if h else p.strip()
    return Path(str(got))


def _cache():
    h = host()
    if h is None:
        return {}
    c = getattr(h, 'cache', None)
    if c is None:
        c = h.cache = {}
    return c


# ---- フォルダー ------------------------------------------------------------------
_FILE_COLUMNS = ['Content', 'Name', 'Extension', 'Date accessed', 'Date modified', 'Date created', 'Attributes',
                 'Folder Path']
_FILE_TYPES = [MType('binary'), MType('text'), MType('text'), MType('datetime'), MType('datetime'),
               MType('datetime'), MType('record'), MType('text')]


def _attributes(p, st, name):
    attrs = getattr(st, 'st_file_attributes', None)
    hidden = bool(attrs & stat_mod.FILE_ATTRIBUTE_HIDDEN) if attrs is not None else name.startswith('.')
    system = bool(attrs & stat_mod.FILE_ATTRIBUTE_SYSTEM) if attrs is not None else False
    readonly = bool(attrs & stat_mod.FILE_ATTRIBUTE_READONLY) if attrs is not None else not os.access(p, os.W_OK)
    archive = bool(attrs & stat_mod.FILE_ATTRIBUTE_ARCHIVE) if attrs is not None else False
    ext = os.path.splitext(name)[1]
    return MRecord.of({'Content Type': _content_type(ext), 'Kind': _kind(ext), 'Size': st.st_size,
                       'ReadOnly': readonly, 'Hidden': hidden, 'System': system, 'Directory': False,
                       'Archive': archive, 'Device': False, 'Normal': False, 'Temporary': False})


def _content_type(ext):
    return {'.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            '.xlsm': 'application/vnd.ms-excel.sheet.macroEnabled.12', '.xls': 'application/vnd.ms-excel',
            '.csv': 'text/csv', '.txt': 'text/plain', '.json': 'application/json'}.get(ext.lower(), '')


def _kind(ext):
    e = ext.lower()
    if e in ('.xlsx', '.xlsm', '.xls', '.xlsb'):
        return 'Excel File'
    if e in ('.csv', '.txt', '.tsv'):
        return 'Text'
    return 'File' if e else ''


def _stamp(t):
    return datetime.fromtimestamp(t).replace(microsecond=0)


def _folder_path(d):
    s = str(d)
    return s if s.endswith(('\\', '/')) else s + ('\\' if '\\' in s or os.name == 'nt' else '/')


def _file_row(p, folder):
    st = p.stat()
    name = p.name
    return [MBinary(p), name, p.suffix, _stamp(st.st_atime), _stamp(st.st_mtime),
            _stamp(getattr(st, 'st_birthtime', st.st_ctime)), _attributes(p, st, name), _folder_path(folder)]


def _missing_folder(p):
    raise MError(f'フォルダーが見つかりませんでした: {p}', 'DataSource.NotFound', str(p))


# Folder.Files で並べるファイルの上限（2.5.3）。ドライブ全体（Folder.Files("C:\\")）のような指定で、際限なく歩いて
# 画面ごと止まらないようにする。業務のフォルダーはこれより十分少ない。
FOLDER_FILES_MAX = 200_000


@m('Folder.Files')
def folder_files(path, options=None):
    """フォルダーの中のファイル（下のフォルダーも含む）。並びは場所と名前の順（毎回同じ順にする）。"""
    root = _resolve(path)
    if not root.is_dir():
        _missing_folder(root)
    rows = []
    for d, subdirs, files in os.walk(root):
        subdirs.sort(key=str.casefold)
        if len(rows) + len(files) > FOLDER_FILES_MAX:
            raise MError(f'フォルダーの中のファイルが {FOLDER_FILES_MAX:,} 件を超えました。もっと狭いフォルダーを指定してください: {root}',
                         'Expression.Error', str(root))
        for name in sorted(files, key=str.casefold):
            p = Path(d) / name
            try:
                rows.append(_file_row(p, d))
            except OSError:
                continue
    h = host()
    if h is not None:
        h.files_read.append(str(root))
    return MTable(_FILE_COLUMNS, rows, _FILE_TYPES)


@m('Folder.Contents')
def folder_contents(path, options=None):
    """フォルダーの直下（ファイルとフォルダー）。フォルダーの Content は、その中身の表。"""
    root = _resolve(path)
    if not root.is_dir():
        _missing_folder(root)
    rows = []
    for p in sorted(root.iterdir(), key=lambda x: x.name.casefold()):
        try:
            if p.is_dir():
                st = p.stat()
                attrs = _attributes(p, st, p.name)
                attrs = MRecord(attrs.names, [True if n == 'Directory' else v for n, v in zip(attrs.names, attrs.vals)])
                rows.append([Thunk(lambda p=p: folder_contents(str(p))), p.name, '', _stamp(st.st_atime),
                             _stamp(st.st_mtime), _stamp(getattr(st, 'st_birthtime', st.st_ctime)), attrs,
                             _folder_path(root)])
            else:
                rows.append(_file_row(p, root))
        except OSError:
            continue
    return MTable(_FILE_COLUMNS, rows, [MType('any')] + _FILE_TYPES[1:])


@m('File.Contents')
def file_contents(path, options=None):
    p = _resolve(path)
    if not p.is_file():
        raise MError(f'ファイルが見つかりませんでした: {p}', 'DataSource.NotFound', str(p))
    return MBinary(p)


# ---- Excel ----------------------------------------------------------------------
_WB_COLUMNS = ['Name', 'Data', 'Item', 'Kind', 'Hidden']
_WB_TYPES = [MType('text'), MType('table'), MType('text'), MType('text'), MType('logical')]


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise MError('Excel を読む部品（openpyxl）が入っていません。config\\requirements.txt を入れ直してください。',
                     'DataSource.Error')


def _xls_message(where):
    return MError(f'古い形式の Excel（.xls）は読めません: {where}。Excel で .xlsx として保存し直してください。',
                  'DataFormat.Error', where)


def _x_value(v):
    """navi_mxlsx の値 → M の値（_cell_value と同じ: 誤りのセルは Error、空の文字は null）。"""
    if isinstance(v, navi_mxlsx.XError):
        return ErrorCell(MError(f'セルの値が誤りです（{v}）。', 'DataFormat.Error', str(v)).mark_origin())
    if v == '':
        return None
    return v


def _cell_value(c):
    v = c.value
    if c.data_type == 'e' and v is not None:
        return ErrorCell(MError(f'セルの値が誤りです（{v}）。', 'DataFormat.Error', v).mark_origin())
    if isinstance(v, str) and v == '':
        return None
    return v


@m('Excel.Workbook')
def excel_workbook(workbook, use_headers=None, delay_types=None):
    """ブックのシート・テーブル・名前付き範囲の一覧（Name・Data・Item・Kind・Hidden）。Data は使うときに読む。"""
    b = force(workbook)
    if not getattr(b, 'is_binary', False):
        raise MError(f'Excel.Workbook にはファイルの中身（File.Contents や [Content]）を渡します（{type_name(b)} が渡されました）。')
    where = b.path or '（メモリー上のブック）'
    if b.path and b.path.lower().endswith('.xls'):
        raise _xls_message(where)
    headers = force(use_headers)
    if isinstance(headers, MRecord):
        headers = headers.get('UseHeaders', None)
    headers = bool(headers)
    key = ('wb', b.path, _mtime(b.path), headers) if b.path else None
    cache = _cache()
    if key and key in cache:
        return cache[key]
    book = _Book(b, where)
    rows = []
    for ws_name, state in book.sheets():
        rows.append([ws_name, Thunk(lambda n=ws_name: book.sheet_table(n, headers)), ws_name, 'Sheet',
                     state != 'visible'])
    for name, sheet, ref in book.tables():
        rows.append([name, Thunk(lambda s=sheet, r=ref: book.range_table(s, r, True)), name, 'Table', False])
    for name, sheet, ref in book.defined_names():
        rows.append([name, Thunk(lambda s=sheet, r=ref: book.range_table(s, r, headers)), name, 'DefinedName', False])
    t = MTable(_WB_COLUMNS, rows, _WB_TYPES)
    if key:
        cache[key] = t
    return t


def _mtime(path):
    try:
        return os.stat(path).st_mtime_ns
    except (OSError, TypeError):
        return None


_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)(?::\$?([A-Za-z]{1,3})\$?(\d+))?$")


def _col_index(letters):
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


class _Book:
    """1つのブック。開くのは1回だけ（シートを読むたびに開き直さない）。
    値は navi_mxlsx で中の XML から直接・1回で読む（速い）。その形で読めないブックだけ openpyxl で読む。"""

    def __init__(self, binary, where):
        self.binary = binary
        self.where = where
        self._wb = None
        self._x = None
        self._zip_tables = None

    def x(self):
        """速い読み口（navi_mxlsx.Book）。読めない形なら None（openpyxl で読む）。"""
        if self._x is None:
            _openpyxl()
            try:
                self._x = navi_mxlsx.Book(self.binary.open())
            except navi_mxlsx.Unsupported:
                self._x = False
        return self._x or None

    def _x_grid(self, x, name, lo_r, lo_c, hi_r, hi_c):
        h = host()
        cap = h.row_cap if h else None
        grid, capped = x.grid(name, lo_r, lo_c, hi_r, hi_c, cap)
        if capped and h:
            h.note(f'下読みのため、各シートの先頭 {cap} 行だけを読んでいます（実行ではすべて読みます）。')
        return [[_x_value(v) for v in r] for r in grid]

    def wb(self):
        if self._wb is None:
            openpyxl = _openpyxl()
            src = self.binary.open()
            try:
                # 場所を渡すと、読み取り専用の openpyxl はファイルを開いたまま持つ（Windows では元の Excel を保存・削除できなくなる）。
                # 中身を渡して、ファイルは読み終えたら閉じる
                self._wb = openpyxl.load_workbook(io.BytesIO(self.binary.data()), read_only=True, data_only=True)
            except zipfile.BadZipFile:
                if isinstance(src, str) and src.lower().endswith('.xls') or self.binary.data()[:4] == b'\xd0\xcf\x11\xe0':
                    raise _xls_message(self.where)
                raise MError(f'Excel のファイルとして読めませんでした: {self.where}', 'DataFormat.Error', self.where)
            except MError:
                raise
            except Exception as e:
                raise MError(f'Excel のファイルとして読めませんでした: {self.where}（{e}）', 'DataFormat.Error', self.where)
        return self._wb

    def sheets(self):
        x = self.x()
        if x:
            return [(n, state) for n, state, _ in x.sheets]
        wb = self.wb()
        return [(ws.title, getattr(ws, 'sheet_state', 'visible')) for ws in wb.worksheets]

    def _grid(self, ws, min_row=None, min_col=None, max_row=None, max_col=None):
        h = host()
        cap = h.row_cap if h else None
        rows = []
        if max_row is not None and cap:
            max_row = min(max_row, min_row + cap - 1)
        for r in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            rows.append([_cell_value(c) for c in r])
            if cap and len(rows) >= cap:
                if h:
                    h.note(f'下読みのため、各シートの先頭 {cap} 行だけを読んでいます（実行ではすべて読みます）。')
                break
        return rows

    def sheet_table(self, name, headers):
        x = self.x()
        if x:
            if not x.has(name):
                raise MError(f'シート「{name}」がありません。', 'Expression.Error')
            lo_r, lo_c, hi_r, hi_c = x.used(name)
            if hi_r is None:
                return MTable([], [])
            return _make_table(self._x_grid(x, name, lo_r, lo_c, hi_r, hi_c), hi_c - lo_c + 1, headers)
        wb = self.wb()
        try:
            ws = wb[name]
        except KeyError:
            raise MError(f'シート「{name}」がありません。', 'Expression.Error')
        if not hasattr(ws, 'iter_rows'):
            return MTable([], [])
        lo_r, lo_c, hi_r, hi_c = self._bounds(ws)
        if hi_r is None:
            return MTable([], [])
        grid = self._grid(ws, lo_r, lo_c, hi_r, hi_c)
        width = hi_c - lo_c + 1
        grid = [(r + [None] * width)[:width] for r in grid]
        return _make_table(grid, width, headers)

    def _bounds(self, ws):
        """使っている範囲。ブックが書いている範囲（dimension）を使い、無ければ中身から求める。"""
        try:
            dim = ws.calculate_dimension()
        except Exception:
            dim = None
        m_ = _REF.match(dim or '')
        if m_ and m_.group(3) and dim not in ('A1:A1',):
            return int(m_.group(2)), _col_index(m_.group(1)), int(m_.group(4)), _col_index(m_.group(3))
        try:
            ws.reset_dimensions()
        except Exception:
            pass
        lo_r = lo_c = hi_r = hi_c = None
        for r in ws.iter_rows():
            for c in r:
                if c.value is None or c.value == '':
                    continue
                row, col = c.row, c.column
                lo_r = row if lo_r is None else min(lo_r, row)
                hi_r = row if hi_r is None else max(hi_r, row)
                lo_c = col if lo_c is None else min(lo_c, col)
                hi_c = col if hi_c is None else max(hi_c, col)
        return lo_r, lo_c, hi_r, hi_c

    def range_table(self, sheet, ref, headers):
        m_ = _REF.match(ref.replace('$', ''))
        if not m_:
            raise MError(f'範囲「{ref}」が読めません。')
        c1, r1 = _col_index(m_.group(1)), int(m_.group(2))
        c2, r2 = (_col_index(m_.group(3)), int(m_.group(4))) if m_.group(3) else (c1, r1)
        x = self.x()
        if x:
            if not x.has(sheet):
                raise MError(f'シート「{sheet}」がありません。')
            return _make_table(self._x_grid(x, sheet, r1, c1, r2, c2), c2 - c1 + 1, headers)
        wb = self.wb()
        try:
            ws = wb[sheet]
        except KeyError:
            raise MError(f'シート「{sheet}」がありません。')
        grid = self._grid(ws, r1, c1, r2, c2)
        width = c2 - c1 + 1
        grid = [(r + [None] * width)[:width] for r in grid]
        return _make_table(grid, width, headers)

    def tables(self):
        """テーブル（Excel の「テーブルとして書式設定」）。読み取り専用の形では openpyxl が読まないので、中の XML を読む。"""
        if self._zip_tables is not None:
            return self._zip_tables
        out = []
        try:
            src = self.binary.open()
            with zipfile.ZipFile(src) as z:
                names = set(z.namelist())
                wbx = ElementTree.fromstring(z.read('xl/workbook.xml'))
                ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
                      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
                rels = ElementTree.fromstring(z.read('xl/_rels/workbook.xml.rels'))
                target = {x.get('Id'): x.get('Target') for x in rels}
                for sh in wbx.findall('m:sheets/m:sheet', ns):
                    rid = sh.get('{%s}id' % ns['r'])
                    path = navi_mxlsx.part('xl', target.get(rid, ''))
                    rel_path = path.rsplit('/', 1)[0] + '/_rels/' + path.rsplit('/', 1)[1] + '.rels'
                    if rel_path not in names:
                        continue
                    srels = ElementTree.fromstring(z.read(rel_path))
                    for x in srels:
                        if not x.get('Type', '').endswith('/table'):
                            continue
                        tp = navi_mxlsx.part(path.rsplit('/', 1)[0], x.get('Target'))
                        if tp not in names:
                            continue
                        tx = ElementTree.fromstring(z.read(tp))
                        out.append((tx.get('displayName') or tx.get('name'), sh.get('name'), tx.get('ref')))
        except (KeyError, zipfile.BadZipFile, ElementTree.ParseError, OSError):
            out = []
        self._zip_tables = out
        return out

    def defined_names(self):
        x = self.x()
        if x:
            return x.destinations()
        out = []
        try:
            dn = self.wb().defined_names
            items = dn.items() if hasattr(dn, 'items') else [(d.name, d) for d in getattr(dn, 'definedName', [])]
            for name, d in items:
                if name.startswith('_xlnm') or getattr(d, 'hidden', False):
                    continue
                for sheet, ref in d.destinations:
                    out.append((name, sheet, ref))
                    break
        except Exception:
            return []
        return out


def _make_table(grid, width, headers):
    t = MTable([f'Column{i + 1}' for i in range(width)], grid)
    if headers and grid:
        from navi_mlib import table_promote_headers
        t = table_promote_headers(t, MRecord.of({'PromoteAllScalars': True}))
    return t


@m('Excel.CurrentWorkbook')
def excel_current_workbook():
    raise MError('Excel.CurrentWorkbook() は、Excel の中で動くときだけ使えます。このアプリでは、ブックを '
                 'Excel.Workbook(File.Contents("ファイルの場所"), null, true) の形で読む式に書き換えてください。',
                 'Expression.Error')


# ---- CSV・テキスト・JSON ----------------------------------------------------------
@m('Csv.Document')
def csv_document(source, columns=None, delimiter=None, extra_values=None, encoding=None):
    src = force(source)
    opts = force(columns)
    cols, quote, enc = None, 1, force(encoding)
    delim = force(delimiter)
    if isinstance(opts, MRecord):
        delim = opts.get('Delimiter', delim)
        cols = opts.get('Columns', None)
        enc = opts.get('Encoding', enc)
        q = opts.get('QuoteStyle', None)
        quote = 1 if q is None else int(q)
    else:
        cols = opts
    delim = ',' if delim in (None, '') else str(delim)
    if isinstance(src, str):
        text = src
    elif getattr(src, 'is_binary', False):
        text = _decode(src.data(), enc)
    else:
        raise MError(f'Csv.Document にはファイルの中身か文字を渡します（{type_name(src)} が渡されました）。')
    h = host()
    cap = h.row_cap if h else None
    rows = []
    if quote == 1:
        reader = csv.reader(io.StringIO(text, newline=''), delimiter=delim[0] if len(delim) == 1 else ',')
        if len(delim) != 1:
            reader = ([c for c in line.split(delim)] for line in text.splitlines())
    else:
        reader = (line.split(delim) for line in text.splitlines())
    for r in reader:
        rows.append(list(r))
        if cap and len(rows) >= cap:
            h.note(f'下読みのため、各ファイルの先頭 {cap} 行だけを読んでいます（実行ではすべて読みます）。')
            break
    while rows and rows[-1] == []:
        rows.pop()
    if isinstance(cols, list):
        names = _names(cols)
    elif cols is not None:
        names = [f'Column{i + 1}' for i in range(int(cols))]
    else:
        names = [f'Column{i + 1}' for i in range(max((len(r) for r in rows), default=0))]
    n = len(names)
    return MTable(names, [((r + [None] * n)[:n]) for r in rows], [MType('text')] * n)


@m('Json.Document')
def json_document(source, encoding=None):
    src = force(source)
    text = src if isinstance(src, str) else _decode(src.data(), force(encoding))
    try:
        return _from_json(json.loads(text))
    except json.JSONDecodeError as e:
        raise MError(f'JSON として読めませんでした（{e}）。', 'DataFormat.Error')


def _from_json(v):
    if isinstance(v, dict):
        return MRecord.of({k: _from_json(x) for k, x in v.items()})
    if isinstance(v, list):
        return [_from_json(x) for x in v]
    return v


@m('Binary.FromText')
def binary_from_text(text, encoding=None):
    import base64
    t = force(text)
    return None if t is None else MBinary(data=base64.b64decode(t))


@m('Binary.Length')
def binary_length(b):
    b = force(b)
    return None if b is None else len(b.data())


# ---- 持たない口 ------------------------------------------------------------------
def _unsupported(name, why):
    def f(*args):
        raise MError(f'{name} はこのアプリでは使えません（{why}）。', 'Expression.Error')
    from navi_mlib import _LIB
    _LIB[name] = MFunction(name, py=f, min_=0, max_=12)


for _name in ('Web.Contents', 'Web.Page', 'Web.BrowserContents', 'OData.Feed', 'SharePoint.Files',
              'SharePoint.Contents', 'SharePoint.Tables', 'Sql.Database', 'Sql.Databases', 'Oracle.Database',
              'Odbc.DataSource', 'Odbc.Query', 'OleDb.DataSource', 'OleDb.Query', 'Access.Database',
              'AnalysisServices.Database', 'PowerBI.Dataflows', 'Exchange.Contents', 'ActiveDirectory.Domains',
              'Salesforce.Data', 'Pdf.Tables', 'Xml.Tables', 'Xml.Document', 'Hdfs.Files',
              'AzureStorage.Blobs', 'Python.Execute', 'R.Execute', 'Html.Table', 'Parquet.Document'):
    _unsupported(_name, 'ファイルの場所から読む Folder.Files・File.Contents・Excel.Workbook・Csv.Document だけに対応しています')
del _name
