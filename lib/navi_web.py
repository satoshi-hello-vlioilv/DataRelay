"""HTTP層 ―― 画面からの求めに応える口。

app.py から分けてある。ここにあるのは受け口だけで、実際の処理は本体が持つ。
逆に本体はここを知らない。取り込みは app.py のいちばん最後で、しかもワーカーでは
行わない（抽出しかしないプロセスに受け口は要らない）。

本体の名前は下の import で明示的に受け取る。ルートの本文は分ける前と1文字も
変えていない ―― 書き換えると、f文字列の中は Python 3.11 では位置が正確でなく、
静かに壊れるため（実際に一度壊した）。

実行中に差し替わる変数（heartbeat_total / last_heartbeat_at / tray /
_first_request_logged）に触る4つの受け口は、実体で受け取ると古くなるので
app.py 側に残してある。あれは受け口の形をした本体の状態そのもの。

付け忘れが無いことは t_web.py が静的に確かめている。

なお、ここが受け取るのは取り込んだ時点の実体。あとから app.X を差し替えても
ここには届かない（Pythonでよくある「使われている場所で差し替える」の話）。
確認で本体の関数を差し替えるときは、こちらにも同じものを当てること。
"""
# send_file だけは本体が使っていないので、ここで直接受け取る。
# jsonify / render_template / request は本体経由で来る（下の import）。
import contextlib,shutil
from flask import send_file

import app
app=app.app          # 受け口を足す先（Flask本体）。以降 @app.get(...) は分ける前と同じ書き方
import navi_join     # キーの見当を付けるところだけ、直接呼ぶ（本体を経由する用が無い）
import navi_order    # 結合の順番と待ち合わせの判断
import navi_book     # マスタをEXCELで出し入れする

from app import (
    APP_ID, APP_NAME, APP_RELEASED_AT, APP_VERSION, APP_VERSION_TITLE, BASE, BUILD_VERSION, CHANGELOG,
    CLOSE_GRACE_SECONDS, DEFAULT_DLL_SEARCH_ROOTS, DOCS, INSPECT_ALL_ORDER, INSPECT_ALL_SPEC,
    INSPECT_TASK_SPECS, INSTANCE_ID, LOCAL_RUNTIME, LOG_FILTERS, LOG_PATH, LOG_READ_BYTES,
    PORT, Path, ROW_AXIS_MODE_LABEL, SPLIT_BATCH_MAX, _api_diag_cache_path, _dll_requirement,
    _log_api_exports,
    _read_api_diag_cache, _run_inspect_endpoint, _split_stage_logged, _split_trial_run,
    _viewer_output_path, _write_api_diag_cache, active_workers, active_workers_lock, alerts,
    alerts_lock, api_readiness, axis_balance_scores, blocked_row_axes, calendar,
    settings_sync_lock, settings_revision, SETTINGS_DB,
    cancel_requested, check_path_item, clamp_parallel_lines, clear_row_axis_blocks,
    column_cache_state, column_weights, command_queue, command_queue_lock, compute_period,
    creds, csv, datetime, dll_diagnostic_issues, dll_search_roots, docs_dir, duplicate_columns,
    enqueue_command, expand_rule_occurrences, find_nearby_file, freshness_view,
    has_template_variables, heartbeat_clients, heartbeat_lock, inspect_task_blank,
    inspect_task_lock, inspect_task_percent, inspect_task_seconds, inspect_tasks,
    job_extra_formats, job_output_plan, job_schedule_preview, json, jsonify, last_run_info, load, sqlite3,
    load_column_cache, load_job_runs, load_rne_timing, load_split_trials, log, log_files,
    machine_path_view, normalize_output_format, normalize_split_mode, normalize_split_shape,
    split_trial_options,
    path_setting_roles,
    os, output_extension, pick_anchor_columns, pick_row_axis_by_mode, plan_run_split,
    queue_snapshot, re, read_copy, read_header_names, read_log_lines, read_preview_data,
    recommend_split_parts, render_filename_segments, render_template, request,
    resolve_output_filename, resolve_path, resolve_rne_path, retry_lock, retry_view,
    retry_waiting, rne_master_view, rne_run_stats, rotate_log_if_needed, row_axis_choice,
    row_split_breakdown, row_split_candidates, run_inspect_worker, runtime_split_view, save,
    save_axis_survey, save_column_cache, save_column_classification, save_condition_items,
    settings_connection, split_batch_blank, split_batch_items, split_batch_lock,
    split_batch_run_data, split_batch_state, split_batch_summary, split_breakeven_share,
    split_gain_reason, split_incompatible, split_link_profile, split_payload_profile,
    split_trial_blank, split_trial_lock, split_trial_state, split_useful_parts, status, struct,
    task_target, threading, time, uuid, write_xlsx_direct,
    TEXT_ENCODINGS, TEXT_ENCODING_LABEL, TEXT_UNITS, TEXT_UNIT_LABEL, TRIM_MODES, TRIM_LABEL,
    JOB_SOURCE_LABEL, normalize_job_source, normalize_text_layout, validate_text_layout,
    text_layout_width, text_layout_overlaps, text_layout_gaps, preview_text, resolve_text_path,
    load_text_layouts, find_text_layout, save_text_layout, delete_text_layout, text_layout_usage,
    text_layouts_export, text_layouts_import,
    JOIN_TYPES, JOIN_TYPE_LABEL, JOIN_TYPE_NOTE, SOURCE_FORMATS, SOURCE_FORMAT_LABEL,
    JOIN_MAX_SOURCES, join_types_available, normalize_join_recipe, validate_join_recipe,
    preview_join_recipe, join_recipes_export, join_recipes_import, load_join_recipes,
    find_join_recipe, save_join_recipe, delete_join_recipe, join_recipe_usage,
    resolve_join_path, join_reader, join_layouts, read_preview_data,
    join_candidates, join_sample_reader, sampled_recipe, suggest_join_keys,
    job_output_paths, job_dependencies, job_wait_reasons,
    VIEWER_MAX_ROWS, viewer_row_budget, output_column_limit, check_output_columns,
    OUTPUT_FORMAT_LABEL,
    COLUMN_TYPES, COLUMN_TYPE_LABEL, COLUMN_TYPE_NOTE, format_keeps_types, normalize_column_type,
    STAMP_FORMAT_SAMPLES, DATE_FORMAT_DEFAULT, DATETIME_FORMAT_DEFAULT, MAX_SCALE,
    text_layout_column_types)

@app.get('/')
def index():
 response=app.make_response(render_template('index.html'));response.headers['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0';response.headers['Pragma']='no-cache';return response

@app.get('/favicon.ico')
def favicon():
 # アイコン実体は static/ にある。ICOを優先し、無い環境ではSVGへフォールバックする。
 for f,ctype in ((BASE/'static'/'favicon.ico','image/x-icon'),(BASE/'static'/'favicon.svg','image/svg+xml'),(BASE/'favicon.ico','image/x-icon')):
  if f.is_file():
   response=app.make_response(f.read_bytes());response.headers['Content-Type']=ctype;response.headers['Cache-Control']='public, max-age=86400';return response
 return ('',404)

@app.get('/api/config')
def get_config():
 c=load()
 try:*_,s=creds(resolve_path(c['symnavim_conf'])); c['credential_status']='読取可能 ['+s+']'
 except Exception as e:c['credential_status']='未確認: '+str(e)
 now=datetime.now()
 for j in c['jobs']:
  pattern=str(j.get('output_pattern') or '').strip()
  # 変数扱いは「命名モードがtemplate」かつ「既知の変数トークンが実在する」場合のみ。
  is_var=str(j.get('naming_mode') or 'fixed').lower()=='template' and has_template_variables(pattern)
  j['output_is_variable']=is_var
  if is_var:
   try:
    rp=None
    try:rp=str(resolve_rne_path(j,c))
    except Exception:rp=None
    fmt=normalize_output_format(j.get('output_format'),j.get('output_file'))
    j['output_file_preview']=resolve_output_filename(j,c,now)
    segs=render_filename_segments(pattern,job=j,rne_path=rp,now=now)
    segs.append({'text':output_extension(fmt),'var':False})
    j['output_file_segments']=segs
   except Exception:
    j['output_file_preview']='';j['output_file_segments']=[]
  # 1回の実行で実際に何ができるのか。名前まで組み立てて返す（画面で作り直さない）
  try:j['output_plan']=job_output_plan(j,c,now)
  except Exception:j['output_plan']=[]
 return jsonify(c)

@app.put('/api/config')
def put_config():
 # 中身を確かめずに save() へ渡すと、jobs が無いだけで _save_local が
 # 「DELETE FROM jobs」と「DELETE FROM job_runs」まで走らせ、全対象と実績が消える。
 # 画面は必ず全量（jobs を含む）を送るので、形が違う要求は保存せずに断る。
 data=request.get_json(silent=True)
 if not isinstance(data,dict) or not isinstance(data.get('jobs'),list):
  return jsonify(ok=False,error='設定の形が正しくありません（対象の一覧が含まれていません）'),400
 # 保存は全量の置き換え（対象も実績も、送られてきたものに合わせて消す）。画面を2つ開いて
 # いると、片方が持っている古い写しで上書きされ、もう片方が追加した対象が黙って消える。
 # 版が食い違うときは断り、画面に読み直してもらう。版を載せない相手（古い画面）は従来どおり。
 want=data.get('settings_revision')
 with settings_sync_lock:
  if want is not None:
   try:want=int(want)
   except (TypeError,ValueError):want=None
  now=settings_revision(SETTINGS_DB)
  if want is not None and want!=now:
   log.warning('CONFIG_PUT_CONFLICT client_rev=%s server_rev=%s（ほかの画面の変更を上書きしません）',want,now)
   return jsonify(ok=False,conflict=True,settings_revision=now,
                  error='ほかの画面で設定が変わっています。最新の内容を読み直してから、もう一度保存してください'),409
  rev=save(data)
 return jsonify(ok=True,settings_revision=rev)

@app.post('/api/settings/parallel-lines')
def set_parallel_lines():
 # 並列ライン数だけを即時に保存する軽量エンドポイント。ユーザーが変更したら他の未保存編集に触れずその値を確定し、次回起動以降も保持する。
 data=request.get_json(silent=True) or {}
 # 読む→1つ変える→書く、のあいだに全量保存が割り込むと、その内容ごと読んだ時点へ戻る。
 # ひとつながりにして、割り込む余地をなくす。
 with settings_sync_lock:
  c=load(); lines=clamp_parallel_lines(data.get('lines',2),c,default=2)
  c['settings']['api_parallel_lines']=lines; c['settings']['stability_profile']='balanced_api_parallel'; save(c,quiet=True)
 log.info('設定保存 job=(共通) 並列ライン数を保存 api_parallel_lines=%s',lines)
 return jsonify(ok=True,api_parallel_lines=lines)

def normalize_column_layout(v):
 """一覧の列の決め方を、そのまま置いても安全な形に整える。

 どんな列があるかを決めているのは画面の側（static/app.js の COLUMNS）。ここで列の名前を
 並べ直すと、列を1つ増やすたびに両方を直すことになり、片方だけ直した日から
 「保存はできるのに出てこない」が起きる。だから中身の当否は見ず、形と大きさだけ整える。
 """
 v=v if isinstance(v,dict) else {}
 def keys(x):
  out=[]
  for k in (x if isinstance(x,list) else [])[:32]:
   k=str(k)[:32]
   if k and k not in out:out.append(k)
  return out
 opt={}
 for k,o in list((v.get('opt') if isinstance(v.get('opt'),dict) else {}).items())[:32]:
  if not isinstance(o,dict):continue
  one={}
  for name,value in list(o.items())[:16]:
   if isinstance(value,bool) or value is None:one[str(name)[:32]]=value
   elif isinstance(value,(int,float)):one[str(name)[:32]]=value
   elif isinstance(value,str):one[str(name)[:32]]=value[:64]
   elif isinstance(value,list):one[str(name)[:32]]=[str(x)[:32] for x in value[:16]]
  opt[str(k)[:32]]=one
 return {'order':keys(v.get('order')),'hidden':keys(v.get('hidden')),'opt':opt}

@app.post('/api/settings/columns')
def set_columns():
 # 列の決め方だけを即時に保存する軽い受け口。ほかの未保存の編集には触らない
 # （並列ライン数と同じ約束）。
 data=request.get_json(silent=True) or {}
 # layout が無い body をそのまま通すと、normalize が空の3キーを作って保存済みの並びを消す。
 # 「全部既定に戻す」は画面が空の layout を明示して送ってくるので、鍵の有無で分けられる。
 if not isinstance(data.get('layout'),dict):return jsonify(ok=False,error='列の決め方(layout)がありません'),400
 with settings_sync_lock:
  c=load();layout=normalize_column_layout(data.get('layout'));c['settings']['column_layout']=layout;save(c,quiet=True)
 log.info('設定保存 job=(共通) 一覧の列 並び=%s 隠す=%s 見せ方=%s',
          layout['order'],layout['hidden'],layout['opt'])
 return jsonify(ok=True,column_layout=layout)

@app.get('/api/schedule-preview')
def schedule_preview():
 c=load(); now=datetime.now(); runs=load_job_runs()
 return jsonify(items=[{**job_schedule_preview(j,now),**last_run_info(runs.get(j['id']))} for j in c['jobs']])

@app.get('/api/freshness')
def freshness():
 return jsonify(**freshness_view())

@app.get('/api/job-trend/<job_id>')
def job_trend(job_id):
 """直近の実績の並び。件数が急に減った・だんだん遅くなっている、に気づくため。

 run_history には2000件ぶん溜まっているのに、これまではカレンダーの実施記録に
 しか使っていなかった。"""
 limit=max(5,min(200,int(request.args.get('limit') or 30)))
 rows=[]
 try:
  with settings_connection() as conn:
   for r in conn.execute('SELECT finished_at,status,rows,cols,detail FROM run_history WHERE job_id=? ORDER BY id DESC LIMIT ?',(job_id,limit)):
    sec=None
    m=re.search(r'/\s*([\d.]+)秒',str(r['detail'] or ''))
    if m:
     try:sec=float(m.group(1))
     except Exception:sec=None
    rows.append({'at':r['finished_at'],'status':r['status'],'rows':r['rows'],'cols':r['cols'],'seconds':sec})
 except Exception:
  log.exception('JOB_TREND_FAILED job=%s',job_id)
 rows.reverse()
 counts=[x['rows'] for x in rows if isinstance(x['rows'],int)]
 secs=[x['seconds'] for x in rows if isinstance(x['seconds'],float)]
 def spread(v):
  if len(v)<2:return {}
  return {'min':min(v),'max':max(v),'first':v[0],'last':v[-1],
          'change':round((v[-1]-v[0])/v[0]*100,1) if v[0] else None}
 return jsonify(ok=True,items=rows,count=len(rows),
                rows_spread=spread(counts),seconds_spread=spread(secs),
                failed=len([x for x in rows if x['status']=='failed']))

@app.get('/api/calendar')
def calendar_view():
 # カレンダービュー用: 指定月の予定（scheduled）と実施履歴（executed）を日付ごとに返す。
 try:year=int(request.args.get('year')); month=int(request.args.get('month'))
 except Exception:
  now=datetime.now(); year,month=now.year,now.month
 month=max(1,min(12,month))
 first=datetime(year,month,1); last_day=calendar.monthrange(year,month)[1]; last=datetime(year,month,last_day,23,59,59)
 c=load(); now=datetime.now()
 scheduled=[]; interval_summary={}
 for j in c['jobs']:
  if not j.get('enabled'):continue
  for r in j.get('schedules',[]):
   if not r.get('enabled'):continue
   points,interval_days=expand_rule_occurrences(r,first,last)
   for p in points:
    scheduled.append({'date':p.strftime('%Y-%m-%d'),'time':p.strftime('%H:%M'),'datetime':p.isoformat(timespec='minutes'),'job_id':j['id'],'job_name':j['name'],'rule_id':r.get('id'),'rule_name':r.get('name',''),'rule_type':r.get('type',''),'past':p<now})
   for d in interval_days:
    key=d['date']; entry=interval_summary.setdefault(key,{'date':key,'jobs':{},'total':0})
    # 対象IDだけを鍵にしていたため、1つの対象に間隔ルールが2本あると2本目が捨てられ、
    # 日セルの合計回数（全ルールぶん）と内訳が食い違っていた。規則ごとに持つ。
    entry['jobs'].setdefault((j['id'],r.get('id') or ''),{'job_id':j['id'],'job_name':j['name'],'rule_name':r.get('name',''),'minutes':d['minutes'],'count':d['count']})
    entry['total']+=d['count']
 interval_list=[{'date':v['date'],'total':v['total'],'items':list(v['jobs'].values())} for v in interval_summary.values()]
 # 実施履歴（追記式run_history）から当月分を取得。
 executed=[]
 try:
  with settings_connection() as conn:
   for row in conn.execute("SELECT id,job_id,job_name,finished_at,status,trigger,rows,cols,output_file FROM run_history WHERE finished_at>=? AND finished_at<=? ORDER BY finished_at",(first.strftime('%Y-%m-%dT00:00:00'),last.strftime('%Y-%m-%dT23:59:59'))):
    fa=str(row['finished_at'] or '')
    if len(fa)<10:continue
    trig=str(row['trigger'] or ''); kind='schedule' if trig.startswith('schedule') else 'manual'
    executed.append({'id':row['id'],'date':fa[:10],'time':fa[11:16],'datetime':fa,'job_id':row['job_id'],'job_name':row['job_name'],'status':row['status'],'trigger':kind,'rows':row['rows'],'cols':row['cols'],'output_file':row['output_file']})
 except Exception:
  log.exception('CALENDAR_HISTORY_FAILED')
 return jsonify(year=year,month=month,days_in_month=last_day,first_weekday=first.weekday(),today=now.strftime('%Y-%m-%d'),scheduled=scheduled,interval=interval_list,executed=executed)

@app.post('/api/schedule/quick-add')
def schedule_quick_add():
 # カレンダーから特定日の1回実行ルールを素早く追加する。既存の対象へspecific_datesルールを1件加える。
 d=request.get_json(force=True) or {}; job_id=d.get('job_id'); date=str(d.get('date') or '').strip(); tm=str(d.get('time') or '06:00').strip()
 if not job_id or not date:return jsonify(error='対象と日付を指定してください'),400
 # 形も、これから来る日時かも確かめる。過ぎた日時で登録すると、規則は残るのに一度も
 # 発火せず（schedule_key は当日その時刻の猶予内しか鍵を返さない）、有効な予定がある
 # 扱いになるためアプリが自動で終わらなくなる。
 try:_when=datetime.strptime(date+' '+(tm or '06:00'),'%Y-%m-%d %H:%M')
 except ValueError:return jsonify(error='日付は YYYY-MM-DD、時刻は HH:MM で指定してください'),400
 if _when<=datetime.now():return jsonify(error=f'{date} {tm} はすでに過ぎています。これから来る日時を指定してください'),400
 with settings_sync_lock:
  c=load(); j=next((x for x in c['jobs'] if x['id']==job_id),None)
  if not j:return jsonify(error='対象が見つかりません'),404
  rule={'id':uuid.uuid4().hex,'enabled':True,'name':d.get('name') or f'{date} 単発実行','type':'specific_dates','time':tm,'dates':[date]}
  j.setdefault('schedules',[]).append(rule); save(c,quiet=True)
 log.info('CALENDAR_QUICK_ADD job=%s date=%s time=%s',j['name'],date,tm)
 return jsonify(ok=True,rule=rule)

@app.post('/api/run-history/delete')
def delete_run_history():
 # カレンダーの実施記録（run_history）を削除する。id指定（複数可）または日付＋任意の対象指定に対応する。
 d=request.get_json(silent=True) or {}
 ids=[int(x) for x in (d.get('ids') or []) if str(x).strip().isdigit()]
 date=str(d.get('date') or '').strip(); job_id=str(d.get('job_id') or '').strip()
 removed=0
 try:
  with settings_connection() as conn:
   if ids:
    conn.execute('DELETE FROM run_history WHERE id IN ('+','.join('?' for _ in ids)+')',ids); removed=conn.total_changes
   elif date:
    if job_id and job_id!='all':
     cur=conn.execute("DELETE FROM run_history WHERE substr(finished_at,1,10)=? AND job_id=?",(date,job_id))
    else:
     cur=conn.execute("DELETE FROM run_history WHERE substr(finished_at,1,10)=?",(date,))
    removed=cur.rowcount if cur.rowcount is not None else conn.total_changes
   else:
    return jsonify(error='削除対象（idまたは日付）を指定してください'),400
 except Exception as e:
  log.exception('RUN_HISTORY_DELETE_FAILED'); return jsonify(error=str(e)),500
 log.info('RUN_HISTORY_DELETE ids=%s date=%s job_id=%s removed=%s',ids,date or '-',job_id or '-',removed)
 return jsonify(ok=True,removed=removed)

@app.post('/api/preview-filename')
def preview_filename():
 data=request.get_json(force=True) or {}; c=load(); now=datetime.now()
 job={'rne':data.get('rne') or '','rne_path':data.get('rne_path') or '','name':data.get('name') or '','table':data.get('table') or '','output_format':data.get('format') or 'sqlite3','output_file':data.get('output_file') or '','naming_mode':'template','output_pattern':data.get('pattern') or ''}
 rne_found=False; mtime=ctime=None
 try:
  rp=resolve_rne_path(job,c)
  if rp and Path(rp).is_file():
   rne_found=True; st=Path(rp).stat(); mtime=datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d %H:%M'); ctime=datetime.fromtimestamp(getattr(st,'st_ctime',st.st_mtime)).strftime('%Y-%m-%d %H:%M')
 except Exception:pass
 try:filename=resolve_output_filename(job,c,now)
 except Exception as e:return jsonify(ok=False,error=str(e)),400
 pattern=str(job.get('output_pattern') or '').strip()
 is_var=has_template_variables(pattern)
 fmt=normalize_output_format(job.get('output_format'),job.get('output_file'))
 try:
  rp2=str(resolve_rne_path(job,c))
 except Exception:rp2=None
 segments=render_filename_segments(pattern,job=job,rne_path=rp2,now=now) if is_var else []
 if segments:segments.append({'text':output_extension(fmt),'var':False})
 return jsonify(ok=True,filename=filename,is_variable=is_var,segments=segments,rne_found=rne_found,rne_mtime=mtime,rne_ctime=ctime,now=now.strftime('%Y-%m-%d %H:%M:%S'))

@app.post('/api/period-preview')
def period_preview():
 # 相対期間設定から、処理日時基準で実際に抽出される期間を試算して返す。
 d=request.get_json(force=True) or {}
 period={'enabled':True,'unit':(d.get('unit') if d.get('unit') in ('month','day') else 'month'),'from_offset':d.get('from_offset',0),'to_offset':d.get('to_offset',0),'control_point':d.get('control_point') or ''}
 spec=compute_period(period,datetime.now())
 if not spec:return jsonify(ok=False,error='期間を計算できません'),400
 return jsonify(ok=True,now=datetime.now().strftime('%Y-%m-%d %H:%M:%S'),control_point=period['control_point'],**spec)

@app.post('/api/period-control-points')
def period_control_points():
 # 対象RNEを開いて時間型管理ポイントとデータ項目を読み取る。
 # 読み取りは独立プロセスで行う。DLL側で異常終了してもアプリ本体は落ちない。
 data=request.get_json(force=True) or {}; c=load()
 if str(c['settings'].get('extract_engine') or 'api').lower()!='api':
  return jsonify(ok=False,error='この自動検出はNavigator API方式のときに使用できます。DDE方式では管理ポイント名を手入力してください'),200
 job=next((x for x in c['jobs'] if x['id']==data.get('job_id')),None) if data.get('job_id') else None
 rne_value=str(data.get('rne_path') or (job.get('rne_path') if job else '') or '').strip()
 if not rne_value:return jsonify(ok=False,error='RNEファイルを指定してください'),200
 tmp={'rne_path':rne_value,'rne':Path(rne_value).name}
 try:rp=resolve_rne_path(tmp,c)
 except Exception as e:return jsonify(ok=False,error=f'RNEパスの解決に失敗しました: {e}'),200
 if not Path(rp).is_file():return jsonify(ok=False,error=f'RNEが見つかりません: {rp}'),200
 try:
  user,pw,server,_=creds(resolve_path(c['symnavim_conf']))
  known=(load_column_cache(rp) or {}).get('columns') or []
  ins=run_inspect_worker(dict(tmp,id=(job or {}).get('id',''),name=(job or {}).get('name',''),_known_columns=known),c,user,pw,server,['points','items'])
 except Exception as e:
  log.exception('PERIOD_CP_DETECT_FAILED rne=%s',rp);return jsonify(ok=False,error=str(e)),200
 if not ins.get('ok'):
  return jsonify(ok=False,error=ins.get('error') or 'RNEを読み取れませんでした',crashed=bool(ins.get('crashed'))),200
 points=ins.get('points') or [];time_points=ins.get('time_points') or []
 items=ins.get('data_items') or [];layout=ins.get('layout') or {}
 cp_named=sum(1 for p in points if not str(p.get('name','')).startswith('管理ポイント#'))
 di_named=sum(1 for x in items if x.get('named'))
 log.info('PERIOD_CP_DETECT rne=%s total=%s time=%s named=%s',rp,len(points),len(time_points),cp_named)
 log.info('CATALOG_DATA_ITEMS rne=%s count=%s named=%s fields=%s(%s) detail=%s',
          rp,len(items),di_named,ins.get('field_count'),ins.get('field_why'),ins.get('di_diag',''))
 log.info('CATALOG_DATA_ITEM_NAMES %s',' | '.join(f"{x.get('location')}:{x.get('name')}" for x in items[:40]))
 # 直近の出力ファイルのヘッダーは、推測なしで確実に取れる列名の一覧。列分割の割り当てはこれを使う。
 out_cols=[];out_from='';out_error=''
 if job:
  try:
   op=_viewer_output_path(job,c)
   if op.is_file():
    out_cols=read_header_names(read_copy(op),job);out_from=str(op)
    if out_cols:save_column_cache(rp,out_cols,source='output',job=job)
   else:out_error='まだ出力ファイルがありません（1回実行すると列名が読めます）'
  except Exception as oe:out_error=str(oe)
 else:out_error='保存済みの対象を編集すると、直近の出力から列名を読み取れます'
 dupes=duplicate_columns(out_cols)
 if dupes:log.warning('COLUMN_DUPLICATES rne=%s count=%s names=%s',rp,len(dupes),[d['name'] for d in dupes[:10]])
 log.info('CATALOG_OUTPUT_COLUMNS file=%s count=%s duplicates=%s error=%s',out_from,len(out_cols),len(dupes),out_error)
 return jsonify(ok=True,points=points,time_points=time_points,count=len(points),time_count=len(time_points),
                data_items=items,data_item_count=len(items),data_item_named=di_named,data_item_error='',
                data_item_detail=ins.get('di_diag',''),field_count=ins.get('field_count'),control_point_named=cp_named,
                output_columns=out_cols,output_column_source=out_from,output_column_error=out_error,
                layout=layout,column_split_ready=bool(ins.get('column_split_ready')),duplicates=dupes)

@app.post('/api/column-plan')
def column_plan():
 """列分割の下調べ。RNEの列定義を用意し、削除できる列と必ず残る列に分ける。

 列名の入手は安いものから順に試す: キャッシュ → 直近の出力ファイル → 1行だけの問い合わせ。
 いずれも読み取りのみで、RNEファイルも出力ファイルも変更しない。
 """
 data=request.get_json(force=True) or {};c=load()
 job=next((x for x in c['jobs'] if x['id']==data.get('job_id')),None) if data.get('job_id') else None
 if not job:return jsonify(ok=False,error='保存済みの対象を選んでください'),200
 # 分割できるかどうかはRNEを開くだけで分かるので、問い合わせは既定では行わない。
 # 1行だけの問い合わせは、結合に必要な「出力列の並び順」がどうしても要るときだけ明示的に許可する。
 force=bool(data.get('force'));allow_probe=bool(data.get('allow_probe',False))
 try:rp=resolve_rne_path(job,c)
 except Exception as e:return jsonify(ok=False,error=f'RNEパスの解決に失敗しました: {e}'),200
 if not Path(rp).is_file():return jsonify(ok=False,error=f'RNEが見つかりません: {rp}'),200
 started=time.perf_counter();state,cached=column_cache_state(rp)
 columns=[];source='';notes=[];probe_info={}
 if cached and not cached['stale'] and not force:
  columns=cached['columns'];source='cache'
 else:
  if cached and cached['stale']:notes.append('RNEが更新されていたため取り直しました')
  try:
   op=_viewer_output_path(job,c)
   if op.is_file():
    columns=read_header_names(read_copy(op),job)
    if columns:source='output';notes.append(f'直近の出力ファイルから読み取りました（{op.name}）')
  except Exception as e:notes.append(f'出力ファイルから読めませんでした: {e}')
 classify=[];classify_error='';layout={}
 need_classify=force or not (cached and cached.get('classify') and not cached['stale'])
 need_probe=not columns and allow_probe
 try:
  if need_probe or need_classify:
   # 列挙は独立プロセスで行う。DLL側の異常終了でアプリごと落ちないようにするため。
   user,pw,server,_=creds(resolve_path(c['symnavim_conf']))
   ins=run_inspect_worker(dict(job,_known_columns=columns),c,user,pw,server,['items'])
   if not ins.get('ok'):
    return jsonify(ok=False,error=ins.get('error') or 'RNEを読み取れませんでした',crashed=bool(ins.get('crashed'))),200
   layout=ins.get('layout') or {}
   log.info('COLUMN_LAYOUT rne=%s removable=%s fixed=%s condition=%s data_items=%s control_points=%s detail=%s',
            rp,len(layout.get('removable') or []),len(layout.get('fixed') or []),len(layout.get('condition') or []),
            layout.get('data_item_count'),layout.get('control_point_count'),ins.get('di_diag',''))
   log.info('COLUMN_LAYOUT_NAMES removable=%s',' | '.join((layout.get('removable') or [])[:40]))
   # 分類は列挙結果から作れる。出力の並び順にある列が、データ欄の列かどうかを見るだけ。
   if columns:
    rem=set(layout.get('removable') or [])
    classify=[{'name':n,'removable':n in rem,'locate':'データ' if n in rem else '表側など'} for n in columns]
 except Exception as e:
  log.exception('COLUMN_PLAN_FAILED rne=%s',rp)
  return jsonify(ok=False,error=str(e)),200
 if columns and source!='cache':save_column_cache(rp,columns,source=source or 'unknown',job=job)
 if classify:save_column_classification(rp,classify)
 if layout.get('condition') is not None and columns:save_condition_items(rp,layout.get('condition') or [])
 elif cached and cached.get('classify') and not cached['stale']:classify=cached['classify']
 # 数え上げが成功していればそれを使う。RNEだけで分かる確かな値なので、名前の突き合わせより優先する。
 if layout.get('removable') or layout.get('fixed'):
  removable=layout['removable'];fixed=layout['fixed'];basis='layout'
 else:
  removable=[x['name'] for x in classify if x.get('removable')]
  fixed=[x['name'] for x in classify if not x.get('removable')];basis='classify'
 elapsed=time.perf_counter()-started
 timing=load_rne_timing(rp);incompatible=split_incompatible(rp)
 # データ量の内訳（固定列と分割できる列）が、効き目の上限を決める。接続なしで測れる。
 pw=None;panchors=[];pcov=0.0
 try:
  op=_viewer_output_path(job,c)
  if op.is_file() and columns:
   pw=column_weights(read_copy(op),job,columns)
   if pw:panchors,pcov=pick_anchor_columns(removable,pw,int(c['settings'].get('split_anchor_limit',3) or 3))
 except Exception as pe:
  log.warning('COLUMN_PLAN_WEIGHTS_FAILED rne=%s error=%s',rp,pe)
 dupes=duplicate_columns(columns)
 if dupes:
  log.warning('COLUMN_DUPLICATES rne=%s count=%s names=%s',rp,len(dupes),[d['name'] for d in dupes[:10]])
 payload=split_payload_profile(columns,removable,pw)
 trials=load_split_trials(rp)
 link=split_link_profile(rp)
 if not link.get('samples'):link=split_link_profile()      # このRNEの実測が無ければ全体の実測を使う
 best,best_gain,detail_rows=recommend_split_parts(columns,removable,int(c['settings'].get('api_parallel_max_lines',4) or 4),trials,timing,pw,c['settings'],link)
 reason=split_gain_reason(columns,removable,max(2,best),timing,pw)
 tp={'samples':link.get('samples',0),'slope':(link['headroom']-1) if link.get('headroom') else None,'points':link.get('points',[]),'link':link}
 breakeven=split_breakeven_share(tp.get('slope'))
 log.info('SPLIT_PAYLOAD rne=%s unit=%s fixed=%s(%.0f%%) splittable=%s(%.0f%%) anchors=%s coverage=%.4f',
          rp,payload['unit'],payload['fixed'],payload['fixed_share']*100,payload['splittable'],payload['splittable_share']*100,
          panchors or '(なし)',pcov)
 if incompatible or dupes:best,best_gain=1,1.0
 # 「次に実行したらどうなるか」。この対象を単独で実行したときの持ち分で判定して見せる。
 solo_budget=max(1,int(c['settings'].get('api_parallel_lines',6) or 6))
 chosen,why=plan_run_split(rp,job,c,normalize_output_format(job.get('output_format'),job.get('output_file')),solo_budget)
 runtime_split=runtime_split_view(rp,job,c,chosen,why,solo_budget)
 log.info('COLUMN_PLAN rne=%s job=%s source=%s basis=%s columns=%s removable=%s fixed=%s recommend=%s gain=%s incompatible=%s elapsed=%.2fs',
          rp,job['name'],source,basis,len(columns),len(removable),len(fixed),best,best_gain,incompatible,elapsed)
 return jsonify(ok=True,rne=str(rp),job=job['name'],source=source,cache_state=state,columns=columns,column_count=len(columns),
                classify=classify,removable=removable,fixed=fixed,removable_count=len(removable),fixed_count=len(fixed),
                basis=basis,layout=layout,classify_error=classify_error,captured_at=(cached or {}).get('captured_at',''),
                notes=notes,probe=probe_info,elapsed=round(elapsed,2),
                recommended_parts=best,predicted_gain=best_gain,gain_detail=detail_rows,gain_reason=reason,
                timing=timing,incompatible=incompatible,payload=payload,anchors=panchors,anchor_coverage=round(pcov,4),
                throughput=tp,breakeven_fixed_share=breakeven,duplicates=dupes,link=link,
                useful_parts=split_useful_parts(link),runtime_split=runtime_split)

@app.post('/api/rne-master')
def rne_master():
 """このRNEについて控えてあるもの一式と、実績の集計。

 「同じファイルなら調べ直さなくてよい」ことを画面で示すための口。
 サーバーへは一切問い合わせない（控えを読むだけ）。
 """
 data=request.get_json(silent=True) or {};c=load()
 job=next((x for x in c['jobs'] if x['id']==data.get('job_id')),None) if data.get('job_id') else None
 if not job:return jsonify(ok=False,error='保存済みの対象を選んでください'),200
 try:view=rne_master_view(job,c)
 except Exception as e:
  log.exception('RNE_MASTER_FAILED');return jsonify(ok=False,error=str(e)),200
 stats=rne_run_stats(view['rne']) if data.get('stats') else None
 log.info('RNE_MASTER job=%s state=%s 列=%s 軸=%s 割り当て=%s%s',job['name'],view['state'],
          view['columns']['count'],view['axes']['count'],len(view['plans']),
          f" 実績={stats['total']}件" if stats else '')
 return jsonify(ok=True,job=job['name'],master=view,stats=stats)

@app.post('/api/rne-master/clear-blocks')
def rne_master_clear_blocks():
 """「使わない」と記録した軸を取り消す。条件が変われば通ることもあるため、戻せるようにしておく。"""
 data=request.get_json(silent=True) or {};c=load()
 job=next((x for x in c['jobs'] if x['id']==data.get('job_id')),None) if data.get('job_id') else None
 if not job:return jsonify(ok=False,error='保存済みの対象を選んでください'),200
 rp=resolve_rne_path(job,c);n=clear_row_axis_blocks(rp)
 log.info('ROW_AXIS_BLOCK_CLEAR rne=%s removed=%s',rp,n)
 return jsonify(ok=True,removed=n,master=rne_master_view(job,c))

@app.post('/api/run-split-state')
def run_split_state():
 """次に本番で実行したら、どの形で取るのか。判定だけを返す（サーバーへは問い合わせない）。

 手順4はこれ1本で完結させる。手順2や手順3の副産物として更新されるだけだと、
 「いまどうなっているのか」を見るために別の操作が要ることになる。
 """
 data=request.get_json(silent=True) or {};c=load()
 job=next((x for x in c['jobs'] if x['id']==data.get('job_id')),None) if data.get('job_id') else None
 if not job:return jsonify(ok=False,error='保存済みの対象を選んでください'),200
 # 画面で選び直した直後でも、その設定で判定できるようにする（保存前でも見える）。
 if data.get('split_mode'):job=dict(job,split_mode=normalize_split_mode(data.get('split_mode')))
 if data.get('split_shape'):job=dict(job,split_shape=normalize_split_shape(data.get('split_shape')))
 try:rp=resolve_rne_path(job,c)
 except Exception as e:return jsonify(ok=False,error=f'RNEパスの解決に失敗しました: {e}'),200
 budget=max(1,int(c['settings'].get('api_parallel_lines',6) or 6))
 fmt=normalize_output_format(job.get('output_format'),job.get('output_file'))
 try:chosen,why=plan_run_split(rp,job,c,fmt,budget)
 except Exception as e:
  log.exception('RUN_SPLIT_STATE_FAILED rne=%s',rp);chosen,why=None,f'判定に失敗しました: {e}'
 view=runtime_split_view(rp,job,c,chosen,why,budget)
 log.info('RUN_SPLIT_STATE job=%s 形=%s 動作=%s 指定=%s → %s（%s）',job['name'],view['shape_used'] or '-',
          view['mode'],view['shape'],view['how'] or '分割しない',why or '裏付けあり')
 return jsonify(ok=True,rne=str(rp),job=job['name'],runtime_split=view)

@app.post('/api/row-split-plan')
def row_split_plan():
 """行分割の下調べ。分割できる列を探し、実行と転送の内訳を測る。

 出力ファイルには一切触れない。サーバーへ行くのは内訳の測定だけで、それも転送は発生させない
 （NAVI_DOWNLOADLATER で問い合わせだけ実行し、受信を始めずに降りる）。
 """
 data=request.get_json(force=True) or {};c=load()
 job=next((x for x in c['jobs'] if x['id']==data.get('job_id')),None) if data.get('job_id') else None
 if not job:return jsonify(ok=False,error='保存済みの対象を選んでください'),200
 try:rp=resolve_rne_path(job,c)
 except Exception as e:return jsonify(ok=False,error=f'RNEパスの解決に失敗しました: {e}'),200
 if not Path(rp).is_file():return jsonify(ok=False,error=f'RNEが見つかりません: {rp}'),200
 parts=max(2,min(8,int(data.get('parts') or 2)))
 started=time.perf_counter()
 # 行を絞れるのは管理ポイントだけ。まずRNEから軸を読む。ここは直近の出力に依存しないので、
 # どのRNEでも同じように調べられる。
 axes=[];axes_error='';axis=None;ranked=[]
 try:
  user,pw,server,_=creds(resolve_path(c['symnavim_conf']))
  ins=run_inspect_worker(dict(job,_read_names=True),c,user,pw,server,['axes'],
                         timeout=int(c['settings'].get('split_trial_timeout_seconds',1800) or 1800))
  if ins.get('ok'):axes=ins.get('axes') or []
  else:axes_error=ins.get('error') or '管理ポイントを読み取れませんでした'
 except Exception as e:
  axes_error=str(e);log.exception('ROW_AXES_FAILED rne=%s',rp)
 parts0=max(2,min(8,int(data.get('parts') or 2)))
 choice=row_axis_choice(data,job)
 scores=axis_balance_scores(job,c,[a.get('name') for a in axes])
 axis_blocks=blocked_row_axes(rp)
 axis,ranked,axis_why=pick_row_axis_by_mode(axes,parts0,choice['mode'],choice['index'],choice['name'],scores,axis_blocks)
 log.info('ROW_AXES rne=%s 読めた軸=%s 使える=%s 決め方=%s 既定=%s（%s）',rp,len(axes),
          sum(1 for x in ranked if x['usable'] and x['enough']),choice['mode'],
          (axis or {}).get('name','(なし)'),axis_why)
 for x in ranked:
  sc=scores.get(x.get('name'))
  if sc:
   x['balance']={'top_share':round(sc['top_share'],4),'distinct':sc['distinct'],'blank_rows':sc['blank_rows']}
   # 値が空の行はどのカテゴリにも入らない。この軸で分けるとその行が落ちる。
   if sc['blank_rows']:
    x['usable']=False
    x['reason']=f'値が空の行が{sc["blank_rows"]}行あります。この軸で分けるとその行が結果から落ちます'
 # ここで読んだ一覧を覚えておく。影実行が同じ問い合わせを繰り返さずに済む（実行直前の読み直しは別途行う）。
 save_axis_survey(rp,axes,parts0)
 state,cached=column_cache_state(rp)
 if not cached or not cached.get('columns'):
  return jsonify(ok=True,rne=str(rp),job=job['name'],parts=parts0,axes=ranked,axis=axis,axes_error=axes_error,axis_mode=choice['mode'],axis_why=axis_why,
                 axis_modes=[{'id':k,'label':v} for k,v in ROW_AXIS_MODE_LABEL.items()],
                 columns=0,fixed=0,removable_count=0,candidates=[],examined=0,sample_rows=0,
                 candidate_error='列定義が未取得です（手順2の「列を調べる」を実行すると、列の情報も出せます）',
                 breakdown={'known':False},condition_items=[]),200
 columns=cached['columns']
 classify=cached.get('classify') or []
 removable=[x['name'] for x in classify if x.get('removable')]
 fixed=[x['name'] for x in classify if not x.get('removable')]
 # 絞り込みの条件は「条件欄」の項目に付く。出力される列（データ欄）に同じ条件を設定しても
 # rc=OK が返るだけで1行も絞られない（2026-08-10の実測）。どれが条件欄にあるのかを先に出す。
 cond_items=[str(x) for x in (cached.get('condition') or [])]
 # 1) どの列でどう割ると均等になるか。直近の出力を読むだけで、サーバーには触れない。
 cand=None;cand_error=''
 try:
  op=_viewer_output_path(job,c)
  if op.is_file():cand=row_split_candidates(read_copy(op),job,columns,removable,parts)
  else:cand_error='直近の出力ファイルがありません。1回実行すると候補を探せます'
 except Exception as e:
  cand_error=str(e);log.warning('ROW_SPLIT_CANDIDATES_FAILED rne=%s error=%s',rp,e)
 if cand:
  log.info('ROW_SPLIT_CANDIDATES rne=%s 調べた列=%s 候補=%s 最良=%s',rp,cand['examined'],len(cand['candidates']),
           (cand['candidates'][0]['column']+f" 偏り{cand['candidates'][0]['balance']}") if cand['candidates'] else '(なし)')
 # 2) 所要時間の内訳。ここだけサーバーへ行く（転送はしない）。
 timing=load_rne_timing(rp);deferred=None;probe_error=''
 if data.get('probe'):
  try:
   user,pw,server,_=creds(resolve_path(c['symnavim_conf']))
   ins=run_inspect_worker(dict(job,_read_names=False),c,user,pw,server,['timing'],
                          timeout=int(c['settings'].get('split_trial_timeout_seconds',1800) or 1800))
   if ins.get('ok') and ins.get('deferred'):deferred=ins['deferred']
   else:probe_error=ins.get('error') or '内訳を測れませんでした'
  except Exception as e:
   probe_error=str(e);log.exception('ROW_SPLIT_PROBE_FAILED rne=%s',rp)
 # 前回の実行で運んだバイト数を足す。速度（KB/s）を出すのに要る。
 lastm=(load_job_runs().get(job['id']) or {}).get('metrics') or {}
 if timing and lastm.get('transfer_bytes') and not lastm.get('split_parts'):
  timing=dict(timing,bytes=int(lastm['transfer_bytes']))
 breakdown=row_split_breakdown(timing,deferred)
 if deferred:
  # 転送なしの実行は行数を返さない（受信して初めて分かる）。行数は前回の実行の実績を使う。
  log.info('ROW_SPLIT_PROBE rne=%s 転送なしの実行=%.2fs（返り行数=%s ※受信前なので0のことがある） '
           '通常の実行=%ss 保存=%ss → サーバー側=%ss 受信=%ss(%s KB/s) 整形=%ss(%s KB/s) 行数=%s',
           rp,deferred['execute_seconds'],deferred.get('rows'),
           (timing or {}).get('execute'),(timing or {}).get('save'),
           breakdown.get('server_seconds'),breakdown.get('download_seconds'),breakdown.get('download_kbs'),
           breakdown.get('format_seconds'),breakdown.get('format_kbs'),breakdown.get('rows'))
 return jsonify(ok=True,rne=str(rp),job=job['name'],parts=parts,columns=len(columns),fixed=len(fixed),
                removable_count=len(removable),
                axes=ranked,axis=axis,axes_error=axes_error,axis_mode=choice['mode'],axis_why=axis_why,
                 axis_modes=[{'id':k,'label':v} for k,v in ROW_AXIS_MODE_LABEL.items()],
                condition_items=cond_items,
                candidates=[dict(x,in_condition=(x['column'] in set(cond_items))) for x in (cand or {}).get('candidates',[])],
                examined=(cand or {}).get('examined',0),
                sample_rows=(cand or {}).get('rows',0),candidate_error=cand_error,
                timing=timing,deferred=deferred,probe_error=probe_error,breakdown=breakdown,
                link=split_link_profile(rp),elapsed=round(time.perf_counter()-started,2))

# 影実行を1本だけ始める受け口（POST /api/column-split-trial）は v1.73.0 で無くした。
# 測るものを選ぶ場所を1枚にまとめたので、1本だけ測るのは「1つだけ選んで
# /api/split-trial-batch」と同じことになる。入口が2つあると、片方で走らせた結果が
# 順位表に載らない・基準が別々に取られる、という食い違いが起きる。
# 進み具合を見る status は残す。測定の「いま走っている1本」の中身は、
# これまでどおりここから読む。
@app.get('/api/column-split-trial/status')
def column_split_trial_status():
 with split_trial_lock:
  st=dict(split_trial_state)
 if st.get('running'):st['elapsed']=round(time.time()-(st.get('started') or time.time()),1)
 return jsonify(ok=True,**st)

@app.get('/api/split-trial/options')
def split_trial_options_api():
 """この対象で、いま何が測れるのか。画面はこれを見て、選べないものを理由つきで塞ぐ。

 また、使える軸は行の片数で変わる（片数が多いほど条件式が長くなる）。画面が選んだ
 片数を row_parts で受け取り、その片数で通る軸だけを「選べる」と返す。
 """
 c=load();job=next((x for x in c['jobs'] if x['id']==request.args.get('job_id')),None) if request.args.get('job_id') else None
 if not job:return jsonify(ok=False,error='保存済みの対象を選んでください'),200
 try:return jsonify(ok=True,**split_trial_options(job,c,request.args.get('row_parts')))
 except Exception as e:
  log.exception('SPLIT_TRIAL_OPTIONS_FAILED job=%s',job.get('name'));return jsonify(ok=False,error=str(e)),200

@app.post('/api/split-trial-batch')
def split_trial_batch_start():
 data=request.get_json(force=True) or {};c=load()
 job=next((x for x in c['jobs'] if x['id']==data.get('job_id')),None) if data.get('job_id') else None
 if not job:return jsonify(ok=False,error='保存済みの対象を選んでください'),200
 if str(c['settings'].get('extract_engine') or 'api').lower()!='api':
  return jsonify(ok=False,error='速さを測るのはNavigator API方式のときに使用できます'),200
 kind='axes' if str(data.get('kind') or 'methods')=='axes' else 'methods'
 try:items=split_batch_items(kind,data,job)
 except Exception as e:return jsonify(ok=False,error=str(e)),200
 measured=[x for x in items if x['shape']!='normal']
 if not measured:
  return jsonify(ok=False,error='測る対象がありません。分け方か軸を1つ以上選んでください'),200
 # 上限を超えたら黙って切り落とさない。画面が「9本測ります」と言ったのに8本しか
 # 走らないと、出てこなかった条件を利用者が探すことになる。
 if len(measured)>SPLIT_BATCH_MAX:
  return jsonify(ok=False,error=f'一度に測れるのは{SPLIT_BATCH_MAX}本までです'
                              f'（{len(measured)}本を指定されました）。軸か分け方を減らしてください'),200
 with split_batch_lock,split_trial_lock:
  if split_batch_state.get('running'):
   return jsonify(ok=False,error=f'速さの測定が進行中です（{split_batch_state.get("job")}）。終わるまでお待ちください',busy=True),200
  if split_trial_state.get('running'):
   return jsonify(ok=False,error=f'影実行が進行中です（{split_trial_state.get("job")}）。終わるまでお待ちください',busy=True),200
  split_batch_state.clear()
  split_batch_state.update(split_batch_blank(running=True,kind=kind,job=job['name'],
    job_id=str(job.get('id') or ''),rne=str(job.get('rne') or ''),started=time.time(),
    total=len(items),items=items))
 def worker():
  t0=time.perf_counter()
  for i,item in enumerate(items):
   with split_batch_lock:
    if split_batch_state.get('stop'):
     for rest in items[i:]:rest['state']='中止'
     break
    item['state']='実行中';split_batch_state.update(index=i+1,items=items)
   with split_trial_lock:
    split_trial_state.clear()
    split_trial_state.update(split_trial_blank(running=True,stage='準備中',job=job['name'],
      job_id=str(job.get('id') or ''),rne=str(job.get('rne') or ''),started=time.time()))
    _split_stage_logged.update(text='',at=0.0)
   log.info('SPLIT_BATCH_ITEM %s/%s job=%s %s',i+1,len(items),job.get('name'),item['label'])
   try:res=_split_trial_run(split_batch_run_data(item,data),c,job)
   except Exception as e:
    log.exception('SPLIT_BATCH_ITEM_FAILED %s',item['label']);res={'ok':False,'error':str(e)}
   with split_trial_lock:
    split_trial_state.update(running=False,stage='完了',result=res,
      elapsed=round(time.time()-(split_trial_state.get('started') or time.time()),1))
   with split_batch_lock:
    if res.get('ok'):
     # 経過時間と倍率は、必ず対応する組で持たせる。
     #   speedup   … 分割なし ÷ この経過時間（画面の秒数と割り算が合う）
     #   run_speedup … 本番で毎回かかる「軸の読み直し」も足した実力。自動はこちらで決める
     # 以前は「見かけの秒数」と「実力の倍率」を並べていたため、55.3秒と0.76倍が並び、
     # 76.6÷55.3=1.39 と暗算しても画面の数字にならなかった（利用者からの指摘）。
     item.update(state='完了',
       elapsed=res.get('normal_elapsed') if item['shape']=='normal' else res.get('split_elapsed'),
       speedup=res.get('speedup'),
       run_speedup=res.get('run_speedup'),axis_seconds=res.get('axis_seconds'),
       identical=True if item['shape']=='normal' else bool(res.get('identical')),
       rows=res.get('rows'),detail=res.get('how') or '',
       # 軸を名指ししないで測ったときは、サーバーが選んだ軸をここで持ち帰る。
       # 持ち帰らないと「自動で測ったが、何の軸だったか分からない」ままになり、
       # 「本番で使う」を押しても別の軸が選ばれうる（利用者からの指摘）。
       axis_used=res.get('row_column') or '',
       stale_baseline=bool(res.get('stale_baseline')),
       # 「一致しません」は分け方のせいだと読める。基準が古いだけのときはそう言う。
       error=('' if (item['shape']=='normal' or res.get('identical'))
              else ('基準が古いため比べられません（基準を測り直してください）' if res.get('stale_baseline')
                    else '結果が一致しませんでした')))
    else:
     item.update(state='失敗',error=str(res.get('error') or '失敗しました'))
    split_batch_state.update(items=items)
   log.info('SPLIT_BATCH_ITEM_END %s/%s %s → %s %s',i+1,len(items),item['label'],item['state'],
            f"{item.get('elapsed')}s" if item.get('elapsed') else item.get('error',''))
  el=time.perf_counter()-t0
  with split_batch_lock:
   split_batch_state.update(running=False,elapsed=round(el,1),index=len(items),
     finished_at=datetime.now().isoformat(timespec='seconds'),items=items)
  s=split_batch_summary(items)
  log.info('SPLIT_BATCH_END job=%s 件数=%s 所要=%.1fs 最速=%s',job.get('name'),len(items),el,
           (s.get('best') or {}).get('label') or '（なし）')
 threading.Thread(target=worker,daemon=True,name='split-batch').start()
 log.info('SPLIT_BATCH_START job=%s kind=%s 件数=%s 内訳=%s',job.get('name'),kind,len(items),
          ' / '.join(x['label'] for x in items))
 return jsonify(ok=True,started=True,job=job['name'],total=len(items),
                items=[{'key':x['key'],'label':x['label'],'why':x['why']} for x in items])

@app.get('/api/split-trial-batch/status')
def split_trial_batch_status():
 with split_batch_lock:
  st=dict(split_batch_state);st['items']=[dict(x) for x in (st.get('items') or [])]
 if st.get('running'):st['elapsed']=round(time.time()-(st.get('started') or time.time()),1)
 st['summary']=split_batch_summary(st.get('items') or [])
 return jsonify(ok=True,**st)

@app.post('/api/split-trial-batch/stop')
def split_trial_batch_stop():
 """走っている1件は最後まで測る。途中で切ると、その1件が測れていないのか
 遅いのかが分からなくなるため。次の1件へ進む前に止める。"""
 with split_batch_lock:
  if not split_batch_state.get('running'):return jsonify(ok=False,error='速さの測定は動いていません'),200
  split_batch_state['stop']=True
 log.info('SPLIT_BATCH_STOP_REQUEST job=%s',split_batch_state.get('job'))
 return jsonify(ok=True,stopping=True)

@app.post('/api/inspect-task/all')
def start_inspect_all():
 """RNEを1回で調べ切る。中身・列・行を続けて読み、この RNE の控えとして保存する。

 3つに分かれていたのは実装の都合で、利用者から見れば「そのRNEを調べる」1つの用事。
 同じRNEを何度も開き直すことにもなっていた。
 """
 data=request.get_json(force=True) or {}
 title,detail,expected=INSPECT_ALL_SPEC
 with inspect_task_lock:
  if inspect_tasks['all'].get('running'):
   return jsonify(ok=False,error='「RNEを調査」はすでに実行中です。終わるまでお待ちください',busy=True),200
  jid,jname,jrne=task_target(data)
  inspect_tasks['all']=inspect_task_blank(
    kind='all',running=True,title=title,stage=detail,started=time.time(),
    job=jname,job_id=jid,rne=jrne,steps=[],
    expected=float(inspect_task_seconds.get('all') or expected),
    measured='all' in inspect_task_seconds)
 def worker():
  t=time.perf_counter();steps=[];parts={};err=''
  for n,kind in enumerate(INSPECT_ALL_ORDER,1):
   _e,_p,ktitle,kdetail,_s=INSPECT_TASK_SPECS[kind]
   with inspect_task_lock:
    inspect_tasks['all'].update(stage=f'{n}/{len(INSPECT_ALL_ORDER)} {ktitle}: {kdetail}',steps=list(steps))
   kt=time.perf_counter()
   try:out=_run_inspect_endpoint(kind,data)
   except Exception as e:
    log.exception('INSPECT_ALL_STEP_FAILED kind=%s',kind);out={'ok':False,'error':str(e)}
   el=time.perf_counter()-kt
   inspect_task_seconds[kind]=round(el,2)
   parts[kind]=out
   steps.append({'kind':kind,'title':ktitle,'ok':bool(out.get('ok')),
                 'error':'' if out.get('ok') else str(out.get('error') or ''),'elapsed':round(el,1)})
   log.info('INSPECT_ALL_STEP %s/%s kind=%s title=%s ok=%s elapsed=%.2fs',
            n,len(INSPECT_ALL_ORDER),kind,ktitle,out.get('ok'),el)
   # 途中で1つ落ちても残りは続ける。列が読めなくても行は調べられることがある。
   if not out.get('ok') and not err:err=f'{ktitle}: {out.get("error") or "失敗しました"}'
   with inspect_task_lock:inspect_tasks['all'].update(steps=list(steps))
  el=time.perf_counter()-t
  inspect_task_seconds['all']=round(el,2)
  ok_all=all(x['ok'] for x in steps)
  out={'ok':ok_all,'error':'' if ok_all else err,'steps':steps,
       'read':parts.get('read'),'column':parts.get('column'),'row':parts.get('row')}
  # 何が控えとして残ったかを、そのまま画面へ返せる形にする。
  try:
   job=next((x for x in load()['jobs'] if x['id']==data.get('job_id')),None)
   if job:out['master']=rne_master_view(job,load())
  except Exception:log.exception('INSPECT_ALL_MASTER_FAILED')
  with inspect_task_lock:
   inspect_tasks['all'].update(running=False,stage='完了',percent=100.0,result=out,
                               elapsed=round(el,1),error=out['error'],steps=steps)
  log.info('INSPECT_ALL job=%s ok=%s elapsed=%.2fs 内訳=%s',data.get('job_name'),ok_all,el,
           ' / '.join(f"{x['title']}={'OK' if x['ok'] else 'NG'}({x['elapsed']}s)" for x in steps))
 threading.Thread(target=worker,daemon=True,name='inspect-all').start()
 log.info('INSPECT_ALL_START job=%s 見込み=%.0f秒%s',data.get('job_name'),
          inspect_tasks['all']['expected'],'' if inspect_tasks['all']['measured'] else '（まだ実測がないので目安）')
 return jsonify(ok=True,started=True,title=title,stage=detail)

@app.post('/api/inspect-task/<kind>')
def start_inspect_task(kind):
 spec=INSPECT_TASK_SPECS.get(str(kind))
 if not spec:return jsonify(ok=False,error=f'知らない調べもの: {kind}'),200
 endpoint,path,title,detail,default_seconds=spec
 data=request.get_json(force=True) or {}
 with inspect_task_lock:
  if inspect_tasks[kind].get('running'):
   return jsonify(ok=False,error=f'「{title}」はすでに実行中です。終わるまでお待ちください',busy=True),200
  jid,jname,jrne=task_target(data)
  inspect_tasks[kind]=inspect_task_blank(
    kind=kind,running=True,title=title,stage=detail,started=time.time(),
    job=jname,job_id=jid,rne=jrne,
    expected=float(inspect_task_seconds.get(kind) or default_seconds),
    measured=kind in inspect_task_seconds)
 view=app.view_functions.get(endpoint)
 def worker():
  t=time.perf_counter();out={'ok':False,'error':'応答がありませんでした'}
  try:
   if view is None:raise RuntimeError(f'{endpoint} が登録されていません')
   with app.test_request_context(path,method='POST',json=data):
    rv=view()
   resp=rv[0] if isinstance(rv,tuple) else rv
   out=resp.get_json(silent=True) or {}
  except Exception as e:
   log.exception('INSPECT_TASK_FAILED kind=%s',kind);out={'ok':False,'error':str(e)}
  el=time.perf_counter()-t
  inspect_task_seconds[kind]=round(el,2)
  with inspect_task_lock:
   inspect_tasks[kind].update(running=False,stage='完了',percent=100.0,result=out,
                              elapsed=round(el,1),error=('' if out.get('ok') else str(out.get('error') or '')))
  log.info('INSPECT_TASK kind=%s title=%s ok=%s elapsed=%.2fs',kind,title,out.get('ok'),el)
 threading.Thread(target=worker,daemon=True,name=f'inspect-{kind}').start()
 log.info('INSPECT_TASK_START kind=%s title=%s 見込み=%.0f秒%s',kind,title,
          inspect_tasks[kind]['expected'],'' if inspect_tasks[kind]['measured'] else '（まだ実測がないので目安）')
 return jsonify(ok=True,started=True,title=title,stage=detail)

@app.get('/api/inspect-task/<kind>')
def get_inspect_task(kind):
 if str(kind) not in INSPECT_TASK_SPECS and str(kind)!='all':
  return jsonify(ok=False,error=f'知らない調べもの: {kind}'),200
 with inspect_task_lock:
  st=dict(inspect_tasks[str(kind)])
 if st.get('running'):
  st['elapsed']=round(time.time()-(st.get('started') or time.time()),1)
  st['percent']=inspect_task_percent(st)
 return jsonify(ok=True,**st)

@app.get('/api/background-tasks')
def background_tasks():
 """いま裏で走っているものを1か所で答える。

 調べものも影実行も、画面を閉じても続く。どの画面にいても「何がどのRNEで走って
 いるか」が分かるようにしないと、終わったのかどうかを確かめる方法が無くなる。
 """
 out=[]
 with inspect_task_lock:
  for kind,st in inspect_tasks.items():
   if not st.get('running'):continue
   out.append({'type':'inspect','kind':kind,'title':st.get('title') or '調べもの',
               'stage':st.get('stage') or '','job':st.get('job') or '','job_id':st.get('job_id') or '',
               'rne':st.get('rne') or '','percent':inspect_task_percent(st),
               'elapsed':round(time.time()-(st.get('started') or time.time()),1)})
 with split_batch_lock:
  bt=dict(split_batch_state)
 if bt.get('running'):
  cur=next((x for x in (bt.get('items') or []) if x.get('state')=='実行中'),{})
  out.append({'type':'batch','kind':'batch','title':'速さを測る',
              'stage':f"{bt.get('index')}/{bt.get('total')} {cur.get('label') or '準備中'}",
              'job':bt.get('job') or '','job_id':bt.get('job_id') or '','rne':bt.get('rne') or '',
              'percent':round(max(0,(int(bt.get('index') or 1)-1))/max(1,int(bt.get('total') or 1))*100,1),
              'elapsed':round(time.time()-(bt.get('started') or time.time()),1)})
 with split_trial_lock:
  st=dict(split_trial_state)
 if st.get('running') and not bt.get('running'):
  out.append({'type':'trial','kind':'trial','title':'速さを測る（影実行）',
              'stage':st.get('stage') or '','job':st.get('job') or '','job_id':st.get('job_id') or '',
              'rne':st.get('rne') or '','percent':float(st.get('percent') or 0),
              'elapsed':round(time.time()-(st.get('started') or time.time()),1)})
 return jsonify(ok=True,count=len(out),tasks=out)

@app.post('/api/run')
def run_all():
 try:
  data=request.get_json(silent=True) or {};requested=clamp_parallel_lines(data.get('parallel_lines',1))
  item,position=enqueue_command(data.get('job_ids'),'manual',requested)
  return jsonify(ok=True,queued=True,queue_id=item['id'],position=position,parallel_lines=requested,job_names=item['job_names'])
 except Exception as e:return jsonify(error=str(e)),400

@app.post('/api/run/<job_id>')
def run_one(job_id):
 try:
  item,position=enqueue_command([job_id],'manual-single',1)
  return jsonify(ok=True,queued=True,queue_id=item['id'],position=position,parallel_lines=1,job_names=item['job_names'])
 except Exception as e:return jsonify(error=str(e)),400

@app.post('/api/run/cancel')
def cancel_run():
 with command_queue_lock:
  queue_cleared=len(command_queue); command_queue.clear()
 if not status.get('running'):
  log.info('CANCEL_REQUESTED_IDLE queue_cleared=%s',queue_cleared)
  return jsonify(ok=True,was_running=False,workers_terminated=0,queue_cleared=queue_cleared)
 cancel_requested.set()
 with active_workers_lock:workers=list(active_workers.values())
 for p in workers:
  try:p.terminate()
  except Exception:pass
 log.info('CANCEL_REQUESTED workers_terminated=%s queue_cleared=%s',len(workers),queue_cleared)
 return jsonify(ok=True,was_running=True,workers_terminated=len(workers),queue_cleared=queue_cleared)

@app.get('/api/execution-queue')
def get_execution_queue():
 response=jsonify(queue_snapshot());response.headers['Cache-Control']='no-store';return response

@app.delete('/api/execution-queue/<queue_id>')
def delete_execution_queue(queue_id):
 with command_queue_lock:
  for i,item in enumerate(command_queue):
   if item['id']==queue_id:
    removed=command_queue.pop(i);log.info('COMMAND_QUEUE_DELETE id=%s jobs=%s',queue_id,removed['job_names']);return jsonify(ok=True)
 return jsonify(error='実行中または対象が見つかりません'),409

@app.post('/api/execution-queue/<queue_id>/move')
def move_execution_queue(queue_id):
 data=request.get_json(silent=True) or {};direction=data.get('direction')
 with command_queue_lock:
  idx=next((i for i,x in enumerate(command_queue) if x['id']==queue_id),None)
  if idx is None:return jsonify(error='待機キューにありません'),404
  dest=idx-1 if direction=='up' else idx+1
  if dest<0 or dest>=len(command_queue):return jsonify(ok=True)
  command_queue[idx],command_queue[dest]=command_queue[dest],command_queue[idx]
  log.info('COMMAND_QUEUE_MOVE id=%s direction=%s new_position=%s',queue_id,direction,dest+1)
 return jsonify(ok=True)

@app.get('/api/status')
def get_status():
 response=jsonify(status);response.headers['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0';return response

@app.get('/api/navigator-api-status')
def navigator_api_status():
 c=load();force=request.args.get('full')=='1'
 if not force:
  cached=_read_api_diag_cache(c)
  if cached:
   log.info('API_DIAG cache_hit=1 ok=%s age=%ss dll=%s dll_bits=%s',cached.get('ok'),cached.get('cache_age_seconds','-'),cached.get('dll'),cached.get('dll_bits'));_log_api_exports(cached.get('dll'),cached.get('exports'),cached.get('exports_bound'))
   cached=dict(cached);cached['cached']=True;cached['requirement']=_dll_requirement(c);cached['search_roots']=[str(x) for x in dll_search_roots(c)];cached['issues']=dll_diagnostic_issues(cached.get('attempts') or [],cached.get('python_bits'),cached.get('exports'),cached.get('exports_bound'),cached['requirement']);cached['readiness']=api_readiness(cached);return jsonify(cached)
 try:
  from navigator_api import NavigatorApi
  started=time.perf_counter();api=NavigatorApi(resolve_path(c.get('symnavi_exe','')),log,resolve_path(c.get('navigator_api_dll')) if c.get('navigator_api_dll') else None,base_dir=BASE,search_roots=dll_search_roots(c));info=api.info();api.close();info['elapsed']=round(time.perf_counter()-started,3);info['cached']=False;info['issues']=dll_diagnostic_issues(info.get('attempts') or [],info.get('python_bits'),info.get('exports'),info.get('exports_bound'),info.get('requirement'));info['readiness']=api_readiness(info);log.info('API_DIAG cache_hit=0 elapsed=%.3fs dll=%s dll_bits=%s attempts=%s selection=%s',info.get('elapsed'),info.get('dll'),info.get('dll_bits'),len(info.get('attempts') or []),info.get('selection_reason'));_log_api_exports(info.get('dll'),info.get('exports'),info.get('exports_bound'));_write_api_diag_cache(info,c);return jsonify(info)
 except Exception as e:
  exports=[];exports_dll=''
  try:
   from navigator_api import candidate_dlls,pe_bits,pe_exports
   pybits=struct.calcsize('P')*8;attempts=[]
   for p in candidate_dlls(resolve_path(c.get('symnavi_exe','')),resolve_path(c.get('navigator_api_dll')) if c.get('navigator_api_dll') else None,[BASE/'Config'/'NAVIAP',BASE/'NAVIAP'],search_roots=dll_search_roots(c)):
    exists=p.is_file();bits=pe_bits(p) if exists else None;attempts.append({'path':str(p),'exists':exists,'dll_bits':bits,'python_bits':pybits,'result':'bit_mismatch' if exists and bits and bits!=pybits else 'not_found'})
    # DLLを読み込めなくても、エクスポート表はファイルを読むだけで分かる。
    # 読み込みに失敗する端末こそ、そのDLLに何ができるのかを知りたい。
    if exists and not exports:
     found=[x for x in pe_exports(p) if x.lower().startswith('navi')]
     if found:exports=found;exports_dll=str(p)
  except Exception:pybits=None;attempts=[]
  log.warning('API_DIAG cache_hit=0 result=failed attempts=%s error=%s',len(attempts),e)
  _log_api_exports(exports_dll,exports,[])
  req=_dll_requirement(c)
  payload={'ok':False,'error':str(e),'mode':'Navigator API','cached':False,'attempts':attempts,
           'exports':exports,'exports_bound':[],'requirement':req,'python_bits':pybits,
           'search_roots':req.get('search_roots') or [],
           'issues':dll_diagnostic_issues(attempts,pybits,exports,[],req)}
  payload['readiness']=api_readiness(payload)
  # 失敗も短い間だけ覚えておく。探し方を変えれば無効になるので、直したのに
  # 古い結果が出続けることはない。
  _write_api_diag_cache(payload,c)
  return jsonify(payload),200

@app.get('/api/navigator-api/requirement')
def navigator_api_requirement():
 """「どのDLLを持ってくればよいか」だけを返す。診断を走らせる前に読めるようにする。"""
 c=load();req=_dll_requirement(c)
 return jsonify(ok=bool(req),requirement=req,configured=c.get('navigator_api_dll',''),
                search_roots=[str(x) for x in dll_search_roots(c)],
                default_roots=list(DEFAULT_DLL_SEARCH_ROOTS),
                depth=int(c.get('navigator_api_search_depth',3) or 3))

@app.post('/api/navigator-api/search-roots')
def navigator_api_search_roots():
 """DLLを探す範囲を保存する。別のPCではNAVIAPの置き場所が変わるため、範囲そのものを設定にしている。"""
 body=request.get_json(silent=True) or {}
 roots=[str(x).strip() for x in (body.get('roots') or []) if str(x or '').strip()]
 try:depth=max(1,min(6,int(body.get('depth',3) or 3)))
 except (TypeError,ValueError):depth=3
 c=load()
 c['navigator_api_search_roots']=roots or list(DEFAULT_DLL_SEARCH_ROOTS)
 c['navigator_api_search_depth']=depth
 c.pop('credential_status',None);save(c)
 log.info('DLL_SEARCH_ROOTS_SAVED count=%s depth=%s roots=%s',len(c['navigator_api_search_roots']),depth,' | '.join(c['navigator_api_search_roots']))
 return jsonify(ok=True,roots=c['navigator_api_search_roots'],resolved=[str(x) for x in dll_search_roots(c)],depth=depth)

@app.post('/api/navigator-api/scan')
def navigator_api_scan():
 """指定した範囲を実際に歩いてSymNaviA.dllを集め、そのまま使えるものを先頭に返す。

 名前ではなくファイルで探すので、配布フォルダーの名前と違う場所に置かれていても見つかる。
 """
 body=request.get_json(silent=True) or {}
 c=load()
 raw=[str(x).strip() for x in (body.get('roots') or []) if str(x or '').strip()]
 roots=[resolve_path(x) for x in raw] if raw else dll_search_roots(c)
 try:depth=max(1,min(6,int(body.get('depth') or c.get('navigator_api_search_depth',3) or 3)))
 except (TypeError,ValueError):depth=3
 started=time.perf_counter()
 try:
  from navigator_api import scan_dll_roots
  result=scan_dll_roots(roots,max_depth=depth)
 except Exception as e:
  log.exception('DLL_SCAN_FAILED');return jsonify(ok=False,error=str(e)),200
 result['elapsed']=round(time.perf_counter()-started,3)
 result['requirement']=_dll_requirement(c)
 result['configured']=c.get('navigator_api_dll','')
 result['recommended']=(result['usable'][0]['path'] if result.get('usable') else '')
 log.info('DLL_SCAN roots=%s depth=%s found=%s usable=%s elapsed=%.2fs recommended=%s',
          len(result.get('roots') or []),depth,len(result.get('found') or []),len(result.get('usable') or []),result['elapsed'],result['recommended'] or '-')
 for x in (result.get('found') or []):
  log.info('DLL_FOUND path=%s dll_bits=%s python_bits=%s usable=%s folder=%s',x.get('path'),x.get('dll_bits'),x.get('python_bits'),int(bool(x.get('usable'))),x.get('folder'))
 return jsonify(ok=True,**result)

@app.post('/api/navigator-api/select')
def navigator_api_select():
 """検索結果や参照ダイアログで選んだDLLを、手動指定として保存する。"""
 body=request.get_json(silent=True) or {}
 path=str(body.get('path') or '').strip()
 if not path:return jsonify(ok=False,error='DLLのパスが指定されていません'),400
 resolved=resolve_path(path)
 if resolved.is_dir():resolved=resolved/'SymNaviA.dll'
 if not resolved.is_file():return jsonify(ok=False,error=f'指定されたファイルがありません: {resolved}'),200
 try:
  from navigator_api import pe_bits
  bits=pe_bits(resolved)
 except Exception:bits=None
 pybits=struct.calcsize('P')*8
 c=load();c['navigator_api_dll']=str(resolved);c.pop('credential_status',None);save(c)
 try:_api_diag_cache_path().unlink()
 except OSError:pass
 log.info('DLL_SELECTED path=%s dll_bits=%s python_bits=%s match=%s',resolved,bits,pybits,int(bits==pybits))
 return jsonify(ok=True,path=str(resolved),dll_bits=bits,python_bits=pybits,match=bool(bits==pybits),
                warning=('' if bits==pybits else f'このDLLは{bits or "不明"}bitです。Pythonは{pybits}bitのため、このままでは読み込めません。'))

@app.get('/api/log')
def get_log():
 try:limit=int(request.args.get('limit') or 1200)
 except Exception:limit=1200
 lines,total,size,clipped=read_log_lines(limit,request.args.get('q') or '',request.args.get('filter') or '')
 older=[{'name':p.name,'mb':round(p.stat().st_size/1024/1024,1)} for p in log_files()[1:]]
 return jsonify(text='\n'.join(lines),total=total,matched=len(lines),
                size_mb=round(size/1024/1024,1),clipped=clipped,
                window_mb=round(LOG_READ_BYTES/1024/1024,1),older=older,
                filters=[{'id':k,'label':v[1]} for k,v in LOG_FILTERS.items()])

@app.post('/api/log/clear')
def clear_log():
 p=LOG_PATH; p.parent.mkdir(exist_ok=True)
 p.write_text('',encoding='utf-8')
 return jsonify(ok=True)

@app.post('/api/log/rotate')
def rotate_log_now():
 """いますぐ付け替える。実行中はワーカーが書いているのでできない。"""
 if status.get('running'):return jsonify(ok=False,error='実行中は付け替えできません（ワーカーが同じファイルへ書いています）'),409
 size=rotate_log_if_needed(force=True)
 return jsonify(ok=True,rotated=bool(size),size_mb=round((size or 0)/1024/1024,1))

@app.post('/api/log/delete-lines')
def delete_log_lines():
 data=request.get_json(silent=True) or {}; remove=set(data.get('lines') or [])
 if status.get('running'):return jsonify(ok=False,error='実行中は削除できません（ワーカーが同じファイルへ書いています）'),409
 p=LOG_PATH
 if not p.exists():return jsonify(ok=True,removed=0)
 lines=p.read_text(encoding='utf-8',errors='replace').splitlines()
 kept=[x for x in lines if x not in remove]
 p.write_text('\n'.join(kept)+('\n' if kept else ''),encoding='utf-8')
 return jsonify(ok=True,removed=len(lines)-len(kept))

@app.post('/api/log/delete-old')
def delete_old_log():
 data=request.get_json(silent=True) or {}
 days=max(1,min(3650,int(data.get('days',30) or 30)))
 if status.get('running'):return jsonify(ok=False,error='実行中は削除できません（ワーカーが同じファイルへ書いています）'),409
 cutoff=datetime.now().timestamp()-(days*86400)
 p=LOG_PATH
 if not p.exists():return jsonify(ok=True,removed=0,kept=0,days=days)
 lines=p.read_text(encoding='utf-8',errors='replace').splitlines()
 kept=[];removed=0
 for line in lines:
  try:
   ts=datetime.strptime(line[:23],'%Y-%m-%d %H:%M:%S,%f').timestamp()
   if ts<cutoff:
    removed+=1;continue
  except Exception:
   pass
  kept.append(line)
 p.write_text('\n'.join(kept)+('\n' if kept else ''),encoding='utf-8')
 log.info('LOG_RETENTION_DELETE days=%s removed=%s kept=%s',days,removed,len(kept))
 return jsonify(ok=True,removed=removed,kept=len(kept),days=days)

@app.post('/api/pick-file')
def pick_file():
 try:
  import tkinter as tk
  from tkinter import filedialog
  data=request.get_json(silent=True) or {}; initial=str(resolve_path(data.get('initial') or str(BASE))); types=data.get('types') or [['すべてのファイル','*.*']]
  root=tk.Tk(); root.withdraw(); root.attributes('-topmost',True)
  path=filedialog.askopenfilename(initialdir=str(Path(initial).parent if Path(initial).suffix else Path(initial)),filetypes=[tuple(x) for x in types])
  root.destroy(); return jsonify(path=path)
 except Exception as e:return jsonify(error=str(e)),500

@app.post('/api/pick-folder')
def pick_folder():
 try:
  import tkinter as tk
  from tkinter import filedialog
  data=request.get_json(silent=True) or {}; initial=str(resolve_path(data.get('initial') or str(BASE))); root=tk.Tk(); root.withdraw(); root.attributes('-topmost',True); path=filedialog.askdirectory(initialdir=initial); root.destroy(); return jsonify(path=path)
 except Exception as e:return jsonify(error=str(e)),500

@app.post('/api/open-path')
def open_path():
 # 出力先は社内共有パス(UNC)やローカルパスであり、Webアドレスではない。
 # ブラウザーを遷移させず、このサーバー(ローカルPC)側でエクスプローラーを開く。
 data=request.get_json(silent=True) or {}; raw=str(data.get('path') or '').strip()
 if not raw:return jsonify(ok=False,error='出力先が指定されていません'),400
 if os.name!='nt':return jsonify(ok=False,error='フォルダーを開けるのはWindowsのみです'),400
 try:
  target=resolve_path(raw)
  # ファイル指定ならその親フォルダーを開く。存在しない場合は明示エラー(Web遷移させない)。
  if target.is_file():target=target.parent
  if not target.exists():return jsonify(ok=False,error=f'出力先が見つかりません: {target}'),404
  os.startfile(str(target))  # type: ignore[attr-defined]
  log.info('OPEN_PATH path=%s resolved=%s',raw,target)
  return jsonify(ok=True,resolved=str(target))
 except Exception as e:
  log.exception('OPEN_PATH_FAILED path=%s',raw); return jsonify(ok=False,error=str(e)),500

@app.post('/api/path-convert')
def path_convert():
 data=request.get_json(silent=True) or {}; value=str(data.get('value') or '').strip(); mode=data.get('mode','absolute')
 if not value:return jsonify(value='',resolved=str(BASE),base=str(BASE),is_relative=False)
 resolved=resolve_path(value).resolve()
 if mode=='relative':
  try:converted='.\\'+str(resolved.relative_to(BASE.resolve())).replace('/','\\')
  except ValueError:
   try:converted=os.path.relpath(str(resolved),str(BASE.resolve())).replace('/','\\')
   except ValueError:return jsonify(error='別ドライブまたはUNC経路のため、このアプリ基準の相対パスへ変換できません。絶対パスを使用してください。'),400
 else:converted=str(resolved)
 return jsonify(value=converted,resolved=str(resolved),base=str(BASE.resolve()),is_relative=not Path(converted).is_absolute())

@app.get('/api/output-capabilities')
def output_capabilities():
 result={'sqlite3':{'ok':True,'detail':'Python標準機能'},'csv':{'ok':True,'detail':'Python標準機能'},'txt':{'ok':True,'detail':'Python標準機能'}}
 try:
  import openpyxl
  result['xlsx']={'ok':True,'detail':'openpyxl '+openpyxl.__version__}
 except Exception as e:result['xlsx']={'ok':False,'detail':'openpyxl未導入: '+str(e)}
 pythoncom=None
 try:
  import pythoncom, win32com.client
  pythoncom.CoInitialize(); conn=win32com.client.Dispatch('ADODB.Connection'); template=resolve_path(load().get('accdb_template','.\\assets\\empty.accdb')); result['accdb']={'ok':template.is_file(),'detail':('テンプレート方式: '+str(template)) if template.is_file() else ('ACCDBテンプレート未配置: '+str(template))}
 except Exception as e:result['accdb']={'ok':False,'detail':'pywin32 COM利用不可: '+str(e)}
 finally:
  if pythoncom:
   try:pythoncom.CoUninitialize()
   except:pass
 return jsonify(result)

@app.get('/api/machine-paths')
def machine_paths():
 return jsonify(**machine_path_view())

@app.get('/api/alerts')
def get_alerts():
 """未確認の知らせと、待機中の取り直し。画面の常設バッジがこれを見る。"""
 with alerts_lock:items=list(alerts)
 worst=('error' if any(x['kind']=='error' for x in items) else
        'warn' if any(x['kind']=='warn' for x in items) else
        'info' if items else '')
 return jsonify(ok=True,alerts=items,count=len(items),worst=worst,retries=retry_view())

@app.post('/api/alerts/ack')
def ack_alerts():
 """読んだ知らせを消す。idを指定しなければ全部。"""
 data=request.get_json(silent=True) or {}
 ids=set(data.get('ids') or [])
 with alerts_lock:
  before=len(alerts)
  if ids:alerts[:]=[x for x in alerts if x['id'] not in ids]
  else:alerts.clear()
  removed=before-len(alerts)
 return jsonify(ok=True,removed=removed,count=len(alerts))

@app.post('/api/retries/cancel')
def cancel_retries():
 """待機中の取り直しを取り消す（原因を直してから自分で流したいとき）。"""
 with retry_lock:
  n=len(retry_waiting);retry_waiting.clear()
 log.info('RETRY_CANCELLED count=%s',n)
 return jsonify(ok=True,cancelled=n)

@app.post('/api/path-check')
def path_check():
 d=request.get_json(force=True); item=d.get('item'); job_id=d.get('job_id'); c=load()
 if item in ('symnavim_conf','symnavim_def','accdb_template'):
  expected={'symnavim_conf':'symnavim.conf','symnavim_def':'symnavim.def','accdb_template':'empty.accdb'}[item]; result=check_path_item(c.get(item,''),'file',expected); result.update(item=item,label=expected)
 elif item=='rne':
  j=next((x for x in c['jobs'] if x['id']==job_id),None)
  if not j:return jsonify(error='対象が見つかりません'),404
  if d.get('value') is not None:j=dict(j,rne_path=d.get('value'),rne=d.get('expected_name') or j.get('rne'))
  rp=resolve_rne_path(j,c); result={'ok':rp.is_file(),'configured':str(j.get('rne_path') or j.get('rne')),'resolved':str(rp),'candidates':[],'needs_reselect':False}; result['candidates']=find_nearby_file(j.get('rne')) if not result['ok'] else []; result['needs_reselect']=not result['ok'] and not result['candidates']; result.update(item='rne',job_id=job_id,label=j.get('rne'))
 else:return jsonify(error='確認対象が不正です'),400
 return jsonify(result)

@app.post('/api/apply-path-suggestion')
def apply_path_suggestion():
 d=request.get_json(force=True); item=d.get('item'); candidate=d.get('candidate'); job_id=d.get('job_id')
 if not candidate or not Path(candidate).is_file():return jsonify(error='修正候補が存在しません'),400
 with settings_sync_lock:
  c=load()
  try:stored=str(Path(candidate).resolve().relative_to(BASE.resolve()))
  except ValueError:stored=str(Path(candidate).resolve())
  if not stored.startswith('.') and not Path(stored).is_absolute():stored='.\\'+stored.replace('/','\\')
  if item in ('symnavim_conf','symnavim_def','accdb_template'):c[item]=stored
  elif item=='rne':
   j=next((x for x in c['jobs'] if x['id']==job_id),None)
   if not j:return jsonify(error='対象が見つかりません'),404
   j['rne_path']=stored; j['rne']=Path(candidate).name
  else:return jsonify(error='修正対象が不正です'),400
  save(c,quiet=True)
 return jsonify(ok=True,path=stored)

@app.get('/api/data-viewer/jobs')
def data_viewer_jobs():
 c=load();items=[]
 for job in c['jobs']:
  path=_viewer_output_path(job,c);items.append({'id':job['id'],'name':job['name'],'format':normalize_output_format(job.get('output_format'),path.name),'exists':path.is_file()})
 return jsonify(items=items)

@app.post('/api/data-viewer/export-pivot')
def data_viewer_export_pivot():
 # 集計表(ピボット)の表示内容をそのままファイルへ書き出す。matrixは1行目がヘッダーの2次元配列。
 # 出力の区切り・文字コードは通常の抽出出力(export_data)と同じ規則へそろえる。
 data=request.get_json(silent=True) or {}
 fmt=normalize_output_format(data.get('format') or 'csv')
 if fmt not in ('csv','txt','xlsx'):return jsonify(ok=False,error=f'集計結果の出力に未対応の形式です: {fmt}'),400
 matrix=data.get('matrix')
 if not isinstance(matrix,list) or not matrix:return jsonify(ok=False,error='出力する集計結果がありません'),400
 if len(matrix)>1048576:return jsonify(ok=False,error=f'行数が多すぎます（{len(matrix):,}行）。絞り込んでから出力してください'),400
 rows=[['' if v is None else str(v) for v in (r if isinstance(r,list) else [r])] for r in matrix]
 width=max(len(r) for r in rows); rows=[r+['']*(width-len(r)) for r in rows]
 headers,body=rows[0],rows[1:]
 stem=re.sub(r'[\\/:*?"<>|]','_',str(data.get('name') or '集計結果')).strip() or '集計結果'
 filename=f'{stem}_{datetime.now().strftime("%Y%m%d_%H%M%S")}{output_extension(fmt)}'
 tmp=LOCAL_RUNTIME/f'pivot_export_{uuid.uuid4().hex}{output_extension(fmt)}'
 try:
  if fmt=='xlsx':
   write_xlsx_direct(tmp,'集計結果',headers,body)
   ctype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
  else:
   delimiter=',' if fmt=='csv' else '\t'
   with tmp.open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f,delimiter=delimiter,quoting=csv.QUOTE_MINIMAL);w.writerow(headers);w.writerows(body)
   ctype='text/csv; charset=utf-8' if fmt=='csv' else 'text/plain; charset=utf-8'
  payload=tmp.read_bytes()
 except Exception as e:
  log.exception('PIVOT_EXPORT_FAILED format=%s rows=%s',fmt,len(rows));return jsonify(ok=False,error=str(e)),500
 finally:
  try:tmp.unlink()
  except OSError:pass
 from urllib.parse import quote
 log.info('PIVOT_EXPORT format=%s rows=%s columns=%s size=%s file=%s',fmt,len(body),width,len(payload),filename)
 response=app.make_response(payload)
 response.headers['Content-Type']=ctype
 response.headers['Content-Disposition']="attachment; filename*=UTF-8''"+quote(filename)
 response.headers['Cache-Control']='no-store'
 return response

@app.get('/api/data-viewer/<job_id>')
def data_viewer(job_id):
 c=load();job=next((x for x in c['jobs'] if x['id']==job_id),None)
 if not job:return jsonify(ok=False,error='対象が見つかりません'),404
 path=_viewer_output_path(job,c)
 if not path.is_file():return jsonify(ok=False,error=f'出力ファイルが見つかりません: {path}'),404
 try:
  # 公開先は直接開かない（開くと差し替えられなくなる）。写してから読む。
  fmt,headers,rows,total,total_cols=read_preview_data(read_copy(path),job,limit=VIEWER_MAX_ROWS)
  # 列は何本あっても全部渡す。多いときに控えるのは行のほう（表の形は変えない）。
  keep=viewer_row_budget(len(headers));limited=len(rows)>keep
  if limited:rows=rows[:keep]
  return jsonify(ok=True,name=job['name'],format=fmt,path=str(path),modified=datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec='seconds'),headers=headers,rows=rows,total_rows=total,preview_limit=len(rows),truncated=total is not None and total>len(rows),
                 total_columns=total_cols,row_budget=keep,row_budget_limited=limited)
 except Exception as e:log.exception('DATA_VIEWER_FAILED job=%s path=%s',job.get('name'),path);return jsonify(ok=False,error=str(e)),200

@app.post('/api/validate')
def validate():
 c=load();checks=[]
 def add(group,label,ok,detail,level=None,**extra):
  checks.append({'group':group,'label':label,'ok':bool(ok),'level':level or ('ok' if ok else 'error'),'detail':str(detail),**extra})
 engine=str(c.get('settings',{}).get('extract_engine') or 'api').lower()
 # 選択中の抽出方式に必要な構成だけを必須診断する。未使用方式の不足でNGにしない。
 if engine=='api':
  try:
   from navigator_api import candidate_dlls,pe_bits
   import struct
   pybits=struct.calcsize('P')*8
   roots=dll_search_roots(c)
   dll_candidates=candidate_dlls(resolve_path(c.get('symnavi_exe','')),resolve_path(c.get('navigator_api_dll')) if c.get('navigator_api_dll') else None,[BASE/'Config'/'NAVIAP',BASE/'NAVIAP'],search_roots=roots)
   found=[x for x in dll_candidates if x.is_file()]
   usable=[x for x in found if pe_bits(x)==pybits];selected=usable[0] if usable else None
   # bit数の話は必ず添える。別のPCで詰まるのはたいていここ。
   if selected:detail=f'利用候補: {selected} / このPCで必要なのは {pybits}bit版 / 一致'
   elif found:detail='見つかったDLLは%s。必要なのは%dbit版です'%(', '.join(sorted({str(pe_bits(x) or "不明")+"bit" for x in found})),pybits)
   else:detail='検索範囲にSymNaviA.dllがありません（範囲: %s）'%(' / '.join(str(x) for x in roots))
   add('実行環境','Navigator API DLL',bool(selected),detail,configured=c.get('navigator_api_dll',''),item='navigator_api_dll',candidates=[str(x) for x in found],needs_reselect=not bool(selected))
   req=_dll_requirement(c);missing=req.get('runtime_missing') or []
   add('実行環境','Visual C++ ランタイム',not missing,('不足: '+', '.join(missing)+f'（{pybits}bit版の再頒布可能パッケージが必要です）') if missing else '必要な%dbit版ランタイムは揃っています: %s'%(pybits,', '.join(req.get('runtime') or [])),item='vc_runtime')
  except Exception as e:add('実行環境','Navigator API DLL',False,e,item='navigator_api_dll',candidates=[],needs_reselect=True)
 else:
  exe=resolve_path(c.get('symnavi_exe',''));add('実行環境','SymNavi.exe',exe.is_file(),exe,configured=c.get('symnavi_exe',''))
  # DDEはpywin32のdde拡張が要る。無ければ実行開始直後に必ず失敗するので、事前に出す。
  try:
   import win32ui,dde  # noqa: F401
   add('実行環境','pywin32のDDE機能',True,'win32ui / dde を読み込めます')
  except Exception as e:
   add('実行環境','pywin32のDDE機能',False,f'読み込めません: {e}（pip install pywin32 が必要です）',item='pywin32')
 # 共通接続ファイル。いま要るかどうかは path_setting_roles に合わせる（画面の役割表示と食い違わせない）。
 conn_roles=path_setting_roles(c)
 for label,key in [('symnavim.conf','symnavim_conf'),('symnavim.def','symnavim_def')]:
  role,why=conn_roles.get(key,('required',''))
  raw=str(c.get(key) or '').strip();p=resolve_path(raw) if raw else None
  if role=='unused':add('接続設定',label,True,why,level='ok',configured=raw,item=key)
  elif raw:add('接続設定',label,bool(p and p.is_file()),p if (p and p.is_file()) else f'{p}（{why}）',configured=raw,item=key,candidates=find_nearby_file(label) if p and not p.is_file() else [],needs_reselect=bool(p and not p.is_file()))
  else:add('接続設定',label,role!='required',f'未設定です。{why}',level='error' if role=='required' else 'warning',item=key,needs_reselect=role=='required')
 # RNE基本フォルダーは補助設定。各ジョブの実ファイルが解決できればフォルダー不足をNGにしない。
 configured_root=resolve_path(c.get('rne_folder','.\\rne'))
 standard_roots=[configured_root,BASE/'Config'/'rne',BASE/'config'/'rne',BASE/'rne']
 existing_roots=[]
 for root in standard_roots:
  try:
   if root.is_dir() and str(root).lower() not in [str(x).lower() for x in existing_roots]:existing_roots.append(root)
  except OSError:pass
 # 固定長テキストの対象はRNEを持たない。ここで一緒に見ると「RNEがありません」と
 # 出続けて、直しようのない赤が並ぶ。入力の種類で分けて、それぞれの入口を確かめる。
 job_results=[]
 for j in c.get('jobs',[]):
  # 結合(join)も手元のファイルから作るので、RNEは持たない。text だけを外していたため、
  # 結合対象が1件でもあると既定の .\rne が「ありません」として error で積まれ、診断は
  # 永久に「修正が必要な項目があります」になっていた（しかも利用者には直しようがない）。
  if normalize_job_source(j.get('source')) in ('text','join'):continue
  rp=resolve_rne_path(j,c);exists=rp.is_file();job_results.append((j,rp,exists))
 all_jobs_ok=all(x[2] for x in job_results) if job_results else True
 root_ok=bool(existing_roots) or all_jobs_ok
 if existing_roots:root_detail='利用可能: '+' / '.join(str(x) for x in existing_roots)
 elif all_jobs_ok and job_results:root_detail='基本フォルダーは未検出ですが、登録済みRNEはすべて個別パスで解決できるため処理可能です。'
 elif not job_results:root_detail='対象未登録のため基本フォルダーは任意です。'
 else:root_detail=f'設定先がありません: {configured_root}。未解決RNEがあるため基本フォルダーまたは個別パスを修正してください。'
 add('RNE配置','RNE参照構成',root_ok,root_detail,level='ok' if root_ok else 'error',configured=c.get('rne_folder',''),item='rne_folder',needs_reselect=not root_ok)
 for j,rp,exists in job_results:
  candidates=find_nearby_file(j.get('rne') or rp.name) if not exists else []
  add('RNE配置',j.get('name','対象')+' RNE',exists,rp,configured=j.get('rne_path'),item='rne',job_id=j.get('id'),candidates=candidates,needs_reselect=not exists and not candidates)
  op=resolve_path(j.get('output_folder') or c.get('default_output_folder','.\\output'))
  add('出力先',j.get('name','対象')+' 出力先',op.is_dir(),op,item='')
 # 結合の対象。要るのは「繋ぎ方（結合マスタ）」と、そこに並ぶファイルが在ること。
 for j in c.get('jobs',[]):
  if normalize_job_source(j.get('source'))!='join':continue
  rp=find_join_recipe(j.get('recipe_id'))
  if not rp:
   add('ファイル結合',j.get('name','対象')+' 結合マスタ',False,'結合マスタが選ばれていません（または削除されています）',job_id=j.get('id'),item='')
  else:
   bad=validate_join_recipe(rp)
   missing=[x['path'] for x in rp['sources'] if not resolve_join_path(x['path'],c).is_file()]
   ok_all=not bad and not missing
   add('ファイル結合',j.get('name','対象')+' 結合マスタ',ok_all,
       (f'{rp["name"]} / {len(rp["sources"])}ファイル / つなぎ目 {len(rp["joins"])}' if ok_all
        else '／'.join(bad+([f'見つからないファイル: '+'、'.join(missing[:3])] if missing else []))),
       job_id=j.get('id'),item='',needs_reselect=bool(missing))
  op=resolve_path(j.get('output_folder') or c.get('default_output_folder','.\\output'))
  add('出力先',j.get('name','対象')+' 出力先',op.is_dir(),op,item='')
 # 固定長テキストの対象。要るのは「読むファイル」と「切り方（読取マスタ）」の2つだけ。
 for j in c.get('jobs',[]):
  if normalize_job_source(j.get('source'))!='text':continue
  tp=resolve_text_path(j,c);ok=tp.is_file()
  add('テキスト読取',j.get('name','対象')+' 読取ファイル',ok,tp if ok else f'{tp}（見つかりません）',
      configured=j.get('text_path'),item='text_path',job_id=j.get('id'),needs_reselect=not ok)
  lay=find_text_layout(j.get('layout_id'))
  if not lay:
   add('テキスト読取',j.get('name','対象')+' 読取マスタ',False,'読取マスタが選ばれていません（または削除されています）',job_id=j.get('id'),item='')
  else:
   bad=validate_text_layout(lay)
   add('テキスト読取',j.get('name','対象')+' 読取マスタ',not bad,
       f'{lay["name"]} / {len(lay["columns"])}列 / {text_layout_width(lay)}{"バイト" if lay["unit"]=="byte" else "文字"}' if not bad else '／'.join(bad),
       job_id=j.get('id'),item='')
   # 読取マスタは列の数が先に分かる。出す形式に入らないなら、実行してから落ちる前に言う。
   # 入らない形式はACCESS(255)くらいで、たいていは何も出ない検査になる。
   n=len(lay['columns'])
   for f in [normalize_output_format(j.get('output_format'),j.get('output_file'))]+job_extra_formats(j):
    lim=output_column_limit(f)
    if lim is not None and n>lim:
     add('テキスト読取',j.get('name','対象')+f' 列数（{OUTPUT_FORMAT_LABEL.get(f,f)}）',False,
         f'{OUTPUT_FORMAT_LABEL.get(f,f)}は{lim:,}列までですが、読取マスタは{n:,}列あります。列数に上限の無いCSV／TXTで出すか、読取マスタの列を減らしてください',
         job_id=j.get('id'),item='')
  op=resolve_path(j.get('output_folder') or c.get('default_output_folder','.\\output'))
  add('出力先',j.get('name','対象')+' 出力先',op.is_dir(),op,item='')
 # 同時出力を設定している対象は、1回の実行で何ができるのかをそのまま出す。
 for j in c.get('jobs',[]):
  extras=job_extra_formats(j)
  if not extras:continue
  try:names=' / '.join(x['file'] for x in job_output_plan(j,c))
  except Exception:names=''
  add('出力先',j.get('name','対象')+' 同時出力',True,
      f'1回の抽出から {len(extras)+1}形式を作ります: {names}',item='')
 # 出力形式に応じて必要なテンプレートだけを検査（同時出力のACCDBも含める）。
 accdb_jobs=[j for j in c.get('jobs',[]) if normalize_output_format(j.get('output_format'),j.get('output_file'))=='accdb'
             or 'accdb' in job_extra_formats(j)]
 if accdb_jobs:
  ap=resolve_path(c.get('accdb_template','.\\assets\\empty.accdb'));add('変換環境','ACCDB空テンプレート',ap.is_file(),ap,configured=c.get('accdb_template',''),item='accdb_template',needs_reselect=not ap.is_file())
 counts={'ok':sum(1 for x in checks if x['level']=='ok'),'warning':sum(1 for x in checks if x['level']=='warning'),'error':sum(1 for x in checks if x['level']=='error')}
 return jsonify(ok=counts['error']==0,checks=checks,counts=counts,engine=engine,summary=('実行可能です' if counts['error']==0 else '修正が必要な項目があります'),search_scope='設定値、標準配置 Config\\rne、個別RNEパスを統合して判定')

@app.get('/api/instance')
def instance_info():
 # versionは版だけ、build_versionは版＋ビルド名。起動待ちモーダルは前者を出す（差し込んだ値と
 # 同じ形にして、サーバーが立った瞬間に表示が変わって見えないようにする）。
 r=app.make_response(jsonify(app=APP_ID,instance_id=INSTANCE_ID,display_name=APP_NAME,version=APP_VERSION,build_version=BUILD_VERSION,pid=os.getpid(),port=PORT,path=str(BASE)))
 # 起動待ちモーダル(loading.html)がfile://から状態を確認できるよう、ローカル情報に限りCORSを許可する。
 r.headers['Access-Control-Allow-Origin']='*'; r.headers['Cache-Control']='no-store'; return r

@app.get('/api/docs')
def docs_list():
 out=[]
 for d in DOCS:
  path=docs_dir()/d['file']
  try:st=path.stat()
  except OSError:st=None
  out.append({'id':d['id'],'title':d['title'],'summary':d['summary'],'source':d['source'],
              'file':d['file'],'path':str(path),'available':bool(st),
              'bytes':st.st_size if st else 0,
              'updated_at':datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d %H:%M') if st else ''})
 return jsonify(ok=True,docs=out)

@app.get('/api/docs/<doc_id>')
def docs_read(doc_id):
 d=next((x for x in DOCS if x['id']==doc_id),None)
 if not d:
  log.warning('DOC_UNKNOWN id=%s',doc_id)
  return jsonify(ok=False,error='登録されていない仕様書です'),404
 path=docs_dir()/d['file']
 if not path.is_file():
  log.warning('DOC_MISSING id=%s path=%s',doc_id,path)
  return jsonify(ok=False,error=f'{path} が見つかりません。配布物に含まれているか確認してください。'),404
 try:text=path.read_text(encoding='utf-8')
 except Exception as e:
  log.exception('DOC_READ_FAILED id=%s path=%s',doc_id,path)
  return jsonify(ok=False,error=str(e)),500
 log.info('DOC_READ id=%s bytes=%s lines=%s',doc_id,len(text.encode('utf-8')),text.count(chr(10))+1)
 return jsonify(ok=True,id=d['id'],title=d['title'],summary=d['summary'],source=d['source'],
                path=str(path),text=text)

@app.get('/api/version')
def version_info():
 return jsonify(version=APP_VERSION,build_version=BUILD_VERSION,title=APP_VERSION_TITLE,released_at=APP_RELEASED_AT,changelog=CHANGELOG)

@app.post('/api/browser-closing')
def browser_closing():
 # pagehideはタブ/ブラウザー終了だけでなく再読込等でも発生する。client_idごとに終了候補へ入れ、猶予中の復帰で取り消す。
 global browser_closed_explicit,browser_closing_at
 data=request.get_json(silent=True) or {};app_id=str(data.get('app_id') or '');client_id=str(data.get('client_id') or request.form.get('client_id') or request.headers.get('X-Heartbeat-Client') or '')[:80]
 if app_id!=APP_ID or not client_id:return jsonify(ok=False,error='アプリタブ識別情報が不正です'),400
 now=time.time()
 with heartbeat_lock:
  browser_closed_explicit=True;browser_closing_at=now
  previous=heartbeat_clients.get(client_id,{})
  heartbeat_clients[client_id]={'app_id':APP_ID,'instance_id':INSTANCE_ID,'last_seen':float(previous.get('last_seen') or now),'user_agent':request.headers.get('User-Agent','')[:160],'closing_at':now,'recovered_count':int(previous.get('recovered_count') or 0)}
 log.info('BROWSER_CLOSE_CANDIDATE client_id=%s grace=%ss',client_id,CLOSE_GRACE_SECONDS)
 return jsonify(ok=True,client_id=client_id,grace_seconds=CLOSE_GRACE_SECONDS)


# ==== 読取マスタ（固定長テキストの切り方）=================================
# RNEの登録と同じ扱いにするための受け口。マスタは対象と別に持ち、複数の対象で使い回す。
@app.get('/api/text-layouts')
def text_layouts_list():
 items=load_text_layouts();c=load()
 for x in items:
  x['width']=text_layout_width(x)
  x['used_by']=text_layout_usage(x['id'],c)
  x['overlaps']=[[a,b] for a,b in text_layout_overlaps(x)]
  x['gaps']=[[a,b] for a,b in text_layout_gaps(x)]
 return jsonify(ok=True,items=items,
                encodings=[{'value':k,'label':TEXT_ENCODING_LABEL.get(k,k)} for k in TEXT_ENCODINGS],
                units=[{'value':k,'label':TEXT_UNIT_LABEL.get(k,k)} for k in TEXT_UNITS],
                trims=[{'value':k,'label':TRIM_LABEL.get(k,k)} for k in TRIM_MODES],
                types=[{'value':k,'label':COLUMN_TYPE_LABEL.get(k,k),'note':COLUMN_TYPE_NOTE.get(k,'')} for k in COLUMN_TYPES],
                stamp_formats=list(STAMP_FORMAT_SAMPLES),max_scale=MAX_SCALE,
                date_format_default=DATE_FORMAT_DEFAULT,datetime_format_default=DATETIME_FORMAT_DEFAULT,
                # どの形式なら型をファイルに持てるか。画面で「CSVでは形が揃うだけ」と言うために要る。
                type_formats={k:format_keeps_types(k) for k in OUTPUT_FORMAT_LABEL})

@app.post('/api/text-layouts')
def text_layouts_save():
 d=request.get_json(force=True) or {}
 try:saved=save_text_layout(d)
 except ValueError as e:return jsonify(ok=False,error=str(e)),400
 except Exception as e:
  log.exception('TEXT_LAYOUT_SAVE_FAILED');return jsonify(ok=False,error=str(e)),500
 return jsonify(ok=True,layout=dict(saved,width=text_layout_width(saved)))

@app.delete('/api/text-layouts/<layout_id>')
def text_layouts_delete(layout_id):
 try:gone=delete_text_layout(layout_id)
 except ValueError as e:return jsonify(ok=False,error=str(e)),409
 return jsonify(ok=bool(gone),error='' if gone else 'その読取マスタはありません')


# ---- マスタをEXCELで出し入れする ------------------------------------------
# JSONは機械には正しいが、人には読めない。読取マスタの列定義は300行を超えることが
# あり、そういうものは実際にはEXCELで作られている（仕様書がEXCELなのだから当然）。
# だから、そのまま開いて直して返せる形でも出し入れできるようにする。

def _from_label(mapping,default='',fallback=None):
 """画面の言葉でも、中の言葉でも受ける取り出し口を作る。

 EXCELを直すのは人なので、「整数」と書いても "integer" と書いても通るのが当たり前。
 どちらか片方しか受けないのは、こちらの都合を人に押しつけているだけ。

 対応表のどれにも当たらなかったときは fallback へ渡す。decimal・float のような
 別の言い方をそこで拾えるようにするため ―― ここで既定へ倒してしまうと、書いた人には
 なぜ「文字」になったのかが分からない。
 """
 rev={str(v):k for k,v in mapping.items()}
 def pick(value):
  t=str(value or '').strip()
  if not t:return default
  if t in mapping:return t
  if t in rev:return rev[t]
  # 「Shift-JIS（cp932）」のような、括弧つきの表記も拾う。
  # 1.85.0より前のEXCELにある「小数」は、前方一致で real（桁で位置指定）へ戻る。
  for k,v in mapping.items():
   if t==str(v) or t.startswith(str(v)) or str(v).startswith(t):return k
  return fallback(t) if fallback else default
 return pick

_TYPE_FROM=_from_label(COLUMN_TYPE_LABEL,'text',fallback=normalize_column_type)
_ENC_FROM=_from_label(TEXT_ENCODING_LABEL,'cp932')
_UNIT_FROM=_from_label(TEXT_UNIT_LABEL,'byte')
_TRIM_FROM=_from_label(TRIM_LABEL,'both')
_JOIN_FROM=_from_label(JOIN_TYPE_LABEL,'inner')

def _book_response(path,name):
 r=send_file(path,as_attachment=True,download_name=name,
             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
 r.headers['Cache-Control']='no-store'
 return r

def _book_tmp(name):
 import tempfile
 d=Path(tempfile.mkdtemp(prefix='navi_book_'))
 return d/name

def _uploaded():
 """持ち込まれたファイルを (バイト列, 名前) で受ける。落としたものも、選んだものも同じ道。"""
 f=request.files.get('file')
 if f:return f.read(),str(f.filename or '')
 return request.get_data() or b'', str(request.args.get('filename') or '')

def _is_xlsx(blob,name):
 # 中身で見分ける。拡張子は当てにならない（名前を変えただけのものが来る）。
 return blob[:2]==b'PK' or str(name).lower().endswith(('.xlsx','.xlsm'))

@app.get('/api/text-layouts/export')
def text_layouts_export_file():
 want=[x for x in (request.args.get('ids') or '').split(',') if x.strip()]
 items=[x for x in load_text_layouts() if not want or x['id'] in want]
 stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
 if str(request.args.get('format') or 'json').lower()=='xlsx':
  # EXCELには「型」も画面と同じ言葉で書く。開いた人が読めなければ意味が無い。
  shown=[dict(x,columns=[dict(c,type=COLUMN_TYPE_LABEL.get(c.get('type'),c.get('type','')))
                         for c in (x.get('columns') or [])],
              encoding=TEXT_ENCODING_LABEL.get(x.get('encoding'),x.get('encoding','')),
              unit=TEXT_UNIT_LABEL.get(x.get('unit'),x.get('unit','')),
              trim=TRIM_LABEL.get(x.get('trim'),x.get('trim',''))) for x in items]
  name=f'読取マスタ_{stamp}.xlsx';tmp=_book_tmp(name)
  navi_book.layouts_to_xlsx(shown,tmp)
  log.info('TEXT_LAYOUT_EXPORT format=xlsx count=%s file=%s',len(items),name)
  return _book_response(tmp,name)
 body=json.dumps(text_layouts_export(items),ensure_ascii=False,indent=1)
 name=f'text-layouts-{stamp}.json'
 r=app.make_response(body)
 r.headers['Content-Type']='application/json; charset=utf-8'
 r.headers['Content-Disposition']=f'attachment; filename="{name}"'
 log.info('TEXT_LAYOUT_EXPORT format=json count=%s file=%s',len(items),name)
 return r

@app.post('/api/text-layouts/import')
def text_layouts_import_file():
 """JSONでもEXCELでも受ける。持ち込む側に形式を選ばせない ―― 中身を見れば分かる。"""
 blob,fname=_uploaded()
 if _is_xlsx(blob,fname):
  tmp=_book_tmp('import.xlsx');tmp.write_bytes(blob)
  try:
   layouts,bad=navi_book.layouts_from_xlsx(tmp,label_to_type=_TYPE_FROM,label_to_encoding=_ENC_FROM,
                                           label_to_unit=_UNIT_FROM,label_to_trim=_TRIM_FROM)
  finally:
   try:shutil.rmtree(tmp.parent,ignore_errors=True)
   except Exception:pass
  # EXCELから来たものも、JSONと同じ検査を通す。入口が2つでも、通す門は1つ。
  checked=[];
  for l in layouts:
   l=normalize_text_layout(l);problems=validate_text_layout(l)
   if problems:bad.append(f'「{l["name"] or "名前なし"}」: '+'／'.join(problems));continue
   checked.append(l)
  layouts=checked
 else:
  layouts,bad=text_layouts_import(blob.decode('utf-8-sig',errors='replace'))
 if not layouts:return jsonify(ok=False,error='／'.join(bad) or '取り込めるマスタがありません'),400
 # 同じ名前があれば置き換える。取り込みのたびに増え続けると、どれが最新か分からなくなる。
 have={x['name']:x['id'] for x in load_text_layouts()}
 added=[];replaced=[]
 for l in layouts:
  if l['name'] in have:l['id']=have[l['name']];replaced.append(l['name'])
  else:added.append(l['name'])
  save_text_layout(l)
 log.info('TEXT_LAYOUT_IMPORT 追加=%s 置き換え=%s 読めなかったもの=%s',added,replaced,len(bad))
 return jsonify(ok=True,added=added,replaced=replaced,skipped=bad)

@app.post('/api/text-preview')
def text_preview_route():
 """下読み。位置が合っているかは、数字を見比べるより切った結果を見るほうが早い。"""
 d=request.get_json(force=True) or {}
 layout=normalize_text_layout(d.get('layout') or {})
 c=load()
 raw=str(d.get('path') or layout.get('sample_path') or '').strip()
 if not raw:return jsonify(ok=False,error='下読みするファイルを選んでください'),400
 path=resolve_text_path({'text_path':raw},c)
 out=preview_text(path,layout,int(d.get('lines') or 12))
 out['width']=text_layout_width(layout)
 return jsonify(out) if out.get('ok') else (jsonify(out),400)

# ==== 結合マスタ（複数ファイルをキーで繋ぐ）===============================
# 読取マスタが「1つのファイルをどう切るか」なら、こちらは「複数をどう繋ぐか」。
@app.get('/api/join-recipes')
def join_recipes_list():
 items=load_join_recipes();c=load()
 for x in items:
  x['used_by']=join_recipe_usage(x['id'],c)
 return jsonify(ok=True,items=items,max_sources=JOIN_MAX_SOURCES,
                types=[{'value':k,'label':JOIN_TYPE_LABEL.get(k,k),'note':JOIN_TYPE_NOTE.get(k,'')}
                       for k in join_types_available()],
                formats=[{'value':k,'label':SOURCE_FORMAT_LABEL.get(k,k)} for k in SOURCE_FORMATS],
                layouts=[{'id':x['id'],'name':x['name']} for x in load_text_layouts()])

@app.post('/api/join-recipes')
def join_recipes_save():
 d=request.get_json(force=True) or {}
 try:saved=save_join_recipe(d)
 except ValueError as e:return jsonify(ok=False,error=str(e)),400
 except Exception as e:
  log.exception('JOIN_RECIPE_SAVE_FAILED');return jsonify(ok=False,error=str(e)),500
 return jsonify(ok=True,recipe=saved)

@app.delete('/api/join-recipes/<recipe_id>')
def join_recipes_delete(recipe_id):
 try:gone=delete_join_recipe(recipe_id)
 except ValueError as e:return jsonify(ok=False,error=str(e)),409
 return jsonify(ok=bool(gone),error='' if gone else 'その結合マスタはありません')

@app.get('/api/join-recipes/export')
def join_recipes_export_file():
 want=[x for x in (request.args.get('ids') or '').split(',') if x.strip()]
 items=[x for x in load_join_recipes() if not want or x['id'] in want]
 stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
 if str(request.args.get('format') or 'json').lower()=='xlsx':
  # 固定長テキストを混ぜている結合では、読取マスタをidで持っている。EXCELには
  # 名前で書く ―― 開いた人にidを見せても、何のことか分からない。
  names={x['id']:x['name'] for x in load_text_layouts()}
  shown=[dict(x,sources=[dict(sx,_layout_name=names.get(sx.get('layout_id'),''))
                         for sx in (x.get('sources') or [])]) for x in items]
  name=f'結合マスタ_{stamp}.xlsx';tmp=_book_tmp(name)
  navi_book.recipes_to_xlsx(shown,tmp,join_label=lambda t:JOIN_TYPE_LABEL.get(t,t))
  log.info('JOIN_RECIPE_EXPORT format=xlsx count=%s file=%s',len(items),name)
  return _book_response(tmp,name)
 body=json.dumps(join_recipes_export(items),ensure_ascii=False,indent=1)
 name=f'join-recipes-{stamp}.json'
 r=app.make_response(body)
 r.headers['Content-Type']='application/json; charset=utf-8'
 r.headers['Content-Disposition']=f'attachment; filename="{name}"'
 log.info('JOIN_RECIPE_EXPORT format=json count=%s file=%s',len(items),name)
 return r

@app.post('/api/join-recipes/import')
def join_recipes_import_file():
 """JSONでもEXCELでも受ける。読取マスタ側と同じ約束。"""
 blob,fname=_uploaded()
 if _is_xlsx(blob,fname):
  tmp=_book_tmp('import.xlsx');tmp.write_bytes(blob)
  ids={x['name']:x['id'] for x in load_text_layouts()}
  try:
   recipes,bad=navi_book.recipes_from_xlsx(tmp,label_to_join=_JOIN_FROM,
                                           name_to_layout=lambda n:ids.get(str(n or '').strip(),''))
  finally:
   try:shutil.rmtree(tmp.parent,ignore_errors=True)
   except Exception:pass
  checked=[]
  for r in recipes:
   r=normalize_join_recipe(r);problems=validate_join_recipe(r)
   if problems:bad.append(f'「{r["name"] or "名前なし"}」: '+'／'.join(problems));continue
   checked.append(r)
  recipes=checked
 else:
  recipes,bad=join_recipes_import(blob.decode('utf-8-sig',errors='replace'))
 if not recipes:return jsonify(ok=False,error='／'.join(bad) or '取り込める結合マスタがありません'),400
 have={x['name']:x['id'] for x in load_join_recipes()}
 added=[];replaced=[]
 for r in recipes:
  if r['name'] in have:r['id']=have[r['name']];replaced.append(r['name'])
  else:added.append(r['name'])
  save_join_recipe(r)
 log.info('JOIN_RECIPE_IMPORT 追加=%s 置き換え=%s 読めなかったもの=%s',added,replaced,len(bad))
 return jsonify(ok=True,added=added,replaced=replaced,skipped=bad)

@app.post('/api/join-source')
def join_source_probe():
 """1つのファイルの見出しと件数だけを読む。キーを選ぶ材料になる。

 結合そのものより先に、まず「そのファイルに何という列があるか」が要る。
 全部を読むと重いので、先頭だけを見て列と件数の見当を返す。

 場所が決まっていて実物があるなら、先に一度だけローカルへ写し、以降はその写しを
 読む。組み立てのあいだ、共有フォルダーへ何度も往復しないため。
 """
 import navi_localcopy
 d=request.get_json(force=True) or {}
 c=load()
 src=dict(d or {});raw=str(src.get('path') or '').strip()
 if not raw:return jsonify(ok=False,error='ファイルの場所を入れてください'),400
 fmt=str(src.get('format') or '').strip().lower()
 origin=resolve_join_path(raw,c)
 copy_info={'ok':False,'note':''}
 read_path=origin
 if origin.is_file():
  local,copy_info=navi_localcopy.sample(origin,fmt or origin.suffix.lstrip('.'))
  if local:read_path=local
 if fmt=='fixed':
  lay=find_text_layout(src.get('layout_id'))
  if not lay:return jsonify(ok=False,error='読取マスタを選んでください'),400
  pv=preview_text(read_path,lay,lines=5)
  if not pv.get('ok'):return jsonify(ok=False,error=pv.get('error') or '読み取れませんでした'),400
  return jsonify(ok=True,columns=pv['headers'],rows=None,sample=pv['rows'][:3],
                 path=str(origin),format='fixed',copy=copy_info,
                 note='固定長テキスト（件数は実行時に数えます）')
 if not origin.is_file():return jsonify(ok=False,error=f'ファイルがありません: {origin}'),400
 path=read_path
 try:
  job={'output_format':fmt or path.suffix.lstrip('.'),'table':src.get('table') or '','sheet':src.get('sheet') or ''}
  # 結合元の列は一覧から選ぶためのもの。表示の都合で切ると選べない列ができる。
  used,headers,rows,total,_cols=read_preview_data(path,job,limit=3,max_columns=None)
 except Exception as e:
  return jsonify(ok=False,error=f'読み取れませんでした: {e}'),400
 tables=[];sheets=[]
 if used=='sqlite3':
  try:
   with contextlib.closing(sqlite3.connect(path)) as conn:
    tables=[r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name") if not str(r[0]).startswith('_')]
  except Exception:tables=[]
 elif used=='xlsx':
  try:
   from openpyxl import load_workbook
   wb=load_workbook(path,read_only=True);sheets=list(wb.sheetnames);wb.close()
  except Exception:sheets=[]
 return jsonify(ok=True,columns=headers,rows=total,sample=rows,path=str(origin),
                format=used,tables=tables,sheets=sheets,copy=copy_info,
                partial=bool(copy_info.get('mode')=='prefix' and copy_info.get('bytes',0)<copy_info.get('total',0)))

@app.post('/api/join-order')
def join_order_route():
 """この結合の材料を、どの対象が作るのか。

 順番は実行時に自動で決まるが、決まったことが見えないと信用できない。
 組み立てている画面で、そのまま見えるようにする。
 """
 d=request.get_json(force=True) or {}
 c=load()
 recipe=normalize_join_recipe(d.get('recipe') or {})
 owners=navi_order.output_owners(c,lambda j:job_output_paths(j,c))
 out=[]
 for sx in recipe.get('sources') or []:
  if not sx.get('path'):continue
  path=resolve_join_path(sx['path'],c)
  made=[{'id':j['id'],'name':j.get('name',''),'enabled':bool(j.get('enabled'))}
        for j in owners.get(str(Path(str(path)).as_posix()).lower(),[])]
  out.append({'alias':sx.get('alias',''),'name':Path(str(path)).name,'path':str(path),
              'exists':path.is_file(),'made_by':made})
 opt=navi_order.wait_settings(c)
 return jsonify(ok=True,sources=out,
                made_count=len([x for x in out if x['made_by']]),
                wait=opt)

@app.get('/api/join-candidates')
def join_candidates_route():
 """繋ぐ相手の候補。このアプリ自身が作ったファイルを先に並べる。

 場所を手で打たせるのは、いちばん間違えやすく、いちばん確かめにくい。出力先は
 こちらが知っているのだから、まず出す ―― 打つのは、そこに無いものを指すときだけでよい。
 """
 import navi_localcopy
 c=load()
 return jsonify(items=join_candidates(c),cache=navi_localcopy.stats())

@app.post('/api/join-cache/clear')
def join_cache_clear_route():
 import navi_localcopy
 n=navi_localcopy.clear()
 return jsonify(ok=True,removed=n,cache=navi_localcopy.stats())

@app.post('/api/join-keys')
def join_keys_route():
 """つなぎ目1つぶんの、キーの見当。実データの重なりで探す。

 300列を2つ並べて「突き合わせる列を選んでください」は酷なので、名前ではなく
 中身で探して、強い順に出す。決めるのは人だが、探すのは機械の仕事。
 """
 d=request.get_json(force=True) or {}
 c=load()
 recipe,_info=sampled_recipe(d.get('recipe') or {},c)
 i=max(0,int(d.get('index') or 0))
 srcs=recipe.get('sources') or []
 if i+1>=len(srcs):return jsonify(ok=False,error='つなぎ目がありません'),400
 try:
  reader=join_sample_reader(c);lay=join_layouts()
  lh,lr=navi_join.read_source(srcs[i],reader,lay,limit=navi_join.KEY_SAMPLE_ROWS)
  rh,rr=navi_join.read_source(srcs[i+1],reader,lay,limit=navi_join.KEY_SAMPLE_ROWS)
 except Exception as e:
  return jsonify(ok=False,error=f'読み取れませんでした: {e}'),400
 items=suggest_join_keys(lh,lr,rh,rr,limit=int(d.get('limit') or 5))
 log.info('JOIN_KEY_SUGGEST index=%s 左=%s列/%s行 右=%s列/%s行 候補=%s',
          i,len(lh),len(lr),len(rh),len(rr),
          ' | '.join(f'{x["left"]}={x["right"]}({x["score"]:.0%})' for x in items) or 'なし')
 return jsonify(ok=True,items=items,left_rows=len(lr),right_rows=len(rr),
                sampled=len(lr)>=navi_join.KEY_SAMPLE_ROWS or len(rr)>=navi_join.KEY_SAMPLE_ROWS)

@app.post('/api/join-preview')
def join_preview_route():
 """繋いだ結果と、つなぎ目ごとの一致を返す。結合の失敗は静かなので、必ず数えて見せる。

 読むのはローカルの写し。組み立てのあいだ、共有フォルダーへ何度も往復しないため。
 本番の実行はこの道を通らない（写しは先頭だけのことがある）。
 """
 d=request.get_json(force=True) or {}
 c=load()
 recipe,copies=sampled_recipe(d.get('recipe') or {},c)
 out=preview_join_recipe(recipe,join_sample_reader(c),join_layouts(),lines=int(d.get('lines') or 12))
 out['copies']=copies
 cut=[x for x in copies if x.get('mode')=='prefix' and x.get('bytes') and x.get('total') and x['bytes']<x['total']]
 if cut and out.get('ok'):
  out.setdefault('notes',[]).append(
   'ローカルへ写した先頭だけを読んでいます（'+'、'.join(x['alias'] for x in cut)
   +'）。件数はその範囲での数字で、実行のときは元のファイルを最初から最後まで読みます')
 return jsonify(out) if out.get('ok') else (jsonify(out),400)
